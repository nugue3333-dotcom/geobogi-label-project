from __future__ import annotations

import pytest
from openpyxl import Workbook

from barcode_label_automation.excel_reader import read_labels


def test_read_labels_validates_and_maps_rows(tmp_path):
    path = tmp_path / "labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item_code", "item_name", "barcode", "lot_no", "qty", "print_qty"])
    sheet.append(["A-1", "Name", "123456", "LOT1", 3, 2])
    workbook.save(path)

    rows = read_labels(path)

    assert len(rows) == 1
    assert rows[0].item_code == "A-1"
    assert rows[0].qty == 3
    assert rows[0].print_qty == 2


def test_read_labels_rejects_non_positive_print_qty(tmp_path):
    path = tmp_path / "labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item_code", "item_name", "barcode", "lot_no", "qty", "print_qty"])
    sheet.append(["A-1", "Name", "123456", "LOT1", 3, 0])
    workbook.save(path)

    with pytest.raises(ValueError, match="print_qty must be a positive integer"):
        read_labels(path)


def test_read_labels_reports_missing_required_columns(tmp_path):
    path = tmp_path / "labels.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item_code"])
    workbook.save(path)

    with pytest.raises(ValueError, match="Missing required columns"):
        read_labels(path)


@pytest.mark.parametrize("value", [1001, 1001.0, "1000000"])
def test_excel_accepts_large_integral_counts_and_preserves_barcode(tmp_path, value):
    path = tmp_path / "large.xlsx"
    workbook = Workbook()
    workbook.active.append(["item_code", "item_name", "barcode", "lot_no", "qty", "print_qty"])
    workbook.active.append(["A", "Name", "001234", "LOT", 1, value])
    workbook.save(path)
    row = read_labels(path)[0]
    assert row.print_qty == int(value)
    assert row.barcode == "001234"

@pytest.mark.parametrize("value", [1.5, True, "1.5", "NaN", "inf"])
def test_excel_rejects_non_integral_print_counts(tmp_path, value):
    path = tmp_path / "bad.xlsx"
    workbook = Workbook()
    workbook.active.append(["item_code", "item_name", "barcode", "lot_no", "qty", "print_qty"])
    workbook.active.append(["A", "Name", "001234", "LOT", 1, value])
    workbook.save(path)
    with pytest.raises(ValueError, match="print_qty"):
        read_labels(path)

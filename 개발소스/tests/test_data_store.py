from __future__ import annotations

import pytest
from openpyxl import Workbook

from barcode_label_automation.data_store import (
    DB_HEADERS,
    DEFAULT_DB_HEADERS,
    LABEL_HEADERS,
    canonical_db_field,
    load_db_source,
    load_db_rows,
    load_label_rows,
    lookup_barcode,
    save_db_rows,
    save_label_rows,
    sample_accessory_db_rows,
)
from barcode_label_automation.excel_reader import read_labels


def test_label_rows_round_trip_and_match_print_engine(tmp_path):
    path = tmp_path / "print_queue.xlsx"
    rows = [
        {
            "item_code": "A1001",
            "item_name": "SENSOR BRACKET",
            "barcode": "88023502",
            "lot_no": "LOT250531",
            "qty": "100",
            "print_qty": "1",
        }
    ]

    save_label_rows(path, rows)

    assert load_label_rows(path) == rows
    label_rows = read_labels(path)
    assert label_rows[0].barcode == "88023502"
    assert label_rows[0].print_qty == 1


def test_label_rows_preserve_dynamic_db_fields_for_print_engine(tmp_path):
    path = tmp_path / "print_queue.xlsx"
    rows = [
        {
            "item_code": "10000",
            "item_name": "가나",
            "barcode": "123456789",
            "lot_no": "",
            "qty": "1",
            "print_qty": "1",
            "바코드": "123456789",
            "가격": "10000",
            "품명": "가나",
        }
    ]

    save_label_rows(path, rows)
    label_rows = read_labels(path)

    assert label_rows[0].barcode == "123456789"
    assert label_rows[0].source_fields == (("가격", "10,000원"), ("품명", "가나"), ("상품코드", "10000"))


def test_barcode_db_lookup_fills_label_fields(tmp_path):
    path = tmp_path / "barcode_db.xlsx"
    rows = [
        {
            "barcode": "KOR-TEST-001",
            "item_code": "K1001",
            "item_name": "한글품목테스트",
            "lot_no": "LOT-HANGUL",
            "qty": "10",
            "print_qty": "1",
        }
    ]

    save_db_rows(path, rows)
    loaded = load_db_rows(path)
    result = lookup_barcode(loaded, "KOR-TEST-001")

    assert result is not None
    assert result.item_name == "한글품목테스트"
    assert result.qty == "10"


def test_barcode_db_accepts_korean_headers_and_preserves_extra_columns(tmp_path):
    path = tmp_path / "korean_db.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(["바코드", "품명", "가격", "수량", "출력 매수"])
    sheet.append(["001234", "라벨 프린터", "580000", "2", "3"])
    workbook.save(path)

    loaded = load_db_rows(path)
    result = lookup_barcode(loaded, "001234")

    assert loaded[0]["barcode"] == "001234"
    assert loaded[0]["item_name"] == "라벨 프린터"
    assert loaded[0]["가격"] == "580000"
    assert loaded[0]["qty"] == "2"
    assert loaded[0]["print_qty"] == "3"
    assert result is not None
    assert result.item_name == "라벨 프린터"


def test_db_source_returns_only_the_workbook_headers_for_ui(tmp_path):
    path = tmp_path / "korean_db.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(["바코드", "품명", "가격", "수량", "출력 매수"])
    sheet.append(["001234", "라벨 프린터", "580000", "2", "3"])
    workbook.save(path)

    rows, headers = load_db_source(path)

    assert headers == ("바코드", "품명", "가격", "수량", "출력 매수")
    assert rows[0]["품명"] == "라벨 프린터"
    assert rows[0]["item_name"] == "라벨 프린터"
    assert canonical_db_field("품명") == "item_name"
    assert canonical_db_field("가격") is None


def test_db_source_preserves_excel_zero_padded_numeric_barcode(tmp_path):
    path = tmp_path / "numeric_barcode.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(["바코드", "상품코드", "품명", "가격"])
    sheet.append([123, 7, "테스트 상품", 1200])
    sheet["A2"].number_format = "00000000"
    sheet["B2"].number_format = "0000"
    sheet["D2"].number_format = "#,##0"
    sheet.append(["00000999", "0008", "문자 바코드", 900])
    workbook.save(path)

    rows, headers = load_db_source(path)

    assert headers == ("바코드", "상품코드", "품명", "가격")
    assert rows[0]["barcode"] == rows[0]["바코드"] == "00000123"
    assert rows[0]["item_code"] == rows[0]["상품코드"] == "0007"
    assert rows[0]["가격"] == "1200"
    assert rows[1]["barcode"] == "00000999"
    assert load_db_rows(path)[0]["barcode"] == "00000123"


def test_db_source_rejects_numeric_barcode_format_that_cannot_be_reproduced(tmp_path):
    path = tmp_path / "formatted_barcode.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(["바코드", "품명"])
    sheet.append([12345, "테스트 상품"])
    sheet["A2"].number_format = "000-00"
    workbook.save(path)

    with pytest.raises(ValueError, match="2행.*바코드.*텍스트"):
        load_db_source(path)


def test_barcode_db_reports_missing_barcode_header(tmp_path):
    path = tmp_path / "bad_db.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(["품명", "가격"])
    sheet.append(["라벨 프린터", "580000"])
    workbook.save(path)

    with pytest.raises(ValueError, match="barcode 또는 바코드"):
        load_db_rows(path)


def test_missing_files_are_created_with_required_headers(tmp_path):
    queue_path = tmp_path / "print_queue.xlsx"
    db_path = tmp_path / "barcode_db.xlsx"

    assert load_label_rows(queue_path) == []
    created_db_rows = load_db_rows(db_path)
    assert len(created_db_rows) == 100
    assert set(DB_HEADERS).issubset(created_db_rows[0])
    assert created_db_rows[0]["판매가"].endswith("원")
    _rows, headers = load_db_source(db_path)
    assert headers == DEFAULT_DB_HEADERS
    assert set(load_label_rows(queue_path) or [{header: "" for header in LABEL_HEADERS}][0]) == set(LABEL_HEADERS)


def test_sample_accessory_db_is_deterministic_and_has_unique_barcodes():
    rows = sample_accessory_db_rows()

    assert len(rows) == 100
    assert len({row["barcode"] for row in rows}) == 100
    assert rows[0]["item_code"] == "ACC-001"
    assert rows[-1]["item_code"] == "ACC-100"
    assert all(row["판매가"].endswith("원") for row in rows)
    assert all(not {"lot_no", "qty", "print_qty"}.intersection(row) for row in rows)


def test_hibc_space_checksum_and_custom_binding_round_trip(tmp_path):
    from barcode_label_automation.hibc import compose_hibc, resolve_hibc
    value = next(compose_hibc("A123",str(i),"0") for i in range(1000) if compose_hibc("A123",str(i),"0").endswith(" "))
    rows=[dict.fromkeys(LABEL_HEADERS, "")]
    rows[0].update(barcode=value,complete_hibc=value,qty="1",print_qty="1")
    path=tmp_path/"queue.xlsx"
    save_label_rows(path,rows)
    assert load_label_rows(path)[0]["barcode"]==value
    loaded,headers=load_db_source(path)
    assert loaded[0]["barcode"]==loaded[0]["complete_hibc"]==value
    assert resolve_hibc(loaded[0]["barcode"],"complete")==value
    assert read_labels(path)[0].barcode==value


def test_udi_excel_date_cells_resolve_through_saved_template(tmp_path):
    from datetime import datetime
    from barcode_label_automation.udi_editor import build_udi_patch, resolve_udi_value
    workbook=Workbook();sheet=workbook.active
    sheet.append(["barcode","DI","MADE","SER"])
    sheet.append(["placeholder","09506000134376",datetime(2026,10,13),"000001"])
    path=tmp_path/"dates.xlsx";workbook.save(path)
    rows,_=load_db_source(path)
    assert rows[0]["MADE"]=="2026-10-13"
    data={"gtin":{"source_field":"DI"},"manufacture_date":{"source_field":"MADE","enabled":True},"serial":{"source_field":"SER","enabled":True}}
    patch=build_udi_patch({},data,row=rows[0])
    assert resolve_udi_value(patch,rows[0])=="(01)09506000134376(21)000001(11)261013"

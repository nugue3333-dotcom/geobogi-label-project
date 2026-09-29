from __future__ import annotations

from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from barcode_label_automation.job_runner import JobRunnerError, create_customer_backup, export_print_queue, restore_customer_backup, run_job


def _write_labels_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Labels"
    sheet.append(["item_code", "item_name", "barcode", "lot_no", "qty", "print_qty"])
    sheet.append(["A1001", "테스트 품목", "8801234567890", "LOT-1", 12, ""])
    sheet.append(["A1002", "바코드 없음", "", "LOT-2", 3, 1])
    workbook.save(path)


def test_export_print_queue_creates_queue_without_powershell(tmp_path: Path) -> None:
    _write_labels_workbook(tmp_path / "labels.xlsm")

    queue_path = export_print_queue(tmp_path)

    workbook = load_workbook(queue_path, data_only=True)
    try:
        sheet = workbook.active
        assert sheet.max_row == 2
        assert [sheet.cell(1, column).value for column in range(1, 7)] == [
            "item_code",
            "item_name",
            "barcode",
            "lot_no",
            "qty",
            "print_qty",
        ]
        assert sheet.cell(2, 3).value == "8801234567890"
        assert sheet.cell(2, 6).value == "1"
    finally:
        workbook.close()


def test_export_print_queue_reports_missing_required_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["barcode"])
    workbook.save(tmp_path / "labels.xlsm")

    with pytest.raises(JobRunnerError, match="필수 컬럼"):
        export_print_queue(tmp_path)


def test_run_label_job_cmd_prefers_compiled_runner() -> None:
    cmd_text = Path("run_label_job.cmd").read_text(encoding="utf-8")

    assert 'if exist "%~dp0라벨작업실행기.exe"' in cmd_text
    assert '"%~dp0라벨작업실행기.exe" "%MODE%" %QUIET_ARG%' in cmd_text


def test_print_job_quiet_mode_requires_screen_confirmation(tmp_path: Path) -> None:
    with pytest.raises(JobRunnerError, match="Quiet 모드"):
        run_job("Print", install_dir=tmp_path, quiet=True)


def test_print_job_adds_yes_only_after_confirmation(monkeypatch, tmp_path: Path) -> None:
    actions: list[str] = []
    messages: list[tuple[str, str]] = []
    monkeypatch.setattr("barcode_label_automation.job_runner._confirm_print_send", lambda *, quiet: True)
    monkeypatch.setattr("barcode_label_automation.job_runner._invoke_engine", lambda _install, _base, action: actions.append(action))
    monkeypatch.setattr(
        "barcode_label_automation.job_runner._show_message",
        lambda text, *, title, quiet, error=False: messages.append((title, text)),
    )

    run_job("Print", install_dir=tmp_path, quiet=False)

    assert actions == ["--print --yes"]
    assert messages
    assert messages[0][0] == "프린터 전송 완료"


def test_create_customer_backup_includes_data_and_excludes_runtime_outputs(tmp_path: Path) -> None:
    (tmp_path / "config.ini").write_text("[printer]\nbrand=tsc\n", encoding="utf-8")
    (tmp_path / "barcode_db.xlsx").write_bytes(b"db")
    (tmp_path / "print_queue.xlsx").write_bytes(b"queue")
    (tmp_path / "labels.xlsm").write_bytes(b"labels")
    (tmp_path / "templates").mkdir()
    (tmp_path / "templates" / "default_label.json").write_text('{"elements":[]}', encoding="utf-8")
    (tmp_path / "assets" / "images").mkdir(parents=True)
    (tmp_path / "assets" / "images" / "label.png").write_bytes(b"image")
    (tmp_path / "assets" / "label_designer.ico").write_bytes(b"icon")
    (tmp_path / "out").mkdir()
    (tmp_path / "out" / "label_001.tspl").write_text("SIZE 50 mm,40 mm", encoding="utf-8")
    (tmp_path / "label_job_runner.exe").write_bytes(b"exe")
    (tmp_path / "last_run.log").write_text("log", encoding="utf-8")

    backup_path = create_customer_backup(tmp_path)

    assert backup_path.parent == tmp_path / "out"
    with zipfile.ZipFile(backup_path) as archive:
        names = set(archive.namelist())
        assert "config.ini" in names
        assert "barcode_db.xlsx" in names
        assert "print_queue.xlsx" in names
        assert "labels.xlsm" in names
        assert "templates/default_label.json" in names
        assert "assets/images/label.png" in names
        assert "assets/label_designer.ico" not in names
        assert "backup_manifest.txt" in names
        assert "out/label_001.tspl" not in names
        assert "label_job_runner.exe" not in names
        assert "last_run.log" not in names


def test_create_customer_backup_reports_no_data(tmp_path: Path) -> None:
    with pytest.raises(JobRunnerError, match="백업할 설정"):
        create_customer_backup(tmp_path)


def test_restore_customer_backup_restores_allowed_files_and_preserves_pre_restore_backup(tmp_path: Path) -> None:
    source = tmp_path / "source"
    target = tmp_path / "target"
    source.mkdir()
    target.mkdir()
    (source / "config.ini").write_text("[printer]\nbrand=tsc\n", encoding="utf-8")
    (source / "barcode_db.xlsx").write_bytes(b"new-db")
    (source / "templates").mkdir()
    (source / "templates" / "default_label.json").write_text('{"elements":[]}', encoding="utf-8")
    (target / "config.ini").write_text("[printer]\nbrand=bixolon\n", encoding="utf-8")
    (target / "barcode_db.xlsx").write_bytes(b"old-db")

    backup_path = create_customer_backup(source)
    result = restore_customer_backup(target, backup_path)

    assert result.restored_files >= 3
    assert result.pre_restore_backup is not None
    assert result.pre_restore_backup.is_file()
    assert "[printer]\nbrand=tsc\n" == (target / "config.ini").read_text(encoding="utf-8")
    assert (target / "barcode_db.xlsx").read_bytes() == b"new-db"
    assert (target / "templates" / "default_label.json").read_text(encoding="utf-8") == '{"elements":[]}'
    with zipfile.ZipFile(result.pre_restore_backup) as archive:
        assert archive.read("barcode_db.xlsx") == b"old-db"


def test_restore_customer_backup_rejects_zip_without_manifest(tmp_path: Path) -> None:
    backup_path = tmp_path / "bad.zip"
    with zipfile.ZipFile(backup_path, "w") as archive:
        archive.writestr("config.ini", "[printer]\n")

    with pytest.raises(JobRunnerError, match="backup_manifest"):
        restore_customer_backup(tmp_path / "target", backup_path)


def test_restore_customer_backup_rejects_path_traversal(tmp_path: Path) -> None:
    backup_path = tmp_path / "bad.zip"
    with zipfile.ZipFile(backup_path, "w") as archive:
        archive.writestr("backup_manifest.txt", "채움랩 고객 데이터 백업\n")
        archive.writestr("../config.ini", "[printer]\n")

    with pytest.raises(JobRunnerError, match="허용되지 않는 경로"):
        restore_customer_backup(tmp_path / "target", backup_path)


def test_restore_customer_backup_rejects_unexpected_files(tmp_path: Path) -> None:
    backup_path = tmp_path / "bad.zip"
    with zipfile.ZipFile(backup_path, "w") as archive:
        archive.writestr("backup_manifest.txt", "채움랩 고객 데이터 백업\n")
        archive.writestr("print_labels.exe", b"not data")

    with pytest.raises(JobRunnerError, match="복원 대상이 아닌 파일"):
        restore_customer_backup(tmp_path / "target", backup_path)

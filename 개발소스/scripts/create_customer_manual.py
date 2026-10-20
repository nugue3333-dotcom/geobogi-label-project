from __future__ import annotations

from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from verify_customer_manuals import capture_contract_errors


ROOT = Path(__file__).resolve().parents[1]
CUSTOMER_DIR = ROOT / "고객용_실행폴더"
OUTPUT_DOCX = ROOT / "채움랩_라벨출력패키지_고객용_매뉴얼.docx"

FONT_KO = "맑은 고딕"
# Match the formal print-manual palette used by the three PDF guides.
COLOR_NAVY = "171A1E"
COLOR_BLUE = "143E70"
COLOR_BLUE_DARK = "143E70"
COLOR_MUTED = "68747C"
COLOR_GRID = "D9D9D9"
COLOR_HEADER_FILL = "FFFFFF"
COLOR_LIGHT_FILL = "F3F5F7"
COLOR_ACCENT_FILL = "FFFFFF"
COLOR_WARNING_FILL = "FFF7E7"


def set_run_font(run, size_pt: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_KO
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size_pt: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size_pt=size_pt, bold=bold, color=color)


def set_style_font(style, size_pt: float, color: str = "000000", bold: bool = False) -> None:
    style.font.name = FONT_KO
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, left=120, bottom=80, right=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = COLOR_GRID, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_borders(table, color: str = COLOR_GRID, size: str = "8") -> None:
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 0) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    set_table_borders(table)

    existing_grid = tbl.find(qn("w:tblGrid"))
    if existing_grid is not None:
        tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx >= len(widths_dxa):
                continue
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, widths_dxa: list[int], header_rows: int = 1) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, widths_dxa)
    for r_idx, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if r_idx < header_rows:
            marker = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(marker)
        for cell in row.cells:
            if r_idx < header_rows:
                set_cell_shading(cell, COLOR_BLUE)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, size_pt=10.2, bold=True, color="FFFFFF")
            else:
                set_cell_shading(cell, "FFFFFF" if r_idx % 2 else "F3F5F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, size_pt=10.2, color="111827")
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.3


def fill_cell(cell, text: str, bold: bool = False, color: str = "111827", size_pt: float = 10.2) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, size_pt=size_pt, bold=bold, color=color)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(12)
        set_paragraph_font(p, size_pt=18, bold=True, color=COLOR_BLUE)
    elif level == 2:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        set_paragraph_font(p, size_pt=14, bold=True, color=COLOR_BLUE)
    else:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
        set_paragraph_font(p, size_pt=12, bold=True, color=COLOR_BLUE_DARK)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size_pt=11, bold=True, color=COLOR_NAVY)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2, size_pt=11, color="111827")
    else:
        r = p.add_run(text)
        set_run_font(r, size_pt=11, color="111827")


def add_callout(doc: Document, title: str, body: str, fill: str = COLOR_LIGHT_FILL) -> None:
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(3)
    r_title = p_title.add_run(title)
    set_run_font(r_title, size_pt=11, bold=True, color=COLOR_NAVY)
    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.line_spacing = 1.3
    r_body = p_body.add_run(body)
    set_run_font(r_body, size_pt=10.2, color="263241")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("채움랩 라벨 출력 패키지 고객용 매뉴얼 | Rev. 2026.10.20.02")
    set_run_font(run, size_pt=8.5, color=COLOR_MUTED)
    paragraph.add_run(" | ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    paragraph._p.append(page)


def add_cover(doc: Document) -> None:
    doc.add_picture(str(ROOT / "assets" / "brand" / "chaeumlab_logo_header_2x.png"), width=Inches(2.5))
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("채움랩 라벨 출력 패키지")
    set_run_font(r, size_pt=30, bold=True, color="000000")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(26)
    r2 = p2.add_run("고객용 프로그램 매뉴얼")
    set_run_font(r2, size_pt=18, bold=True, color=COLOR_BLUE)

    metadata = doc.add_table(rows=4, cols=2)
    rows = [
        ("문서 용도", "프린터 설정, 라벨 디자인, 상품 엑셀 연결과 일상 출력을 한 흐름으로 익히는 사용 설명서"),
        ("장비 확인", "BIXOLON, TSC, Zebra, SEWOO 브랜드 설정은 실제 모델의 지원 명령과 장착 옵션을 별도로 확인합니다."),
        ("배포 폴더", "고객용_실행폴더"),
        ("문서 버전", "3.0 / 2026.10.20.02"),
    ]
    for row, (label, value) in zip(metadata.rows, rows):
        fill_cell(row.cells[0], label, bold=True, color=COLOR_NAVY)
        fill_cell(row.cells[1], value)
    style_table(metadata, [2700, 6660], header_rows=0)
    for row in metadata.rows:
        set_cell_shading(row.cells[0], COLOR_HEADER_FILL)

    add_callout(
        doc,
        "운영 핵심",
        "처음 설치한 PC에서는 시작하기.cmd menu의 처음 실행 점검 또는 처음실행_점검.cmd를 먼저 실행합니다. "
        "이 점검은 실행 전 점검, .cllabel 저장파일 연결, 출력 파일 생성 테스트, 고객 데이터 백업을 한 번에 수행합니다. "
        "PC 교체 후에는 고객데이터_복원.cmd로 백업 ZIP을 선택한 뒤 01_output_check.cmd로 복원 결과를 확인합니다. "
        "고객환경점검.exe로 배포 폴더 상태를 확인한 뒤 프린터설정.exe로 장비 환경을 저장합니다. "
        "config.ini는 설정 프로그램이 자동으로 갱신합니다. "
        "업데이트 배포본을 다시 실행해도 이미 저장된 고객 프린터 설정은 자동으로 덮어쓰지 않습니다.",
        COLOR_ACCENT_FILL,
    )
    doc.add_page_break()


def add_quick_start(doc: Document) -> None:
    add_heading(doc, "1 빠른 시작", 1)
    add_body(
        doc,
        "처음 설치한 PC에서는 아래 순서대로 준비합니다. 평소에는 시작하기.cmd가 여는 라벨 출력 관리에서 상품을 선택해 인쇄합니다.",
    )

    table = doc.add_table(rows=1, cols=3)
    headers = ("순서", "실행 내용", "완료 기준")
    for cell, header in zip(table.rows[0].cells, headers):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("1", "고객용_실행폴더를 고객 PC의 원하는 위치에 복사합니다.", "폴더 안에 실행 파일, 설정, DB, 템플릿이 함께 있습니다."),
        ("2", "처음실행_점검.cmd를 실행합니다. 시작 메뉴는 시작하기.cmd menu로 열 수 있습니다.", "필수 파일, 설정, 엑셀, .cllabel 파일 연결, 출력 파일 생성 테스트와 고객 데이터 백업을 점검합니다."),
        ("3", "필요하면 시작하기.cmd menu > 버전 정보 보기를 확인합니다.", "지원 문의 시 버전정보.txt와 지원 ZIP을 함께 전달합니다."),
        ("4", "Excel 매크로 방식도 사용할 경우 00_install_trusted_location.cmd를 실행합니다.", "Excel 보안 경고 없이 매크로 버튼을 사용할 준비가 됩니다."),
        ("5", "프린터설정.exe를 열어 브랜드, 연결 방식, 용지 크기, 용지 유형, 인쇄 방식을 입력하고 설정 점검 후 저장합니다.", "config.ini가 자동으로 갱신되고 열린 디자이너·출력 관리에 공통 규격이 반영됩니다."),
        (
            "6",
            "시작하기.cmd로 라벨 출력 관리를 열어 상품을 선택하고 매수를 확인합니다.",
            "선택 항목이 없으면 전체 항목이 대상입니다. 인쇄 매수 선택 창에서 대상 건수와 총 매수를 확인합니다.",
        ),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_data):
            fill_cell(cell, text)
    style_table(table, [900, 4300, 4160])

    add_callout(
        doc,
        "고객 안내 문구",
        "실제 프린터로 보내기 전에는 00_고객PC_실행전점검.cmd와 01_output_check.cmd를 순서대로 실행해 주세요. "
        "처음 실행 점검은 저장파일 연결도 PowerShell 없이 register_label_filetype.cmd로 처리합니다. "
        "01/02 배치파일은 라벨작업실행기.exe를 우선 사용하므로 일반적인 출력 점검과 인쇄 명령 생성은 PowerShell 없이 처리됩니다. "
        "일상 사용 중에도 라벨출력관리.exe의 도구 > 실행 전 점검 메뉴로 out\\customer_preflight_report.txt를 즉시 갱신할 수 있습니다. "
        "업데이트 배포본을 복사해도 기존 config.ini가 있으면 고객 프린터 설정은 유지됩니다. 기본값으로 되돌리고 싶을 때만 백업 후 config.example.ini를 참고하세요. "
        "오류 문의 시 라벨출력관리.exe의 도움말 > 지원 패키지 생성 메뉴로 out\\customer_support_package.zip을 새로 만든 뒤 함께 전달하면 원인 확인이 빠릅니다. "
        "지원 ZIP에는 원본 DB, 인쇄 데이터, labels.xlsm이 포함되지 않습니다.",
        COLOR_WARNING_FILL,
    )


def add_product_screens(doc: Document) -> None:
    """Use only actual source captures from the current integrated release."""
    add_heading(doc, "공통 메뉴와 파일 형식", 1)
    add_body(doc, "파일은 현재 작업과 저장, 보기는 해당 화면의 표시, 도구는 장비와 업무, 도움말은 고객 안내를 다룹니다. 편집·이동과 정밀편집 메뉴는 표시하지 않습니다. 앱별 역할 아이콘으로 편집·출력·설정을 구분하고 도안 문서와 이동용 프로젝트는 별도 파일 아이콘을 사용합니다.")
    add_body(doc, "세 프로그램의 도움말 > 고객용 매뉴얼은 현재 프로그램에 맞는 PDF 설명서를 엽니다. 출력 관리의 빠른 사용안내와 지원 패키지도 도움말 메뉴에서 사용합니다.")
    add_body(doc, "새 도안 기본 확장자는 .cllabel, 이동용 프로젝트는 .clproject입니다. 기존 .gblabel·.gbproject도 열 수 있습니다. 구형 도안은 원래 파일에 저장하며 새 형식으로 옮길 때는 다른 이름으로 저장을 사용합니다. 이동 프로젝트에는 원본 상품 엑셀이 포함되지 않으므로 별도로 옮겨 연결합니다.")
    screens = (
        ("라벨디자이너", "label-designer-final.png", "chaeumlab_designer_icon.png"),
        ("라벨출력관리", "label-manager-final.png", "chaeumlab_manager_icon.png"),
        ("프린터설정", "printer-settings.png", "chaeumlab_settings_icon.png"),
    )
    for title, filename, icon in screens:
        doc.add_page_break()
        add_heading(doc, title + " 전체 화면", 1)
        doc.add_picture(str(ROOT / "assets" / "brand" / icon), width=Inches(0.4))
        path = ROOT / "outputs" / "ui-redesign-preview" / filename
        if not path.exists():
            raise FileNotFoundError("최신 실제 화면 캡처가 없습니다: " + str(path))
        doc.add_picture(str(path), width=Inches(6.5))
        add_body(doc, "2026.10.20.02 실제 프로그램 화면. 창 크기와 Windows 배율에 따라 패널 폭과 표시 위치가 달라질 수 있습니다.")


def add_file_map(doc: Document) -> None:
    add_heading(doc, "2 배포 폴더 구성", 1)
    add_body(doc, "아래 파일은 같은 폴더 안에 있어야 합니다. 고객 사용 중에는 파일 이름을 바꾸거나 삭제하지 않는 것이 원칙입니다.")
    table = doc.add_table(rows=1, cols=3)
    for cell, header in zip(table.rows[0].cells, ("파일", "용도", "고객 사용 여부")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("시작하기.cmd", "더블클릭하면 라벨 출력 관리가 열립니다. 명령줄에 menu를 붙이면 점검·설정·디자인·백업·복원 메뉴가 열립니다.", "일상 사용"),
        ("처음실행_점검.cmd", "실행 전 점검, .cllabel 저장파일 연결, 출력 파일 생성 테스트, 고객 데이터 백업을 순서대로 실행하는 최초 설치용 점검 파일", "최초 설치 시 실행"),
        ("register_label_filetype.cmd", ".cllabel 저장파일 아이콘, 더블클릭 열기, 우클릭 인쇄 연결을 PowerShell 없이 등록", "최초 설치/문제 복구 시 실행"),
        ("버전정보.txt", "패키지명, 배포 생성시간, 고객 첫 실행 순서, 지원 문의 시 전달할 정보를 적은 버전 확인 파일", "지원/문의용"),
        ("고객데이터_백업.cmd", "설정, 상품 DB, 인쇄 데이터, 템플릿, assets/images 도안 이미지를 out 폴더의 ZIP으로 백업", "PC 교체 전 실행"),
        ("고객데이터_복원.cmd", "백업 ZIP에서 설정, 상품 DB, 인쇄 데이터, 템플릿, assets/images 도안 이미지를 현재 PC로 복원", "PC 교체 후 실행"),
        ("labels.xlsm", "품목 입력 및 라벨 출력 버튼이 들어 있는 Excel 파일", "사용"),
        ("라벨작업실행기.exe", "01/02/03/04 배치파일에서 사용하는 작업 실행기", "내부 실행 파일"),
        ("라벨출력엔진.exe", "라벨 명령 생성 및 프린터 전송 프로그램", "직접 실행하지 않음"),
        ("고객환경점검.exe", "필수 파일, 설정, 엑셀, 인쇄 데이터 바코드/매수, 출력 폴더, dry-run 점검과 지원 ZIP 저장", "최초 설치 또는 지원 요청 시 사용"),
        ("release_manifest.json", "배포 파일 구성과 해시 검증용 목록", "지원/검증용"),
        ("배포_파일목록.txt", "고객이 바로 읽을 수 있는 배포 파일 목록", "확인용"),
        ("라벨출력관리.exe", "DB 조회, 출력 데이터 선택, 인쇄 실행, 도구 메뉴의 실행 전 점검과 도움말 메뉴의 사용안내/지원 패키지 생성", "일상 사용"),
        ("라벨디자이너.exe", "빈 라벨·예제로 시작해 개체를 배치하고 .cllabel 파일을 저장하거나 .clproject로 이동용 프로젝트를 내보냅니다.", "양식 제작/수정 시 사용"),
        ("templates\\sample_excel_product.cllabel", "기본 빈 라벨과 별도로 제공하는 상품 엑셀 연결 예제 도안", "처음 디자인 시 사용"),
        ("프린터설정.exe", "프린터 브랜드, 연결 방식, 용지 크기, 용지 유형, 인쇄 방식, 저장 전 점검 설정 프로그램", "설치 담당자 사용"),
        ("config.ini", "프린터 설정 저장 파일", "직접 수정하지 않음"),
        ("00_고객PC_실행전점검.cmd", "고객 환경 점검과 out\\customer_support_package.zip 저장", "최초 설치 시 실행"),
        ("00_install_trusted_location.cmd", "Excel 보안 신뢰 위치 자동 등록", "최초 1회 실행"),
        ("01_output_check.cmd", "프린터 전송 없이 출력 파일만 생성하는 점검용 파일", "문제 점검 시 사용"),
        ("03_open_output_folder.cmd", "생성된 라벨 명령 파일 폴더 열기", "문제 점검 시 사용"),
        ("04_open_last_log.cmd", "마지막 실행 로그 열기", "문제 점검 시 사용"),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_data):
            fill_cell(cell, text)
    style_table(table, [2300, 5100, 1960])


def add_excel_usage(doc: Document) -> None:
    add_heading(doc, "3 Excel 입력 방법", 1)
    add_body(doc, "labels.xlsm의 표 한 줄이 라벨 한 종류입니다. 출력 매수는 print_qty 값으로 결정됩니다.")
    table = doc.add_table(rows=1, cols=4)
    for cell, header in zip(table.rows[0].cells, ("열 이름", "입력 내용", "예시", "비고")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("item_code", "품목 코드", "A1001", "라벨 상단 식별값"),
        ("item_name", "품목명", "센서 브라켓", "한글 입력 가능"),
        ("barcode", "바코드 값", "A1001-250531", "바코드로 실제 인쇄되는 값"),
        ("lot_no", "LOT 번호", "LOT250531", "추적용 LOT 정보"),
        ("qty", "수량", "100", "라벨에 표시되는 기준 수량"),
        ("print_qty", "출력 매수", "1", "해당 행을 몇 장 출력할지 입력"),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_data):
            fill_cell(cell, text)
    style_table(table, [1900, 2500, 2100, 2860])

    add_heading(doc, "출력 흐름", 2)
    workflow = doc.add_table(rows=1, cols=4)
    for cell, header in zip(workflow.rows[0].cells, ("입력", "버튼", "전송", "기록")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    row = workflow.add_row().cells
    for cell, text in zip(
        row,
        (
            "Excel 표 내용을 수정",
            "라벨 출력 클릭",
            "프린터로 즉시 전송",
            "last_run.log와 print_log.xlsx에 결과 저장",
        ),
    ):
        fill_cell(cell, text)
    style_table(workflow, [2340, 2340, 2340, 2340])

    add_callout(
        doc,
        "입력 시 주의",
        "barcode가 비어 있으면 정상 라벨을 만들 수 없습니다. 앞자리 0이 있는 바코드는 Excel에서 문자로 저장하는 것이 가장 안전합니다. "
        "순수한 0 채우기 숫자 서식은 화면 표시값을 읽지만, 복잡한 숫자 서식이나 15자리를 넘는 숫자 값은 확인 후 문자로 고쳐야 합니다. "
        "print_qty는 실제 출력 매수이므로 테스트 후 고객 운영 수량에 맞춰 입력해 주세요.",
    )


def add_daily_operation(doc: Document) -> None:
    add_heading(doc, "4 일상 사용 절차", 1)
    table = doc.add_table(rows=1, cols=2)
    for cell, header in zip(table.rows[0].cells, ("작업", "상세 설명")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("1. 출력 화면 열기", "시작하기.cmd를 더블클릭해 라벨 출력 관리를 엽니다."),
        ("2. 상품 찾기", "상품 엑셀 연결에서 상품 엑셀을 선택하고 상품명·바코드를 검색합니다."),
        ("3. 대상 선택", "인쇄 데이터에서 필요한 행을 직접 선택합니다. 체크한 행이 없으면 전체 항목을 인쇄 대상으로 사용합니다."),
        ("4. 매수와 전송", "인쇄를 눌러 1 이상의 정수 매수를 지정하고 대상 건수를 확인한 뒤 인쇄 시작을 누릅니다."),
        ("5. 실물 확인", "전송 완료 안내 후 프린터에서 라벨 배출·위치·바코드 판독을 확인합니다."),
    ]
    for row_data in rows:
        row = table.add_row().cells
        fill_cell(row[0], row_data[0], bold=True, color=COLOR_BLUE_DARK)
        fill_cell(row[1], row_data[1])
    style_table(table, [2300, 7060])

    add_body(
        doc,
        "중요: 선택 없이 인쇄 버튼을 누르면 전체 항목이 대상입니다. 인쇄 전에는 항상 양의 정수를 입력하는 인쇄 매수 선택 창이 열립니다. 100장·1,000장 같은 임의 상한은 없습니다. 대상·총 매수를 확인하고 시작 전 취소하면 프린터 명령을 보내지 않습니다. 전송 중 중단은 미전송 작업만 멈추며 이미 프린터가 받은 작업의 중단은 보장하지 않습니다. 이번 매수는 임시 큐에만 적용되며 원본 DB와 기본 print_queue.xlsx의 저장값은 바뀌지 않습니다.",
        bold_prefix="중요:",
    )
    add_body(
        doc,
        "인쇄 전 점검: 라벨출력관리.exe는 실제 인쇄 직전에 바코드 누락, 출력 매수 오류, config.ini 프린터 설정 오류를 먼저 확인합니다. 오류가 표시되면 해당 행 또는 프린터 설정을 수정한 뒤 다시 인쇄하세요.",
        bold_prefix="인쇄 전 점검:",
    )
    add_body(
        doc,
        "도움말: 라벨출력관리.exe의 도구 메뉴에서 실행 전 점검을, 도움말에서 빠른 사용안내·고객용 매뉴얼·지원 패키지를 엽니다. 고객 데이터 백업·복원은 파일 메뉴에 있습니다. 라벨 디자인은 상단의 별도 버튼입니다. 01/02 배치파일은 라벨작업실행기.exe를 우선 사용합니다.",
        bold_prefix="도움말:",
    )
    add_body(
        doc,
        "Excel 매크로 사용: labels.xlsm을 사용하는 사업장은 상품 정보를 수정하고 저장한 뒤 라벨 출력 버튼을 사용합니다. 먼저 00_install_trusted_location.cmd로 Excel 신뢰 위치를 등록하세요. 디자이너의 인쇄파일 생성은 프린터로 보내지 않는 사전 확인입니다.",
        bold_prefix="디자이너 출력:",
    )


def add_label_designer_usage(doc: Document) -> None:
    add_heading(doc, "5 라벨 디자이너", 1)
    add_body(doc, "라벨디자이너.exe는 양식을 만들거나 수정할 때 사용합니다. 처음 화면에서 빈 라벨, 예제로 시작, 최근 라벨 열기 중 하나를 고릅니다. 기본 템플릿은 개체가 없는 빈 라벨입니다.")

    table = doc.add_table(rows=1, cols=2)
    for cell, header in zip(table.rows[0].cells, ("순서", "작업 방법")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("1. 시작", "예제로 시작을 누르면 상품 엑셀과 예제 도안이 연결됩니다. 자신의 양식은 새 라벨에서 만들고, 작업 파일은 저장 버튼 또는 Ctrl+S로 .cllabel에 저장합니다."),
        ("2. 크기와 개체", "용지 설정에서 실제 가로·세로(mm)를 저장합니다. 디자이너 상단은 공통 규격을 읽기 전용으로 보여 줍니다. 아이콘과 이름이 함께 표시된 개체 도구에서 텍스트, 1D/2D 바코드, 그림, 박스, 선, 표를 넣습니다. 개체를 더블클릭해 크기와 위치를 확인합니다. 오른쪽은 상품 DB 메뉴만 표시합니다. 상단 좌우 가운데 정렬·상하 가운데 정렬로 가로·세로 중심에 배치합니다."),
        ("3. 상품 연결", "상품 엑셀 연결에서 .xlsx/.xlsm을 열고 개체의 열을 연결합니다. 상품 선택에서 출력할 행만 고릅니다."),
        ("4. 출력 전 확인", "인쇄파일을 먼저 생성하고 미리보기·행·매수를 확인합니다. 경계 밖 개체, 비어 있는 필수 값, 작은 QR·글자 축소 경고를 확인한 뒤 실제 인쇄합니다."),
    ]
    for row_data in rows:
        row = table.add_row().cells
        fill_cell(row[0], row_data[0], bold=True, color=COLOR_BLUE_DARK)
        fill_cell(row[1], row_data[1])
    style_table(table, [2200, 7160])

    add_heading(doc, "편집과 복구", 2)
    add_body(doc, "Ctrl+클릭으로 여러 개체를 선택합니다. Ctrl+C로 복사, Ctrl+V로 붙여넣기, Ctrl+D로 복제합니다. 글자 입력칸에 포커스가 있으면 해당 글자를 복사·붙여넣기합니다. 그림과 DB 연결도 개체 복사에 포함됩니다. 개체는 더블클릭으로 수정하고 좌우·상하 가운데 정렬은 상단에서 실행합니다. 방향키는 1mm, Shift+방향키는 0.1mm씩 이동합니다. 화면 맞춤과 확대·축소는 보기 크기만 바꿉니다.")
    add_body(doc, "다른 규격의 도안을 열면 기존 규격 유지 또는 공통 규격 적용을 선택합니다. 개체가 영역 밖으로 나가는 크기 변경은 취소되고 기존 도안을 유지합니다. Ctrl+Z는 실행취소, Ctrl+Y는 다시실행입니다. 저장하지 않은 변경사항은 파일명과 창 제목의 *로 보입니다. 비정상 종료 후 복구 제안이 나타나면 저장한 원본을 덮어쓰지 않는 작업 사본으로 확인합니다.")

    add_heading(doc, "도안 인식과 이동", 2)
    add_body(doc, "왼쪽 도안 가져오기 카드의 도안 불러오기에서 PNG/JPG/PSD를 선택하고 도안 적용을 누르면 편집 가능한 후보 개체가 생성됩니다. 인식 값 전체 확인으로 검토창을 열고 원본 이미지와 후보를 나란히 비교해 잘못된 글자·바코드를 먼저 수정한 뒤 전체 값 확인을 누르면 한 번에 확정됩니다. 빈 값·임시 바코드가 있으면 전체 확인을 막습니다.")
    add_body(doc, "다른 PC로 옮길 때는 파일 메뉴의 이동 프로젝트 내보내기로 .clproject를 만듭니다. 가져온 PC에서는 상품 엑셀 경로를 다시 연결하고, 그림·글꼴·라벨 크기·인쇄 설정을 확인합니다. 프로젝트 가져오기는 기존 고객 데이터를 덮어쓰지 않습니다. 기존 .gblabel·.gbproject도 읽을 수 있습니다. 구형 도안은 저장 시 원래 파일을 유지하고 새 .cllabel로 바꾸려면 다른 이름으로 저장을 사용합니다.")
    add_body(doc, "BarTender .btw 파일은 직접 열 수 없습니다. 파일명만 .png로 바꾸지 말고 원본 프로그램에서 실제 이미지로 내보내어 참고 도안으로 불러온 뒤 개체와 데이터 열을 다시 연결합니다. 수식·변수·프린터 설정까지 자동 변환되는 기능은 아닙니다.")

    add_callout(
        doc,
        "실제 출력 확인",
        "프린터 전송 완료는 라벨이 실제로 나왔다는 뜻이 아닙니다. 프린터에서 배출·위치·스캔을 확인하세요. 전송 중단으로 결과가 미확인인 항목은 실물을 먼저 확인하고 화면에서 전송됨 또는 재전송 대상으로 직접 결정해야 중복 출력을 줄일 수 있습니다.",
        COLOR_ACCENT_FILL,
    )


def add_installer_settings(doc: Document) -> None:
    add_heading(doc, "6 설치 담당자용 설정", 1)
    add_body(doc, "고객에게 전달하기 전 프린터설정.exe에서 프린터 환경만 맞춰 두면 됩니다. config.ini는 직접 열지 않아도 됩니다.")

    table = doc.add_table(rows=1, cols=3)
    for cell, header in zip(table.rows[0].cells, ("설정 항목", "값", "설명")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        ("브랜드", "BIXOLON / TSC / Zebra / SEWOO(ZPL)", "고객이 사용하는 프린터 브랜드 한 가지를 선택합니다."),
        ("연결 방식", "LAN / USB", "LAN은 IP 직접 전송, USB는 Windows 프린터 큐를 통해 전송합니다."),
        ("인쇄 방식", "감열 / 열전사", "감열은 리본 없음, 열전사는 리본 사용 환경입니다."),
        ("ip, port", "예: 192.168.0.130 / 9100", "network 방식일 때 프린터 IP와 포트를 입력합니다."),
        ("windows_printer_name", "auto 또는 프린터 이름", "windows_raw 방식일 때 Windows 프린터 이름을 지정합니다."),
        ("width_mm, height_mm", "예: 50.5 / 30.25", "실제 라벨 용지 폭·길이(mm). 소수점 지원."),
        ("print_orientation", "normal / landscape / rotate_180 / landscape_180", "세로 / 가로 / 세로 180도 / 가로 180도."),
        ("offset_x_mm, offset_y_mm", "예: 1.25 / -0.5", "내용 전체 보정. +오른쪽 / +아래쪽. 영역 초과는 차단."),
        ("gap_offset_mm, feed_offset_mm", "auto 또는 mm", "BIXOLON 마크-절취선은 SL p4(200도트 예시)이며 p4 제조사 최대값은 확인되지 않아 용지 길이·SL p1 페이지 길이(최대 2432도트) 상한을 사용합니다. TA 절취·커터 정지는 별도 ±100도트이며 영구 저장됩니다. TSC GAP/BLINE 추가 이송은 0~용지 길이, OFFSET은 ±25.4mm입니다. TSC auto는 GAP/BLINE 전체를 생략하고 명시적 0은 선택한 용지 유형과 이송 0을 적용합니다. Zebra ~TA는 ±120단계(600 DPI는 2도트/단계)."),
        ("media_type", "gap / black_mark / continuous", "갭 용지, 블랙마크 용지, 연속 용지 중 실제 라벨지를 선택합니다."),
        ("gap_mm", "예: 3", "갭 용지와 블랙마크 용지는 0보다 크게 입력합니다. 연속 용지는 0을 권장합니다."),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_data):
            fill_cell(cell, text)
    style_table(table, [2300, 2700, 4360])

    add_callout(
        doc,
        "브랜드별 기본값",
        "BIXOLON은 cp949 계열 명령 인코딩을 사용하고, TSC, Zebra, SEWOO(ZPL)는 UTF-8 기반 명령을 사용합니다. "
        "language와 command_encoding은 auto 상태로 두면 brand 값에 맞춰 자동 적용됩니다.",
        COLOR_ACCENT_FILL,
    )
    add_callout(
        doc,
        "저장 전 점검",
        "프린터설정.exe의 설정 점검 버튼은 IP/포트, Windows 프린터 이름, 용지 유형, 간격, DPI, 브랜드별 인쇄후작업 지원 여부를 확인합니다. "
        "오류가 표시되면 저장하지 말고 실제 장비 설정과 라벨지를 먼저 맞춘 뒤 다시 저장하세요.",
        COLOR_WARNING_FILL,
    )
    add_callout(
        doc,
        "저장하지 않은 변경사항",
        "프린터설정.exe에서 값을 바꾸면 창 제목 끝에 * 표시가 붙습니다. 저장하지 않은 상태로 현재 설정을 다시 불러오거나 창을 닫으려 하면 확인창이 표시되므로, 실수로 작업 중인 설정을 잃지 않도록 안내에 따라 선택하세요.",
        COLOR_ACCENT_FILL,
    )


def add_troubleshooting(doc: Document) -> None:
    add_heading(doc, "7 문제 해결", 1)
    table = doc.add_table(rows=1, cols=3)
    for cell, header in zip(table.rows[0].cells, ("증상", "확인할 내용", "조치")):
        fill_cell(cell, header, bold=True, color=COLOR_NAVY)
    rows = [
        (
            "처음 설치 후 정상 여부를 모르겠음",
            "배포 폴더 필수 파일, 설정, 엑셀, 인쇄 데이터 바코드/매수, dry-run",
            "00_고객PC_실행전점검.cmd를 실행해 오류 항목을 확인하고 out\\customer_support_package.zip을 보관합니다.",
        ),
        (
            "Excel에서 매크로 보안 경고가 나옴",
            "신뢰 위치 등록 여부",
            "Excel을 닫고 00_install_trusted_location.cmd를 실행한 뒤 다시 엽니다.",
        ),
        (
            "라벨 출력 버튼을 눌러도 출력되지 않음",
            "인쇄 전 점검 메시지, 프린터 전원, 네트워크, config.ini IP",
            "먼저 바코드 누락, 출력 매수 오류, 프린터 설정 오류 메시지를 수정합니다. 이후 04_open_last_log.cmd로 마지막 로그를 열고 IP와 포트 9100 연결 상태를 확인합니다.",
        ),
        (
            "한글이 깨져 출력됨",
            "config.ini의 brand 값과 실제 프린터 브랜드",
            "브랜드 값을 실제 장비에 맞추고, 프린터가 한글 폰트/코드페이지를 지원하는지 확인합니다.",
        ),
        (
            "바코드가 비어 있거나 스캔되지 않음",
            "Excel의 barcode 열",
            "barcode 값을 비우지 말고 스캐너가 읽을 수 있는 값으로 입력합니다.",
        ),
        (
            "원하는 매수보다 많이 출력됨",
            "Excel의 print_qty 열",
            "행별 print_qty 값을 실제 필요한 매수로 수정합니다.",
        ),
        (
            "인쇄할 항목이 없다는 안내가 나옴",
            "인쇄 데이터의 선택 칸",
            "필요한 행을 직접 체크합니다. 체크가 없으면 전체 항목이 대상이므로 인쇄 매수 선택 창에서 총 매수를 다시 확인합니다.",
        ),
        (
            "도안 인식 뒤 출력이 차단됨",
            "인식 값 검토의 미확인 글자·바코드",
            "원본 이미지와 값을 대조하고 인식 값 검토에서 확인합니다. 임시 바코드 값은 원본으로 교체합니다.",
        ),
        (
            "이전 PC의 도안 그림이나 상품 엑셀이 보이지 않음",
            ".cllabel만 복사했는지, .clproject로 옮겼는지",
            "파일 메뉴의 이동용 프로젝트를 가져오고 새 PC에서 상품 엑셀을 다시 연결합니다. 원본 엑셀 파일은 별도 보관해야 합니다.",
        ),
        (
            "BarTender .btw 파일이 열리지 않음",
            "지원 파일 종류",
            ".btw 직접 열기는 지원하지 않습니다. 원본 프로그램에서 실제 이미지로 내보내거나 인쇄 미리보기를 저장한 뒤 도안 불러오기로 재작성하고 필드·바코드·인쇄 결과를 확인합니다.",
        ),
        (
            "출력 파일만 생성하고 싶음",
            "프린터 전송 전 사전 점검 필요 여부",
            "01_output_check.cmd를 실행한 뒤 out 폴더의 명령 파일을 확인합니다.",
        ),
        (
            "지원 담당자에게 원인 확인을 요청함",
            "점검 보고서, 마지막 실행 로그, 출력 이력",
            "라벨출력관리.exe의 도움말 > 지원 패키지 생성 메뉴를 눌러 out\\customer_support_package.zip을 새로 만든 뒤 전달합니다. "
            "ZIP 안의 environment_summary.txt, file_inventory.txt, release_manifest_summary.txt로 설치 위치와 누락 파일을 확인할 수 있고 원본 DB와 인쇄 데이터는 포함되지 않습니다.",
        ),
        (
            "PC 교체나 재설치 전 데이터 백업",
            "config.ini, DB, 인쇄 데이터, 템플릿, 도안 이미지",
            "라벨출력관리.exe의 파일 > 고객 데이터 백업 또는 고객데이터_백업.cmd를 실행해 out\\chaeumlab_customer_backup_*.zip을 별도 보관합니다.",
        ),
        (
            "PC 교체나 재설치 후 데이터 복원",
            "백업 ZIP, 기존 현재 데이터",
            "고객데이터_복원.cmd에서 백업 ZIP을 선택합니다. 복원 전 현재 데이터는 out 폴더에 pre-restore 백업으로 먼저 저장됩니다.",
        ),
    ]
    for row_data in rows:
        row = table.add_row().cells
        for cell, text in zip(row, row_data):
            fill_cell(cell, text)
    style_table(table, [2700, 2800, 3860])


def add_support(doc: Document) -> None:
    add_heading(doc, "8 고객 지원", 1)
    add_body(doc, "구매처 또는 설치 담당자에게 오류 내용과 아래 점검 자료를 함께 전달합니다. 지원 패키지는 원본 상품 DB를 포함하지 않습니다.")
    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("프로그램", "채움랩 라벨 출력 패키지"),
        ("문의처", "구매처 또는 설치 담당자"),
        ("전달 자료", "버전정보.txt, 오류 화면, out/customer_support_package.zip"),
        ("장비 확인", "프린터 모델명, 연결 방식, 라벨 크기와 DPI, 문제 발생 시간"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        fill_cell(row.cells[0], label, bold=True, color=COLOR_NAVY)
        fill_cell(row.cells[1], value)
    style_table(table, [2300, 7060], header_rows=0)
    for row in table.rows:
        set_cell_shading(row.cells[0], COLOR_HEADER_FILL)


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(ROOT / "assets" / "brand" / "chaeumlab_logo_header_2x.png"), width=Inches(0.85))

    styles = doc.styles
    set_style_font(styles["Title"], 30, "000000", True)
    normal = styles["Normal"]
    set_style_font(normal, 11, "111827")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.35

    set_style_font(styles["Heading 1"], 18, COLOR_BLUE, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(22)
    styles["Heading 1"].paragraph_format.space_after = Pt(12)
    styles["Heading 1"].paragraph_format.line_spacing = 1.25

    set_style_font(styles["Heading 2"], 14, COLOR_BLUE, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(16)
    styles["Heading 2"].paragraph_format.space_after = Pt(8)
    styles["Heading 2"].paragraph_format.line_spacing = 1.25

    set_style_font(styles["Heading 3"], 12, COLOR_BLUE_DARK, True)
    styles["Heading 3"].paragraph_format.space_before = Pt(10)
    styles["Heading 3"].paragraph_format.space_after = Pt(5)
    styles["Heading 3"].paragraph_format.line_spacing = 1.25

    return doc


def main() -> None:
    capture_errors = capture_contract_errors(ROOT)
    if capture_errors:
        raise RuntimeError("최신 실제 화면 캡처 확인 실패: " + "; ".join(capture_errors))
    CUSTOMER_DIR.mkdir(parents=True, exist_ok=True)
    doc = configure_document()
    add_cover(doc)
    add_quick_start(doc)
    add_product_screens(doc)
    add_file_map(doc)
    add_excel_usage(doc)
    add_daily_operation(doc)
    add_label_designer_usage(doc)
    add_installer_settings(doc)
    add_troubleshooting(doc)
    add_support(doc)
    doc.add_page_break()
    doc.add_heading("GS1 응용식별자와 바코드 세부 옵션", level=1)
    doc.add_paragraph("바코드 종류를 선택하고 바코드 세부 옵션에서 해당 종류의 크기, 체크숫자, 오류복원, 버전, 행·열 등 지원 옵션을 설정합니다. GS1 데이터 구성 (AI)에서 AI를 검색하고 고정값 또는 DB 열을 연결한 뒤 데이터 검증과 적용을 누릅니다.")
    doc.add_paragraph("표시문자에는 AI 괄호를 표시하고 실제 데이터에는 필요한 GS 구분자를 자동으로 넣습니다. GTIN 체크숫자, 길이, 날짜와 AI 조합을 검증하며 URL이나 긴 문자열을 임의로 제거하거나 잘라내지 않습니다. 라벨에 실제 도트 크기가 들어가지 않으면 출력을 차단합니다.")
    doc.add_paragraph("80개 선택 항목을 제공하며 BarTender 공식 69개 형식 주제 중 57개에 대응합니다. iQR, 2D Pharmacode, Anker, 원형 바코드 4종, Code 93i, NEC 2 of 5, OPC, PosiCode, Singapore Post는 미지원입니다. BarTender의 모든 참조 옵션과 완전 동등한 제품은 아닙니다. 실제 프린터와 스캐너로 첫 장을 확인하세요.")
    doc.add_paragraph("새 AI와 선택한 GS1 행의 입력 변경은 데이터 검증·적용 전에 자동 반영합니다. 취소하면 원본은 유지됩니다. AI21 일련번호의 필수 조합 오류는 필수 AI 선택에서 01·03·8006 중 해당 상품의 실제 식별번호 또는 DB 열을 추가하여 해결합니다. 자주 사용하는 AI와 한국어 검색을 사용할 수 있습니다.")
    doc.add_page_break()
    doc.add_heading("현재 도안과 반복 출력 연결", level=1)
    doc.add_paragraph("디자이너에서 도안을 저장한 뒤 보기 > 반복 출력 화면 또는 상단 반복 출력을 누릅니다. 출력 관리에서 선택 도안 이름과 현재 도안 미리보기를 확인합니다. 출력 도안 선택으로 .cllabel을 직접 지정할 수 있으며 자동 기본 배치 사용은 프린터 설정의 기본 배치로 변경합니다. 두 모드를 구분하고 출력 대상·매수를 확인하세요.")
    doc.add_paragraph("상품 엑셀 연결에서 사용할 파일과 시트를 고릅니다. 도안에는 연결 파일 경로·시트·열·변경 확인값을 저장하며 엑셀의 전체 고객 행을 복사하지 않습니다. 원본 참조와 가져온 복사본, 연결 시각을 확인하고 파일이 이동하거나 바뀌면 재연결 또는 엑셀 새로고침 후 열 연결을 검증합니다. 전체 데이터 검증에서 행·개체·원인을 한 번에 확인합니다.")
    doc.add_heading("일련번호·날짜·문자 생성", level=2)
    doc.add_paragraph("디자이너 도구 > 값 생성 규칙에서 일련번호, 날짜 / 기한, 문자 붙이기, 부분 추출을 추가합니다. 생성 필드는 {{필드 이름}} 또는 GS1의 DB 열로 연결합니다. 일련번호는 증가·감소, 숫자·영문 대문자·영문 소문자·숫자+영문 순번, 자리수 유지·자동 자리수와 번호당 복사 수를 지정합니다. Z 다음은 AA로 이어집니다. 작업마다 시작·계속 발급을 선택하고 날짜는 작업 날짜 또는 DB 날짜에 일수를 더해 형식을 정합니다. 규칙을 적용한 뒤 도안을 저장합니다.")
    doc.add_paragraph("DB 열을 변경 대상으로 지정하면 1부터 세는 시작 위치와 글자 수만큼 순번으로 바꾸고 나머지 원문을 보존합니다. 예: ABC00042에서 4번째부터 5글자를 00001로 바꾸면 ABC00001입니다. 바꿀 글자 수와 순번 자리수는 같아야 합니다. 4장 미리보기는 번호를 소비하지 않으며 범위 초과와 잘못된 DB 부분 범위는 출력 전에 차단합니다. 인쇄 확인 후 확정한 값과 기준일은 작업 기록으로 보존하며 재개·재인쇄 시 새 날짜나 번호를 자동 계산하지 않습니다. 전송 결과가 불확실하면 실제 출력물을 먼저 확인하고 기존 값 재인쇄와 새 번호 발급을 구분하세요.")
    doc.add_paragraph("고객 데이터 백업에는 out/value_jobs/ledger.json의 발급·재인쇄 기록을 포함합니다. 나머지 out 로그와 지원 ZIP은 포함하지 않습니다. 발급 중 잠금, 기록 손상, 현재 번호보다 오래된 백업, 누락된 작업 기록은 안전을 위해 복원을 차단합니다. 기록을 지워 우회하지 말고 최신 백업과 실제 사용 번호를 확인하세요.")
    doc.core_properties.title = "채움랩 라벨 출력 패키지 고객용 프로그램 매뉴얼"
    doc.add_heading("GS1 데이터 표시와 재인쇄", level=2)
    doc.add_paragraph("GS1 DataMatrix의 표시문자를 켜면 AI별 데이터를 표시합니다. 새 배치에서는 문자 위치를 오른쪽 또는 아래로 선택합니다. GS1-128은 아래에 AI 괄호를 포함한 전체 값을 표시합니다. 값은 자동으로 잘라내거나 문자를 제거하지 않습니다. 괄호는 표시용이며 바코드에 인코딩하지 않습니다.")
    doc.add_paragraph("마지막 생성 작업 재인쇄는 저장된 값과 기준일을 재사용하며 번호를 소비하지 않습니다. 원본 엑셀이 없어도 원래 도안 버전으로 재인쇄할 수 있습니다. 번호 발급 초기화는 중복 위험을 확인하고 필요한 필드에만 적용하세요.")
    doc.add_heading("2D 바코드 크기 우선과 표시문자 배치", level=2)
    doc.add_paragraph("2D 개체를 더블클릭하여 2D 크기·문자 배치의 크기 방식을 확인합니다. 새 2D 개체는 바코드 크기 우선으로 시작합니다. 기존 도안은 기존 배치를 유지하며 바코드 크기 우선을 선택하여 명시적으로 전환합니다. 1D의 기존 배치는 유지합니다.")
    doc.add_paragraph("바코드 한 변(mm) 또는 바코드 가로(mm)·바코드 세로(mm)는 필수 여백을 포함하고 표시문자를 제외한 바코드 영역의 목표입니다. 코드 원래 비율과 정수 모듈 도트를 보존하므로 실제 크기는 목표와 다를 수 있습니다. 실제 크기 확인에서 현재 데이터와 출력 DPI에 따른 실제 mm·모듈 도트를 확인합니다. 모듈 크기 고정은 도트를 기준으로 하며 mm 목표와 동시에 독립 편집하지 않습니다.")
    doc.add_paragraph("문자 위치는 오른쪽 또는 아래, 문자 간격(mm)은 필수 여백 바깥 간격입니다. 문자 크기는 자동 맞춤 또는 지정 크기를 선택합니다. 표시문자·글꼴·위치를 바꾸어도 확정한 바코드 크기는 유지됩니다. 개체 전체 가로·세로는 코드와 문자를 함께 담는 영역입니다. 전체 가로·세로는 배치 결과에서 계산합니다. 문자 공간 오류 중에도 코드 크기 조절·이동·실행취소를 할 수 있습니다. 현재 라벨에 맞추기 또는 문자 아래 배치로 해결한 뒤 다시 검증합니다. 바코드를 자동 축소하거나 값을 잘라 해결하지 않습니다.")
    doc.add_paragraph("라벨 경계 또는 데이터별 모듈 크기 오류는 인쇄 전에 수정해야 합니다. 목표 mm는 모든 GS1 용도의 승인값이 아닙니다. 저장 도안과 반복 출력 미리보기·인쇄파일을 확인하고 실제 프린터·스캐너에서 첫 장의 위치·값·판독을 확인하세요.")
    doc.add_heading("AI·의료기기 UDI 통합 편집", level=2)
    doc.add_paragraph("바코드를 더블클릭해 바코드 속성을 엽니다. 데이터 원본에서 일반 데이터를 입력하고 바코드 종류·크기와 표시문자에서 바코드 종류와 배치를 조정합니다. GS1·UDI 구성은 데이터 원본의 GS1 데이터로 설정… 버튼 또는 기타 데이터 옵션…에서 필요할 때만 엽니다. ‘시험용 완성 예제로 시작’은 시험용 데이터이며 실제 운영 전 발급·관리하는 번호 또는 DB 열로 교체합니다. 필수 AI 선택은 부족한 조합을 함께 보여 줍니다. 새 AI와 수정 중인 행은 검증·적용 전에 자동 반영하며 임의의 실제 번호를 채우지 않습니다.")
    doc.add_paragraph("한 AI의 값 칸에는 값만 넣습니다. 전체 문자열의 AI 괄호 또는 스캐너 구분자는 전체 GS1 데이터 가져오기에서 처리하고 결과의 선행 0·항목·값을 확인합니다. AI21과 AI37의 필요한 조합은 서로 다르므로 모든 항목에 같은 식별번호 AI를 넣지 않습니다.")
    doc.add_paragraph("의료기기 UDI는 GS1과 HIBC 프로필을 구분합니다. 제품 식별번호(DI)와 로트·일련번호·제조일·사용기한(PI)을 한글 필드에 입력하고 고정값 또는 DB 열을 선택합니다. HIBC는 원본으로 생성과 완성 코드 가져오기를 구분하며 완성 원문의 +접두·전체 체크문자·공백을 임의로 삭제하거나 중복 추가하지 않습니다. 표시문자와 확정 코드 데이터를 대조합니다.")
    doc.add_paragraph("선택 기호 생성 검증에서 긴 UDI가 GS1-128 용량을 넘으면 데이터를 보존한 채 GS1 DataMatrix 전환을 검토합니다. 필수 필드를 삭제하거나 데이터를 잘라 해결하지 않습니다. 국내 프로필은 기관 등록·규제 승인 완료를 뜻하지 않습니다. Basic UDI-DI는 GTIN과 별도 정보입니다. UDI 선택은 편집 초안에만 반영됩니다. 창 하단의 적용 또는 확인으로 저장합니다. 취소, Esc 또는 창 닫기는 마지막 적용 상태를 보존합니다.")
    doc.add_paragraph("정사각형은 바코드 한 변(mm), 직사각형은 원래 비율로 가로·세로·대각선 손잡이를 조절합니다. 목표 mm와 정수 도트 기준 실제 mm·모듈 도트를 확인합니다. 문자 공간·라벨 경계 오류 중에도 코드 그림·크기 조절·이동·실행취소는 유지합니다. 현재 라벨에 맞추기와 문자 아래 배치를 검토한 뒤 다시 확인합니다. 모듈 크기 고정은 도트 설정으로 조절하며 손잡이는 잠깁니다. 미해결 오류는 인쇄를 차단합니다.")
    doc.add_page_break()
    doc.add_heading("UDI 통합 입력창 실제 화면", level=1)
    doc.add_paragraph("바코드 속성은 데이터 원본부터 열립니다. 직접 입력 또는 DB / 엑셀 연결을 선택하고 첫 화면의 GS1 데이터로 설정…을 사용합니다. 의료기기 UDI는 기타 데이터 옵션…에서 선택합니다. 상단 미리보기 버튼을 눌러 DB 행·실제 인코딩 값·표시문자·DPI·실제 mm를 확인하며 전체 보기·100%·확대 보기를 사용할 수 있습니다. 원본별 입력 초안은 전환해도 유지됩니다. 적용과 확인은 창 하단에서 한 번만 수행합니다.")
    udi_screen = ROOT / "outputs/ui-redesign-preview/frozen-udi-content.png"
    if not udi_screen.is_file():
        raise FileNotFoundError("최신 UDI 편집 화면 캡처가 없습니다.")
    doc.add_picture(str(udi_screen), width=Inches(6.5))
    doc.add_paragraph("UDI 구성 화면에서 DI/PI를 입력하고 바코드 속성의 바코드 종류·크기와 표시문자에서 실제 크기를 확인합니다. 편집창을 적용한 뒤 메인 화면의 전체 데이터 검증으로 모든 DB 행을 검사합니다. 시험용 예제 번호를 실제 고객 발급값으로 사용하지 마세요.")
    doc.add_page_break()
    doc.add_heading("소수 규격과 도안 인식 사용법", level=1)
    guide = (ROOT / "docs/MEDIA_AND_OCR_GUIDE.md").read_text(encoding="utf-8")
    for line in guide.split('## 제조사별 고급 오프셋')[0].splitlines()[2:]:
        if line.strip():
            doc.add_paragraph(line.replace('**', ''))
    doc.add_page_break()
    doc.add_heading("제조사별 오프셋 조정", level=1)
    doc.add_paragraph("센서·절취선 거리는 프린터 설정의 용지 종류 아래에서 조정합니다. 좌우·상하 보정은 인쇄 내용 위치를 이동하고, 공급 보정은 장비의 배출·정지 위치를 조정합니다. 실물 용지 이동은 방향·위치 미리보기에 표시되지 않습니다.")
    offset_table = doc.add_table(rows=1, cols=3)
    for cell, header in zip(offset_table.rows[0].cells, ("제조사", "센서 기준 보정", "배출·절취·커터 보정")):
        fill_cell(cell, header, bold=True)
    for row_data in (
        ("BIXOLON SLCS", "갭·블랙마크 / SL p4\n200도트 예시, p4 최대값 미기재\n프로그램은 용지·페이지 길이 상한", "TA ±100도트\n절취·커터 정지\n프린터 영구 저장"),
        ("TSC TSPL", "갭·블랙마크 GAP/BLINE 추가 이송\n0~용지 길이, auto는 명령 전체 생략", "OFFSET ±25.4mm\n별도 공급 보정"),
        ("Zebra ZPL", "해당 보정은 장비에서 설정", "~TA ±120단계\n600 DPI는 2도트/단계\nHC100은 0 이상"),
        ("SEWOO", "모델별 명령 확인 필요", "확인 전 고급 배출 보정 미제공"),
    ):
        for cell, text in zip(offset_table.add_row().cells, row_data):
            fill_cell(cell, text)
    style_table(offset_table, [1800, 3780, 3780])
    doc.add_paragraph("mm 입력은 장비 단위로 반올림합니다. TSC에서 현재 설정 사용은 GAP/BLINE 명령 전체를 생략하고, 직접 0은 선택한 용지 유형의 추가 이송 0을 적용합니다. BIXOLON p4를 명시하면 SL과 함께 인쇄 전 후진용 SF1을 보냅니다. 이 프로그램의 p4 상한은 제조사 최대 허용값이 아니라 용지·페이지 길이를 기준으로 한 안전 제한입니다.")
    doc.add_paragraph("설정 저장은 config.ini만 변경하며 실제 명령은 출력 작업에 포함됩니다. 빅솔론 TA는 프린터에 영구 저장되므로 기존 값을 기록하세요. 현재 프린터 설정 사용은 장비에 저장됐던 과거 값을 복원하는 기능이 아닙니다. 제조사·용지마다 별도 입력 초안을 사용합니다. 구버전 빅솔론의 비활성 칸에 저장된 기본 0은 장비 설정 유지로 불러옵니다.")
    doc.add_paragraph("Windows RAW 전송에는 드라이버 보정이 자동 적용되지 않습니다. 실제 1장으로 첫 장과 다음 장 시작 위치·배출·커터·스캔을 확인합니다. 제조사 명령 검증은 실물 모델 호환성 승인과 별도입니다.")
    doc.add_page_break()
    doc.add_heading("작은 라벨 자동 배치", level=1)
    doc.add_paragraph("자동 기본 배치에서도 현재 도안 미리보기를 사용할 수 있습니다. 40×20mm의 네 항목은 글자 폭·높이를 측정해 두 열 또는 여러 줄로 배치합니다. 8pt부터 최소 5pt까지 맞추며 값·항목·바코드 필수 여백은 보존합니다. 들어가지 않는 긴 값은 용지 크기나 저장 도안을 조정하세요. 미리보기는 원본 DB·인쇄 데이터·설정을 변경하거나 프린터로 전송하지 않습니다. 닫기 또는 Esc로 돌아갑니다.")
    doc.add_picture(str(ROOT / "outputs/ui-redesign-preview/manager-compact-preview.png"), width=Inches(6.5))
    doc.add_paragraph("자동 기본 배치 사용 → 현재 도안 미리보기에서 첫 출력 대상을 확인합니다. 그림은 시험용 예제이며 실제 상품명·상품 코드·LOT·수량·바코드를 대조하세요. 실제 출력용 흑백 도트를 미리 보여 주며 프린터로 보내지 않습니다. 첫 장에서 글자 크기·인쇄 위치와 스캔 결과를 확인합니다.")
    doc.core_properties.subject = "Excel 기반 라벨 출력 패키지 사용 설명서"
    doc.core_properties.author = "채움랩"
    doc.save(OUTPUT_DOCX)
    copy2(OUTPUT_DOCX, CUSTOMER_DIR / OUTPUT_DOCX.name)
    copy2(OUTPUT_DOCX, ROOT.parent / OUTPUT_DOCX.name)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()

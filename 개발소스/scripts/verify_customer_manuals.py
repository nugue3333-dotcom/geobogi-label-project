from __future__ import annotations

"""Verify PDF structure, readable text and synchronized customer copies."""

import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "고객용_매뉴얼"
COPIES = (
    ROOT / "고객용_실행폴더" / "docs/고객용_매뉴얼",
    ROOT.parent / "docs/고객용_매뉴얼",
)

EXPECTED = {
    "채움랩_라벨디자이너_고객용_매뉴얼.pdf": (14, 12, ("예제로 시작", "상품 엑셀 연결", "전체 값 확인", "좌우 가운데 정렬", "상하 가운데 정렬", ".cllabel", ".clproject", ".gblabel", ".gbproject", ".btw", "프린터 전송 완료", "Ctrl+C", "Ctrl+V", "Ctrl+D", "공통 규격", "용지 설정", "현재 도안 미리보기", "값 생성 규칙", "전체 데이터 검증", "필수 AI 선택", "시험용 완성 예제로 시작", "전체 GS1 데이터 가져오기", "의료기기 UDI", "완성 코드 가져오기", "현재 라벨에 맞추기", "GS1 데이터로 설정", "기타 데이터 옵션", "Esc", "ledger.json", "2D 크기·문자 배치", "바코드 크기 우선", "기존 배치", "모듈 크기 고정", "실제 크기 확인")),
    "채움랩_라벨출력관리_고객용_매뉴얼.pdf": (7, 6, ("중복 선택 창", "체크가 0개일 때", "전체 항목", "임시 큐", "취소하면 명령을 전송하지 않습니다", "인쇄 매수 선택", "지원 패키지 생성", "상한 없이", "양의 정수", "이미 프린터", "출력 도안 선택", "자동 기본 배치", "값 생성 규칙", "전체 데이터 검증", "ledger.json")),
    "채움랩_프린터설정_고객용_매뉴얼.pdf": (7, 6, ("180도 회전", "연결 확인", "BIXOLON/빅솔론", "공통 규격", "영역 밖")),
}

FORBIDDEN_POLICIES = (
    "선택 항목이 없으면 인쇄가 차단됩니다",
    "선택 없이 인쇄 버튼을 누르면 출력이 차단됩니다",
    "선택 항목이 없으면 인쇄가 시작되지 않습니다",
    "체크된 행이 없으면 인쇄가 차단됩니다",
    "선택된 행이 없으면 인쇄를 시작하지 않습니다",
    "인쇄 대상을 선택하지 않으면 전송되지 않습니다",
    "1~100장", "1~100매", "100장까지", "최대 100장", "최대 1000장",
)
FORBIDDEN_MENU_PATHS = (
    "설정 > 지원 패키지 생성", "설정 > 고객 데이터 백업", "설정 > 고객 데이터 복원",
    "설정 > 빠른 사용안내", "설정 > 상세 매뉴얼", "이동용 프로젝트 내보내기", "이동용 프로젝트 가져오기",
    "상세 매뉴얼 열기", "프린터설정 고객 매뉴얼",
)
MANUAL_VERSION = "2026.10.13.05"
REQUIRED_UI_SOURCES = tuple("barcode_label_automation/" + name for name in (
    "label_designer_app.py", "label_manager_app.py", "settings_app.py", "ui_tokens.py", "ui_window.py", "ui_style.py", "brand_assets.py",
    "gs1_editor.py", "barcode_options_dialog.py", "barcode_engine.py", "barcode_catalog.py", "gs1.py", "designer_clipboard.py", "label_media.py", "tool_icons.py", "print_quantity.py", "designer_print_dispatch.py",
    "job_profile.py", "template_output.py", "profile_dialogs.py", "value_rules.py", "value_rules_dialog.py", "value_jobs.py", "saved_output.py", "data_store.py", "barcode_layout.py", "hibc.py", "udi_editor.py", "workflow_probe.py", "barcode_editor_preview.py",
))
REQUIRED_BRAND_SOURCES = tuple("assets/brand/" + name for name in (
    "chaeumlab_logo_header_2x.png", "chaeumlab_designer_icon.png", "chaeumlab_manager_icon.png", "chaeumlab_settings_icon.png",
))


def capture_contract_errors(root: Path = ROOT) -> list[str]:
    """Reject stale captures, missing style provenance, or replaced screenshots."""
    metadata_path = root / "outputs/ui-redesign-preview/customer-manual-capture.json"
    if not metadata_path.exists():
        return ["actual source capture metadata is missing"]
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    errors = []
    if metadata.get("capture_kind") != "actual_source_gui":
        errors.append("manual screenshots must be actual source GUI captures")
    ui_sources = metadata.get("ui_sources", {})
    brand_sources = metadata.get("brand_sources", {})
    errors.extend("capture provenance missing: " + name for name in REQUIRED_UI_SOURCES if name not in ui_sources)
    errors.extend("capture provenance missing: " + name for name in REQUIRED_BRAND_SOURCES if name not in brand_sources)
    for name, expected in {**ui_sources, **brand_sources}.items():
        path = root / name
        if not path.exists() or digest(path) != expected:
            errors.append("capture source changed: " + name)
    required_screens = {"label-designer-final.png", "label-designer-data.png", "label-manager-final.png", "printer-settings.png", "frozen-udi-content.png", "frozen-udi-size.png"}
    screens = metadata.get("screenshots", [])
    errors.extend("capture missing: " + name for name in sorted(required_screens - {item["filename"] for item in screens}))
    for item in screens:
        path = metadata_path.parent / item["filename"]
        if not path.exists() or digest(path) != item["sha256"]:
            errors.append("capture image changed: " + item["filename"])
    return errors


def text_contract_errors(text: str, required: tuple[str, ...]) -> list[str]:
    compact = "".join(text.split())
    errors = [f"missing required phrase: {phrase}" for phrase in required if "".join(phrase.split()) not in compact]
    errors += [f"obsolete print policy: {phrase}" for phrase in FORBIDDEN_POLICIES if "".join(phrase.split()) in compact]
    errors += [f"obsolete menu path: {phrase}" for phrase in FORBIDDEN_MENU_PATHS if "".join(phrase.split()) in compact]
    return errors


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    from docx import Document
    from pypdf import PdfReader

    errors: list[str] = capture_contract_errors()
    for filename, (expected_pages, expected_bookmarks, phrases) in EXPECTED.items():
        source = SOURCE / filename
        if not source.exists():
            errors.append(f"missing: {source}")
            continue
        reader = PdfReader(str(source))
        if not expected_pages <= len(reader.pages) <= expected_pages + 3:
            errors.append(f"{filename}: pages={len(reader.pages)}, expected={expected_pages}..{expected_pages+3}")
        if len(reader.outline) < expected_bookmarks:
            errors.append(f"{filename}: bookmarks={len(reader.outline)}, expected>={expected_bookmarks}")
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
        for index, text in enumerate(texts, 1):
            if len(text) < 80:
                errors.append(f"{filename}: page {index} is blank or unreadable")
            if "\ufffd" in text:
                errors.append(f"{filename}: page {index} contains replacement glyphs")
        full_text = "\n".join(texts)
        if MANUAL_VERSION not in full_text:
            errors.append(f"{filename}: current release version is missing")
        errors += [f"{filename}: {error}" for error in text_contract_errors(full_text, phrases)]
        source_hash = digest(source)
        for copy_dir in COPIES:
            copied = copy_dir / filename
            if not copied.exists() or digest(copied) != source_hash:
                errors.append(f"{filename}: copy mismatch: {copied}")
        print(f"OK {filename}: pages={len(reader.pages)}, bookmarks={len(reader.outline)}, sha256={source_hash[:12]}")

    docx_name = "채움랩_라벨출력패키지_고객용_매뉴얼.docx"
    docx_source = ROOT / docx_name
    if not docx_source.exists():
        errors.append(f"missing: {docx_source}")
    else:
        document = Document(docx_source)
        docx_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        phrases = ("예제로 시작", "전체 값 확인", "좌우 가운데 정렬", "상하 가운데 정렬", ".cllabel", ".clproject", ".gblabel", ".gbproject", ".btw", "전체 항목", "임시 큐", "실제 출력 확인", "시험용 완성 예제로 시작", "전체 GS1 데이터 가져오기", "의료기기 UDI", "완성 코드 가져오기", "현재 라벨에 맞추기", "GS1 데이터로 설정", "기타 데이터 옵션", "Esc", MANUAL_VERSION)
        errors += [f"{docx_name}: {error}" for error in text_contract_errors(docx_text, phrases)]
        if len(document.inline_shapes) < 7:
            errors.append(f"{docx_name}: actual UI captures or brand icons are missing")
        source_hash = digest(docx_source)
        for copy in (ROOT / "고객용_실행폴더" / docx_name, ROOT.parent / docx_name):
            if not copy.exists() or digest(copy) != source_hash:
                errors.append(f"{docx_name}: copy mismatch: {copy}")
        print(f"OK {docx_name}: sha256={source_hash[:12]}")

    for name in ("README_먼저읽기.txt", "설치_및_사용_메뉴얼.txt", "사용안내.txt"):
        path = ROOT / name
        text = path.read_text(encoding="utf-8-sig")
        errors += [f"{name}: {error}" for error in text_contract_errors(text, (".cllabel", ".clproject", ".gblabel", ".gbproject", "임시 큐", "전체 항목"))]
        for copy in (ROOT / "고객용_실행폴더" / name, ROOT.parent / name):
            if not copy.exists() or digest(copy) != digest(path):
                errors.append(f"{name}: copy mismatch: {copy}")

    if errors:
        for error in errors:
            print(f"ERROR {error}")
        return 1
    print("CUSTOMER_MANUALS_VERIFIED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

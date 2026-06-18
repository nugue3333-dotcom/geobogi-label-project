from __future__ import annotations

import shutil
from pathlib import Path

import pythoncom
import win32com.client
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEPLOY_DIR = PROJECT_ROOT / "고객용_실행폴더"
FINAL_WORKBOOK = DEPLOY_DIR / "labels.xlsm"
TEST_WORKBOOK = DEPLOY_DIR / "labels_edit_test.xlsm"
QUEUE_FILE = DEPLOY_DIR / "print_queue.xlsx"

EDITED_ROW = ["EDIT-CODE", "EDITED ITEM", "EDIT-250601", "EDITLOT", 7, 1]


def main() -> int:
    pythoncom.CoInitialize()
    try:
        if TEST_WORKBOOK.exists():
            TEST_WORKBOOK.unlink()
        shutil.copy2(FINAL_WORKBOOK, TEST_WORKBOOK)

        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(TEST_WORKBOOK))
        try:
            sheet = workbook.Worksheets(1)
            for index, value in enumerate(EDITED_ROW, start=1):
                sheet.Cells(2, index).Value = value
            sheet.Range("A2").Select()
            excel.Run("DryRunLabels")
        finally:
            workbook.Close(False)
            excel.Quit()

        queue_row = _read_queue_row()
        print("edited queue row:", queue_row)
        if queue_row != EDITED_ROW:
            raise AssertionError(f"Edited Excel values were not exported: {queue_row!r}")

        final = _find_open_workbook(FINAL_WORKBOOK)
        owns_excel = False
        if final is None:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            final = excel.Workbooks.Open(str(FINAL_WORKBOOK))
            owns_excel = True
        else:
            excel = final.Application

        try:
            final.Worksheets(1).Range("A2").Select()
            excel.Run("DryRunLabels")
        finally:
            if owns_excel:
                final.Close(False)
                excel.Quit()

        print("restored queue row:", _read_queue_row())
        TEST_WORKBOOK.unlink(missing_ok=True)
        return 0
    finally:
        pythoncom.CoUninitialize()


def _read_queue_row() -> list[object]:
    workbook = load_workbook(QUEUE_FILE, data_only=True)
    sheet = workbook.active
    return [sheet.cell(2, column).value for column in range(1, 7)]


def _find_open_workbook(path: Path):
    target = str(path).lower()
    rot = pythoncom.GetRunningObjectTable()
    ctx = pythoncom.CreateBindCtx(0)
    for moniker in rot.EnumRunning():
        try:
            name = moniker.GetDisplayName(ctx, None)
        except Exception:
            continue
        if name.lower() == target:
            obj = rot.GetObject(moniker)
            return win32com.client.Dispatch(obj.QueryInterface(pythoncom.IID_IDispatch))
    return None


if __name__ == "__main__":
    raise SystemExit(main())

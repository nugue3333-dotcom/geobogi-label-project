from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


DB_SHEET_NAME = "BarcodeDB"
DB_HEADERS = ("barcode", "item_code", "item_name", "lot_no", "qty", "print_qty")
DB_SAMPLE_ROWS = (
    ("88023502", "A1001", "SENSOR BRACKET", "LOT250531", 100, 1),
    ("5J3P7YAYWXL5", "A1002", "SENSOR BRACKET", "LOT250532", 100, 2),
    ("A1001-250533", "A1003", "SENSOR BRACKET", "LOT250533", 100, 3),
    ("A1001-250534", "A1004", "SENSOR BRACKET", "LOT250534", 100, 4),
    ("KOR-TEST-001", "K1001", "\ud55c\uae00\ud488\ubaa9\ud14c\uc2a4\ud2b8", "LOT-HANGUL", 10, 1),
)


def rgb(red: int, green: int, blue: int) -> int:
    return red + (green * 256) + (blue * 65536)


def ensure_barcode_db_sheet(workbook: object, workbook_path: Path | None = None) -> object:
    source_path = workbook_path.with_name("barcode_db.xlsx") if workbook_path else None
    rows = load_db_rows(source_path)
    sheet = worksheet_or_none(workbook, DB_SHEET_NAME)

    if sheet is None:
        sheet = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
        sheet.Name = DB_SHEET_NAME
        write_db_rows(sheet, rows)
    elif not sheet_has_rows(sheet):
        write_db_rows(sheet, rows)

    style_db_sheet(sheet)
    arrange_workbook_sheets(workbook, sheet)
    return sheet


def worksheet_or_none(workbook: object, name: str) -> object | None:
    try:
        return workbook.Worksheets(name)
    except Exception:
        return None


def load_db_rows(path: Path | None) -> list[tuple[object, ...]]:
    if path is None or not path.exists():
        return [DB_HEADERS, *DB_SAMPLE_ROWS]

    source = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = source[DB_SHEET_NAME] if DB_SHEET_NAME in source.sheetnames else source.active
        rows: list[tuple[object, ...]] = []
        for values in sheet.iter_rows(values_only=True):
            normalized = tuple("" if value is None else value for value in values[: len(DB_HEADERS)])
            normalized = normalized + ("",) * (len(DB_HEADERS) - len(normalized))
            if any(str(value).strip() for value in normalized):
                rows.append(normalized)
        if not rows or tuple(str(value).strip() for value in rows[0]) != DB_HEADERS:
            rows.insert(0, DB_HEADERS)
        return rows
    finally:
        source.close()


def sheet_has_rows(sheet: object) -> bool:
    return bool(str(sheet.Cells(1, 1).Value or "").strip())


def write_db_rows(sheet: object, rows: list[tuple[object, ...]]) -> None:
    sheet.Range("A1:F500").Clear()
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row[: len(DB_HEADERS)], start=1):
            sheet.Cells(row_index, col_index).Value = value


def style_db_sheet(sheet: object) -> None:
    sheet.Cells.Font.Name = "Malgun Gothic"
    sheet.Cells.Font.Size = 10
    sheet.Tab.Color = rgb(17, 24, 39)

    widths = (22, 15, 28, 18, 10, 12)
    for column, width in enumerate(widths, start=1):
        sheet.Columns(column).ColumnWidth = width

    sheet.Columns("A:D").NumberFormat = "@"
    sheet.Columns("E:F").NumberFormat = "0"
    sheet.Rows("1:1").RowHeight = 28
    sheet.Rows("2:500").RowHeight = 22

    header = sheet.Range("A1:F1")
    header.Interior.Color = rgb(17, 24, 39)
    header.Font.Color = rgb(255, 255, 255)
    header.Font.Bold = True
    header.HorizontalAlignment = -4108
    header.VerticalAlignment = -4108

    body = sheet.Range("A2:F500")
    body.Interior.Color = rgb(255, 255, 255)
    body.Borders.LineStyle = 1
    body.Borders.Color = rgb(226, 232, 240)

    for row in range(2, 501):
        if row % 2 == 0:
            sheet.Range(f"A{row}:F{row}").Interior.Color = rgb(249, 250, 251)

    try:
        if sheet.AutoFilterMode:
            sheet.AutoFilterMode = False
        sheet.Range("A1:F500").AutoFilter()
    except Exception:
        pass


def arrange_workbook_sheets(workbook: object, db_sheet: object) -> None:
    labels = worksheet_or_none(workbook, "Labels")
    if labels is not None:
        labels.Move(Before=workbook.Worksheets(1))
        db_sheet = worksheet_or_none(workbook, DB_SHEET_NAME)
        if db_sheet is not None:
            db_sheet.Move(After=workbook.Worksheets("Labels"))

    workbook.Application.DisplayAlerts = False
    for index in range(workbook.Worksheets.Count, 0, -1):
        sheet = workbook.Worksheets(index)
        if sheet.Name not in ("Labels", DB_SHEET_NAME) and is_blank_sheet(sheet):
            sheet.Delete()


def is_blank_sheet(sheet: object) -> bool:
    try:
        used = sheet.UsedRange
        return used.Rows.Count == 1 and used.Columns.Count == 1 and str(sheet.Cells(1, 1).Value or "").strip() == ""
    except Exception:
        return False

from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from barcode_label_automation.data_store import DEFAULT_DB_HEADERS, sample_accessory_db_rows


CUSTOMER_DIR_NAME = "\uace0\uac1d\uc6a9_\uc2e4\ud589\ud3f4\ub354"
OUTPUTS = (
    PROJECT_ROOT / "barcode_db.xlsx",
    PROJECT_ROOT / "db" / "barcode_db.xlsx",
    PROJECT_ROOT / CUSTOMER_DIR_NAME / "barcode_db.xlsx",
    PROJECT_ROOT / CUSTOMER_DIR_NAME / "db" / "barcode_db.xlsx",
)
HEADERS = DEFAULT_DB_HEADERS


def main() -> int:
    for output in OUTPUTS:
        output.parent.mkdir(parents=True, exist_ok=True)
        create_db(output)
        print(f"created: {output}")
    return 0


def create_db(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"

    for col, header in enumerate(HEADERS, start=1):
        sheet.cell(row=1, column=col, value=header)

    rows = sample_accessory_db_rows()
    for row_index, row in enumerate(rows, start=2):
        for col, header in enumerate(HEADERS, start=1):
            sheet.cell(row=row_index, column=col, value=row.get(header, ""))

    _style(sheet)
    workbook.save(path)


def _style(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="111827")
    body_fill = PatternFill("solid", fgColor="FFFFFF")
    alt_fill = PatternFill("solid", fgColor="F9FAFB")
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    widths = {
        "barcode": 18,
        "item_code": 15,
        "item_name": 28,
        "판매가": 14,
    }
    for col, header in enumerate(HEADERS, start=1):
        width = widths[header]
        sheet.column_dimensions[get_column_letter(col)].width = width

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    sample_rows = sample_accessory_db_rows()
    for row in sheet.iter_rows(
        min_row=2,
        max_row=len(sample_rows) + 1,
        min_col=1,
        max_col=len(HEADERS),
    ):
        fill = alt_fill if row[0].row % 2 == 0 else body_fill
        for cell in row:
            cell.fill = fill
            cell.font = Font(name="Malgun Gothic", size=10, color="111827")
            cell.alignment = Alignment(vertical="center")
            cell.border = border

    for col in ("A", "B", "C"):
        for cell in sheet[col]:
            cell.number_format = "@"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(sample_rows) + 1}"


if __name__ == "__main__":
    raise SystemExit(main())

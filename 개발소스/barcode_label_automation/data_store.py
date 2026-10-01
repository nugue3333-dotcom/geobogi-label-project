from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook


LABEL_HEADERS = ("item_code", "item_name", "barcode", "lot_no", "qty", "print_qty")
DB_HEADERS = ("barcode", "item_code", "item_name", "lot_no", "qty", "print_qty")
# The built-in customer demo DB is intentionally smaller than the internal
# print queue. Lot, quantity, and print quantity must not appear in its UI.
DEFAULT_DB_HEADERS = ("barcode", "item_code", "item_name", "판매가")
SAMPLE_DB_EXTRA_HEADERS = ("판매가",)
FIELD_ALIASES = {
    "barcode": ("barcode", "bar_code", "바코드", "상품바코드", "제품바코드"),
    "item_code": ("item_code", "itemcode", "품목코드", "상품코드", "제품코드", "코드", "품번"),
    "item_name": ("item_name", "itemname", "품목명", "품명", "상품명", "제품명", "이름", "명칭"),
    "lot_no": ("lot_no", "lot", "lot번호", "lot 번호", "로트", "로트번호", "로트 번호", "lotno"),
    "qty": ("qty", "quantity", "수량", "입수", "개수"),
    "print_qty": ("print_qty", "printqty", "출력매수", "출력 매수", "인쇄매수", "인쇄 매수", "매수"),
}


@dataclass(frozen=True)
class BarcodeLookupResult:
    barcode: str
    item_code: str
    item_name: str
    lot_no: str
    qty: str
    print_qty: str


def load_label_rows(path: str | Path) -> list[dict[str, str]]:
    return load_table(path, LABEL_HEADERS, sheet_name="Labels")


def save_label_rows(path: str | Path, rows: list[dict[str, str]]) -> None:
    save_table(path, _headers_with_row_extras(LABEL_HEADERS, rows), rows, sheet_name="Labels")


def load_db_rows(path: str | Path) -> list[dict[str, str]]:
    rows, _headers = load_db_source(path)
    # Print and lookup workflows use these canonical fields, while the data
    # source UI only receives the columns that actually exist in its workbook.
    normalized_rows: list[dict[str, str]] = []
    for row in rows:
        normalized = {header: row.get(header, "") for header in DB_HEADERS}
        normalized.update(row)
        normalized_rows.append(normalized)
    return normalized_rows


def load_db_source(path: str | Path) -> tuple[list[dict[str, str]], tuple[str, ...]]:
    """Load DB rows plus the real workbook headers for UI data-source pickers."""
    workbook_path = Path(path)
    if not workbook_path.exists():
        save_default_db_rows(workbook_path, _sample_db_rows())

    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet = workbook["BarcodeDB"] if "BarcodeDB" in workbook.sheetnames else workbook.active
        header_map = _flexible_header_map(sheet, DB_HEADERS)
        source_headers = [_cell_text(cell.value) for cell in sheet[1]]
        visible_headers = tuple(dict.fromkeys(header for header in source_headers if header))
        rows: list[dict[str, str]] = []
        for row_number in range(2, sheet.max_row + 1):
            row: dict[str, str] = {}
            for header, column_index in header_map.items():
                row[header] = _db_cell_text(sheet.cell(row_number, column_index), row_number, header)
            for column_index, header in enumerate(source_headers, start=1):
                if not header or header in row:
                    continue
                row[header] = _db_cell_text(sheet.cell(row_number, column_index), row_number, header)
            if any(value.strip() for value in row.values()):
                rows.append(row)
        return rows, visible_headers
    finally:
        workbook.close()


def save_db_rows(path: str | Path, rows: list[dict[str, str]]) -> None:
    save_table(path, _headers_with_row_extras(DB_HEADERS, rows), rows, sheet_name="BarcodeDB")


def save_default_db_rows(path: str | Path, rows: list[dict[str, str]]) -> None:
    """Save the customer demo DB without internal lot/quantity columns."""
    save_table(path, DEFAULT_DB_HEADERS, rows, sheet_name="BarcodeDB")


def sample_accessory_db_rows() -> list[dict[str, str]]:
    """Return deterministic accessory test data used by source and customer DBs."""
    product_names = (
        "실버 볼 체인 목걸이",
        "미니 하트 귀걸이",
        "데일리 진주 귀걸이",
        "슬림 레이어드 반지",
        "컬러 비즈 팔찌",
        "오벌 헤어 집게핀",
        "미니 리본 헤어핀",
        "하트 키링",
        "아크릴 키링",
        "카드 수납 지갑",
    )
    rows: list[dict[str, str]] = []
    for index in range(1, 101):
        price = 3900 + ((index * 700) % 16100)
        product_number = f"ACC-{index:03d}"
        rows.append(
            {
                "barcode": f"8801000{index:06d}",
                "item_code": product_number,
                "item_name": product_names[(index - 1) % len(product_names)],
                "판매가": f"{price:,}원",
            }
        )
    return rows


def lookup_barcode(db_rows: list[dict[str, str]], barcode: str) -> BarcodeLookupResult | None:
    target = str(barcode).strip()
    if not target:
        return None
    for row in db_rows:
        if _row_field_value(row, "barcode") == target:
            return BarcodeLookupResult(
                barcode=target,
                item_code=_row_field_value(row, "item_code"),
                item_name=_row_field_value(row, "item_name"),
                lot_no=_row_field_value(row, "lot_no"),
                qty=_row_field_value(row, "qty"),
                print_qty=_row_field_value(row, "print_qty") or "1",
            )
    return None


def load_table(
    path: str | Path,
    headers: tuple[str, ...],
    sheet_name: str,
    sample_rows: list[dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        save_table(workbook_path, headers, sample_rows or [], sheet_name=sheet_name)

    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        header_map = _header_map(sheet, headers)
        rows: list[dict[str, str]] = []
        for row_number in range(2, sheet.max_row + 1):
            row = {header: _cell_text(sheet.cell(row_number, header_map[header]).value) for header in headers}
            if any(value.strip() for value in row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def save_table(path: str | Path, headers: tuple[str, ...], rows: list[dict[str, str]], sheet_name: str) -> None:
    workbook_path = Path(path)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(headers))

    for row in rows:
        if not any(str(row.get(header, "")).strip() for header in headers):
            continue
        sheet.append([str(row.get(header, "")).strip() for header in headers])

    for column_index, header in enumerate(headers, start=1):
        width = 28 if header in {"item_name", "barcode"} else 16
        sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = width
        font = copy(sheet.cell(1, column_index).font)
        font.bold = True
        sheet.cell(1, column_index).font = font
    for column_index in range(1, min(4, len(headers)) + 1):
        for row_index in range(1, sheet.max_row + 1):
            sheet.cell(row_index, column_index).number_format = "@"
    workbook.save(workbook_path)


def _headers_with_row_extras(base_headers: tuple[str, ...], rows: list[dict[str, str]]) -> tuple[str, ...]:
    headers = list(base_headers)
    seen = set(headers)
    for row in rows:
        for header in row:
            if header not in seen:
                headers.append(header)
                seen.add(header)
    return tuple(headers)


def _header_map(sheet: object, expected_headers: tuple[str, ...]) -> dict[str, int]:
    found: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):
        value = _cell_text(cell.value)
        if value:
            found[value] = column_index
    missing = [header for header in expected_headers if header not in found]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    return {header: found[header] for header in expected_headers}


def _flexible_header_map(sheet: object, expected_headers: tuple[str, ...]) -> dict[str, int]:
    by_normalized: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):
        key = _normalize_header_key(_cell_text(cell.value))
        if key and key not in by_normalized:
            by_normalized[key] = column_index

    result: dict[str, int] = {}
    for header in expected_headers:
        for alias in (header, *FIELD_ALIASES.get(header, ())):
            column_index = by_normalized.get(_normalize_header_key(alias))
            if column_index is not None:
                result[header] = column_index
                break
    if "barcode" not in result:
        raise ValueError("Missing required columns: barcode 또는 바코드")
    return result


def _field_for_header(header: str) -> str | None:
    normalized = _normalize_header_key(header)
    for field, aliases in FIELD_ALIASES.items():
        if normalized in {_normalize_header_key(alias) for alias in aliases}:
            return field
    return None


def canonical_db_field(header: str) -> str | None:
    """Return the internal field name for a user-facing DB header, when known."""
    return _field_for_header(header)


def _row_field_value(row: dict[str, str], field: str) -> str:
    for header, value in row.items():
        if _field_for_header(header) == field:
            return str(value).strip()
    return str(row.get(field, "")).strip()


def _normalize_header_key(value: str) -> str:
    return "".join(ch for ch in str(value).strip().lower() if ch.isalnum())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _db_cell_text(cell: object, row_number: int, header: str) -> str:
    value = cell.value
    text = _cell_text(value)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return text
    if isinstance(value, float) and not value.is_integer():
        return text

    number_format = str(cell.number_format or "General")
    if _field_for_header(header) == "barcode" and len(text.lstrip("-")) > 15:
        raise ValueError(f"{row_number}행 '{header}' 바코드는 Excel 숫자 정밀도를 넘습니다. 셀을 텍스트로 저장하세요.")
    if value >= 0 and re.fullmatch(r"0{2,}", number_format):
        return text.zfill(len(number_format))
    if _field_for_header(header) == "barcode" and number_format not in {"General", "0"}:
        raise ValueError(f"{row_number}행 '{header}' 바코드의 숫자 서식 '{number_format}'을 안전하게 읽을 수 없습니다. 셀을 텍스트로 저장하세요.")
    return text


def _legacy_sample_db_rows() -> list[dict[str, str]]:
    values = [
        ("88023502", "A1001", "SENSOR BRACKET", "LOT250531", "100", "1"),
        ("5J3P7YAYWXL5", "A1002", "SENSOR BRACKET", "LOT250532", "100", "2"),
        ("A1001-250533", "A1003", "SENSOR BRACKET", "LOT250533", "100", "3"),
        ("A1001-250534", "A1004", "SENSOR BRACKET", "LOT250534", "100", "4"),
        ("KOR-TEST-001", "K1001", "\ud55c\uae00\ud488\ubaa9\ud14c\uc2a4\ud2b8", "LOT-HANGUL", "10", "1"),
    ]
    return [dict(zip(DB_HEADERS, row, strict=True)) for row in values]
def _sample_db_rows() -> list[dict[str, str]]:
    return sample_accessory_db_rows()

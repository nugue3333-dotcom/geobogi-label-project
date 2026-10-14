"""Workbook provenance saved with a label, without copying customer rows."""
from __future__ import annotations

import hashlib
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .data_store import load_db_source


def workbook_sheets(path: Path) -> tuple[str, ...]:
    if not path.is_file():
        raise FileNotFoundError(f"상품 엑셀을 찾을 수 없습니다: {path.name}")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(book.sheetnames)
    finally:
        book.close()


def default_sheet(path: Path) -> str:
    sheets = workbook_sheets(path)
    return "BarcodeDB" if "BarcodeDB" in sheets else sheets[0]


def file_fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def data_profile(path: Path, *, sheet: str | None = None, headers=(), mode="원본 참조", original_path: Path | None = None) -> dict:
    selected = sheet or default_sheet(path)
    if not headers:
        _, headers = load_db_source(path, sheet_name=selected)
    return {
        "path": str(path.resolve()), "sheet": selected, "headers": list(headers),
        "fingerprint": file_fingerprint(path), "mode": mode,
        "original_path": str(original_path.resolve()) if original_path else "",
        "original_fingerprint": file_fingerprint(original_path) if original_path and original_path.is_file() else "",
        "loaded_at": datetime.now(timezone.utc).isoformat(),
    }


def normalize_data_profile(value: object) -> dict:
    if value is None:
        return {}
    if not isinstance(value, dict):
        raise ValueError("도안의 엑셀 연결 정보가 올바르지 않습니다.")
    headers = value.get("headers", [])
    if not isinstance(headers, list) or not all(isinstance(item, str) for item in headers):
        raise ValueError("도안의 DB 열 목록이 올바르지 않습니다.")
    return {key: str(value.get(key, "")) for key in
            ("path", "sheet", "fingerprint", "mode", "original_path", "original_fingerprint", "loaded_at")} | {"headers": list(headers)}


def profile_problem(profile: dict) -> str | None:
    if profile and not profile.get("path"):
        return "이 도안의 상품 엑셀을 현재 PC에 다시 연결하세요."
    if not profile or not profile.get("path"):
        return None
    path = Path(profile["path"])
    if not path.is_file():
        return "연결한 엑셀을 찾을 수 없습니다. 상품 엑셀 재연결을 눌러 파일과 시트를 확인하세요."
    try:
        if profile.get("fingerprint") and file_fingerprint(path) != profile["fingerprint"]:
            return "연결 후 엑셀 파일이 변경되었습니다. 엑셀 새로고침 후 열 연결과 값을 확인하세요."
        if profile.get("sheet") not in workbook_sheets(path):
            return "연결한 시트가 없습니다. 상품 엑셀을 다시 연결하세요."
        if profile.get("mode") == "가져온 복사본" and profile.get("original_path"):
            original = Path(profile["original_path"])
            if not original.is_file():
                return "가져온 복사본의 원본 엑셀을 찾을 수 없습니다. 원본을 재연결해 최신 데이터를 확인하세요."
            if profile.get("original_fingerprint") and file_fingerprint(original) != profile["original_fingerprint"]:
                return "원본 엑셀이 변경되어 가져온 복사본과 다릅니다. 엑셀 새로고침으로 원본을 다시 가져오세요."
    except (OSError, ValueError) as exc:
        return f"연결한 엑셀을 확인할 수 없습니다: {exc}"
    return None


def describe_profile(profile: dict) -> str:
    if not profile:
        return "상품 엑셀 미연결"
    stamp = str(profile.get("loaded_at", ""))
    try:
        stamp = datetime.fromisoformat(stamp).astimezone(timezone(timedelta(hours=9), "Asia/Seoul")).strftime("%m/%d %H:%M")
    except (ValueError, TypeError):
        stamp = "확인 필요"
    return f"{profile.get('mode') or '원본 참조'} · {Path(profile.get('path', '')).name}\n시트: {profile.get('sheet', '')} · 읽은 시각: {stamp}\n" + (profile_problem(profile) or "연결 파일 변경 없음")

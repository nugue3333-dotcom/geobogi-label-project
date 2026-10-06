from __future__ import annotations

import argparse
import copy
import locale
import os
import shutil
import subprocess
import sys
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook

from .brand_assets import apply_window_icon, load_header_logo
from .config import load_config
from .data_store import (
    DB_HEADERS,
    LABEL_HEADERS,
    _db_cell_text,
    load_db_source,
    load_db_rows,
    load_label_rows,
    save_label_rows,
)
from .runtime_paths import executable_dir, runtime_base_dir
from .print_progress import (
    NEW_JOB_REQUIRED_CODE,
    PROGRESS_FILE_NAME,
    TRANSPORT_UNCERTAIN_CODE,
    UNKNOWN_ITEMS_CODE,
    PrintProgress,
    PrintProgressError,
)
from .settings_app import load_settings as load_printer_settings
from .settings_app import validate_settings as validate_printer_settings
from .ui_tokens import COLORS, SPACING, TYPOGRAPHY
from .ui_style import configure_suite_style
from .ui_window import set_initial_window_size


SELECT_HEADER = "__selected__"
LABEL_DISPLAY_HEADERS = (SELECT_HEADER, *LABEL_HEADERS)
CHECKED = "☑"
UNCHECKED = "☐"
DEFAULT_PRINT_QTY = "1"
SUPPORT_REPORT_FILE = "customer_preflight_report.txt"
SUPPORT_PACKAGE_FILE = "customer_support_package.zip"

COLUMN_LABELS = {
    SELECT_HEADER: "선택",
    "barcode": "상품 바코드",
    "item_code": "상품 코드",
    "item_name": "상품명",
    "lot_no": "LOT",
    "qty": "수량",
    "print_qty": "출력 매수",
    "판매가": "판매가",
}

FIELD_ALIASES = {
    "barcode": ("barcode", "bar_code", "바코드", "상품바코드", "제품바코드"),
    "item_code": ("item_code", "itemcode", "품목코드", "상품코드", "제품코드", "코드", "품번"),
    "item_name": ("item_name", "itemname", "품목명", "품명", "상품명", "제품명", "이름", "명칭"),
    "lot_no": ("lot_no", "lot", "lot번호", "lot 번호", "로트", "로트번호", "로트 번호", "lotno"),
    "qty": ("qty", "quantity", "수량", "입수", "개수"),
    "print_qty": ("print_qty", "printqty", "출력매수", "출력 매수", "인쇄매수", "인쇄 매수", "매수"),
}


def app_base_dir() -> Path:
    return runtime_base_dir(executable_dir())


def _initial_duplicate_selection_positions(match_count: int) -> set[int]:
    return set()


def _duplicate_selection_result(
    matches: list[tuple[int, dict[str, str]]],
    selected_positions: set[int],
) -> list[tuple[int, dict[str, str]]]:
    return [matches[position] for position in sorted(selected_positions) if 0 <= position < len(matches)]


class LabelManagerApp(tk.Tk):
    def __init__(self, base_dir: Path) -> None:
        super().__init__()
        self._display_scale = self._resolve_display_scale()
        self.base_dir = base_dir
        self.install_dir = executable_dir() if getattr(sys, "frozen", False) else base_dir
        self.queue_path = base_dir / "print_queue.xlsx"
        self.db_path = base_dir / "barcode_db.xlsx"
        self.db_upload_dir = base_dir / "db"
        self.config_path = base_dir / "config.ini"
        self.print_exe = _first_existing(base_dir / "라벨출력엔진.exe", self.install_dir / "라벨출력엔진.exe")
        self.settings_exe = _first_existing(
            self.install_dir / "\ud504\ub9b0\ud130\uc124\uc815.exe",
            base_dir / "\ud504\ub9b0\ud130\uc124\uc815.exe",
        )
        self.preflight_exe = _first_existing(
            self.install_dir / "고객환경점검.exe",
            base_dir / "고객환경점검.exe",
        )
        self.designer_exe = _first_existing(
            self.install_dir / "라벨디자이너.exe",
            base_dir / "라벨디자이너.exe",
        )
        self.quick_guide_path = _first_existing(
            base_dir / "사용안내.txt",
            self.install_dir / "사용안내.txt",
        )
        self.manual_path = _first_existing(
            *(root / folder / name for name in ("채움랩_라벨출력관리_고객용_매뉴얼.pdf", "채움LAB_라벨출력관리_고객용_매뉴얼.pdf")
              for root in (base_dir, self.install_dir)
              for folder in (Path("고객용_매뉴얼"), Path("docs/고객용_매뉴얼"), Path("."))),
            base_dir / "채움랩_라벨출력패키지_고객용_매뉴얼.docx",
            self.install_dir / "채움랩_라벨출력패키지_고객용_매뉴얼.docx",
            base_dir / "라벨출력패키지_고객용_매뉴얼.docx",
            self.install_dir / "라벨출력패키지_고객용_매뉴얼.docx",
            base_dir / "설치_및_사용_메뉴얼.txt",
            self.install_dir / "설치_및_사용_메뉴얼.txt",
        )
        self.label_rows: list[dict[str, str]] = []
        self.db_rows: list[dict[str, str]] = []
        self.selected_label_indexes: set[int] = set()
        self.active_editor: tk.Entry | None = None
        self.db_connected = True
        self.db_headers = tuple(DB_HEADERS)
        self.label_headers = tuple(LABEL_HEADERS)
        self.label_display_headers = (SELECT_HEADER, *self.label_headers)
        self.scan_var = tk.StringVar()
        self.status_var = tk.StringVar()
        self.document_name_var = tk.StringVar(value=self.db_path.name)
        self.document_state_var = tk.StringVar(value="불러오는 중")
        self._data_dirty = False
        self._delete_history: list[dict[str, object]] = []
        self._job_panel_visible = True
        self.print_scope_var = tk.StringVar(value="인쇄할 데이터 없음")
        self._print_action_widgets: list[ttk.Button] = []
        self.output_menu: tk.Menu | None = None
        self.brand_logo = load_header_logo(
            self,
            base_dir=self.base_dir,
            install_dir=self.install_dir,
            max_width=self._scaled(168),
            max_height=self._scaled(40),
        )

        self.title("채움랩 라벨 출력 관리")
        set_initial_window_size(
            self,
            preferred_width=self._scaled(1320),
            preferred_height=self._scaled(820),
            minimum_width=self._scaled(1020),
            minimum_height=self._scaled(660),
        )
        self.configure(bg=COLORS.background)

        self._configure_style()
        self._apply_window_icon()
        self._build_ui()
        self.bind_all("<Control-f>", lambda _event: self._focus_scan_input())
        self.bind_all("<Control-s>", lambda _event: self._run_keyboard_command(self.save_files))
        self.bind_all("<Control-p>", lambda _event: self._run_keyboard_command(self.print_with_quantity))
        self.bind_all("<Control-z>", self._on_undo_delete_shortcut)
        self.protocol("WM_DELETE_WINDOW", self._confirm_close)
        self.load_files()

    def _run_keyboard_command(self, command) -> str:
        command()
        return "break"

    def _focus_scan_input(self) -> str:
        self.scan_box.focus_set()
        return "break"

    def _resolve_display_scale(self) -> float:
        try:
            scale = float(self.winfo_fpixels("1i")) / 96.0
        except (tk.TclError, TypeError, ValueError):
            return 1.0
        return max(1.0, min(2.0, scale))

    def _scaled(self, value: int) -> int:
        return max(1, round(value * self._display_scale))

    def _apply_window_icon(self) -> None:
        apply_window_icon(self, base_dir=self.base_dir, install_dir=self.install_dir, app_role="manager")

    def _configure_style(self) -> None:
        style = configure_suite_style(self)
        for name in ("SidePanel",):
            style.configure(f"{name}.TFrame", background=COLORS.surface)
        style.configure("Brand.TLabel", font=TYPOGRAPHY.button_text, foreground=COLORS.surface,
                        background=COLORS.primary, padding=(9, 4))
        style.configure("Title.TLabel", font=TYPOGRAPHY.section_title, foreground=COLORS.text_primary,
                        background=COLORS.surface)
        style.configure("HeaderField.TLabel", font=TYPOGRAPHY.button_text, foreground=COLORS.text_primary,
                        background=COLORS.surface)
        style.configure("Subtitle.TLabel", font=TYPOGRAPHY.caption, foreground=COLORS.text_secondary,
                        background=COLORS.surface)
        style.configure("Hint.TLabel", font=TYPOGRAPHY.caption, foreground=COLORS.text_secondary,
                        background=COLORS.surface)
        for name in ("PrintScope", "ToolbarScope", "OnboardingStep"):
            style.configure(f"{name}.TLabel", font=TYPOGRAPHY.button_text, foreground=COLORS.primary,
                            background=COLORS.accent_soft, padding=(8, 5))
        style.configure("ToolbarStatus.TLabel", font=TYPOGRAPHY.caption, foreground=COLORS.text_secondary,
                        background=COLORS.surface)
        style.configure("Compact.TButton", font=TYPOGRAPHY.button_text, foreground=COLORS.text_primary,
                        background=COLORS.surface, bordercolor=COLORS.primary, padding=(8, 5))
        style.map("Compact.TButton", background=[("disabled", COLORS.surface_muted), ("active", COLORS.surface_muted)],
                  foreground=[("disabled", COLORS.text_secondary)], bordercolor=[("focus", COLORS.primary)])
        style.configure("Menu.TMenubutton", font=TYPOGRAPHY.button_text, foreground=COLORS.text_primary,
                        background=COLORS.surface, padding=(10, 6))

    def _build_ui(self) -> None:
        page_padding = self._scaled(12)
        top = ttk.Frame(self, style="Header.TFrame", padding=(page_padding, self._scaled(8)))
        top.pack(fill="x")
        top.columnconfigure(1, weight=1)
        if self.brand_logo is not None:
            ttk.Label(top, image=self.brand_logo, style="HeaderLogo.TLabel").grid(
                row=0, column=0, rowspan=2, padx=(0, self._scaled(16)), sticky="w")
        else:
            ttk.Label(top, text="채움랩", style="Brand.TLabel").grid(row=0, column=0, rowspan=2, sticky="w")
        ttk.Label(top, textvariable=self.document_name_var, style="HeaderField.TLabel").grid(row=0, column=1, sticky="w")
        ttk.Label(top, textvariable=self.document_state_var, style="Subtitle.TLabel").grid(row=1, column=1, sticky="w")
        ttk.Label(top, text="라벨출력관리", style="HeaderTitle.TLabel").grid(row=0, column=2, rowspan=2, sticky="e")
        self._build_menu_surface()

        body = ttk.Frame(self, style="App.TFrame", padding=(page_padding, 6, page_padding, 8))
        body.pack(fill="both", expand=True)
        body.rowconfigure(2, weight=1)
        body.columnconfigure(0, weight=1)
        toolbar = ttk.Frame(body, style="Toolbar.TFrame", padding=(8, 6))
        toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        actions = (
            ("상품 엑셀 연결", self.connect_db_file), ("저장", self.save_files),
            ("새 행", self.add_row), ("선택 삭제", self.delete_selected),
            ("삭제 실행취소", self.undo_delete), ("전체 선택", self.select_all_labels),
            ("선택 해제", self.clear_label_selection), ("라벨디자이너", self.open_designer),
            ("프린터 설정", self.open_settings), ("실행 전 점검", self.run_preflight_check),
            ("인쇄", self.print_with_quantity),
        )
        buttons = []
        for label, command in actions:
            style = "Primary.TButton" if label == "인쇄" else "Secondary.TButton"
            button = ttk.Button(toolbar, text=label, command=command, style=style)
            buttons.append(button)
            if label == "인쇄":
                self.print_button = button
                self._print_action_widgets.append(button)
            elif label == "삭제 실행취소":
                self.undo_delete_button = button
                button.state(["disabled"])
        self._toolbar_buttons = buttons
        self._toolbar_columns = 0

        def layout_actions(event=None):
            available = max(240, int(event.width) if event is not None else toolbar.winfo_width())
            for column in range(len(buttons)):
                toolbar.columnconfigure(column, weight=0)
            row = column = used = 0
            for button in buttons:
                requested = button.winfo_reqwidth() + 6
                if column and used + requested > available - 16:
                    row += 1
                    column = used = 0
                button.grid(row=row, column=column, sticky="ew", padx=(0, 6), pady=(0, 4))
                used += requested
                column += 1
            self._toolbar_columns = column
        toolbar.bind("<Configure>", layout_actions)
        self.after_idle(layout_actions)

        search = ttk.Frame(body, style="Surface.TFrame", padding=(10, 7))
        search.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        search.columnconfigure(1, weight=1)
        ttk.Label(search, text="상품 DB 조회", style="HeaderField.TLabel").grid(row=0, column=0, padx=(0, 8))
        self.scan_box = ttk.Entry(search, textvariable=self.scan_var)
        self.scan_box.grid(row=0, column=1, sticky="ew")
        self.scan_box.bind("<Return>", self.scan_barcode)
        ttk.Button(search, text="조회", command=self.scan_barcode, style="Secondary.TButton").grid(row=0, column=2, padx=(8, 0))
        ttk.Label(search, textvariable=self.print_scope_var, style="ToolbarScope.TLabel").grid(row=0, column=3, padx=(12, 0))

        work_area = ttk.Frame(body, style="App.TFrame")
        self._work_area = work_area
        work_area.grid(row=2, column=0, sticky="nsew")
        work_area.rowconfigure(0, weight=1)
        work_area.columnconfigure(0, weight=1)
        work_area.columnconfigure(1, minsize=self._scaled(288))
        self.tabs = ttk.Notebook(work_area)
        self.tabs.grid(row=0, column=0, sticky="nsew")
        self.labels_tree = self._create_table(self.tabs, LABEL_DISPLAY_HEADERS, "인쇄 데이터")
        self.db_tree = self._create_table(self.tabs, DB_HEADERS, "원본 DB")

        side_card = tk.Frame(work_area, bg=COLORS.surface, highlightbackground=COLORS.border,
                             highlightthickness=1, bd=0)
        self._job_panel_card = side_card
        side_card.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
        side_card.columnconfigure(0, weight=1)
        side_card.rowconfigure(0, weight=1)
        side_canvas = tk.Canvas(side_card, width=self._scaled(278), bg=COLORS.background,
                                highlightthickness=0, bd=0)
        side_canvas.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(side_card, orient="vertical", command=side_canvas.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        side_canvas.configure(yscrollcommand=scroll.set)
        side_inner = ttk.Frame(side_canvas, style="App.TFrame", padding=8)
        side_window = side_canvas.create_window((0, 0), window=side_inner, anchor="nw")
        side_inner.columnconfigure(0, weight=1)
        side_inner.bind("<Configure>", lambda _event: side_canvas.configure(scrollregion=side_canvas.bbox("all")))
        side_canvas.bind("<Configure>", lambda event: side_canvas.itemconfigure(side_window, width=event.width))
        job_inner = self._panel_card(side_inner, 0, "출력 작업")
        self._build_job_panel(job_inner)
        workflow_inner = self._panel_card(side_inner, 1, "작업 순서")
        self._build_start_checklist(workflow_inner)

        def scroll_panel(event):
            side_canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")
            return "break"
        def bind_scroll(widget):
            widget.bind("<MouseWheel>", scroll_panel, add="+")
            for child in widget.winfo_children():
                bind_scroll(child)
        bind_scroll(side_inner)
        side_canvas.bind("<MouseWheel>", scroll_panel)
        footer = ttk.Frame(self, style="StatusBar.TFrame", padding=(page_padding, 6))
        footer.pack(side="bottom", fill="x", before=body)
        ttk.Label(footer, textvariable=self.status_var, style="FooterStatus.TLabel", anchor="w").pack(side="left", fill="x", expand=True)
        ttk.Button(footer, text="출력 패널 보기/숨기기", command=self.toggle_job_panel, style="Compact.TButton").pack(side="right")
        self._update_print_scope()

    def _build_menu_surface(self) -> None:
        bar = ttk.Frame(self, style="Toolbar.TFrame", padding=(12, 0))
        bar.pack(fill="x")
        definitions = (
            ("파일", (("상품 엑셀 연결", self.connect_db_file), ("저장", self.save_files),
                     ("고객 데이터 백업", self.open_backup), ("고객 데이터 복원", self.open_restore))),
            ("보기", (("인쇄 데이터", lambda: self.tabs.select(0)), ("원본 DB", lambda: self.tabs.select(1)),
                     ("출력 작업 패널 보기/숨기기", self.toggle_job_panel))),
            ("도구", (("인쇄", self.print_with_quantity), ("전체 선택", self.select_all_labels),
                     ("선택 해제", self.clear_label_selection), ("상품 DB 연결 해제", self.disconnect_db_file),
                     ("새 행", self.add_row), ("선택 삭제", self.delete_selected),
                     ("삭제 실행취소", self.undo_delete), ("라벨디자이너", self.open_designer), ("프린터 설정", self.open_settings),
                     ("실행 전 점검", self.run_preflight_check))),
            ("도움말", (("빠른 사용안내 열기", self.open_quick_guide), ("고객용 매뉴얼", self.open_manual),
                      ("지원 패키지 생성", self.create_support_package))),
        )
        self.suite_menus = {}
        for name, commands in definitions:
            button = ttk.Menubutton(bar, text=name, style="Menu.TMenubutton")
            menu = tk.Menu(button, tearoff=0)
            for label, command in commands:
                menu.add_command(label=label, command=command)
            button.configure(menu=menu)
            button.pack(side="left", padx=(0, 4))
            self.suite_menus[name] = menu
            if name == "도구":
                self.output_menu = menu

    def _panel_card(self, parent, row: int, title: str) -> ttk.Frame:
        card = tk.Frame(parent, bg=COLORS.surface, highlightbackground=COLORS.border,
                        highlightthickness=1, bd=0)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 8))
        card.columnconfigure(0, weight=1)
        ttk.Label(card, text=title, style="CardTitle.TLabel").grid(row=0, column=0, sticky="ew")
        inner = ttk.Frame(card, style="SidePanel.TFrame", padding=(10, 8))
        inner.grid(row=1, column=0, sticky="ew")
        inner.columnconfigure(0, weight=1)
        return inner

    def toggle_job_panel(self) -> None:
        self._job_panel_visible = not self._job_panel_visible
        if self._job_panel_visible:
            self._job_panel_card.grid()
        else:
            self._job_panel_card.grid_remove()
        self._work_area.columnconfigure(1, minsize=self._scaled(288) if self._job_panel_visible else 0)

    def _build_start_checklist(self, parent: ttk.Frame) -> None:
        parent.columnconfigure(0, weight=1)
        ttk.Label(parent, text="시작 체크리스트", style="SidePanelTitle.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(
            parent,
            text="처음 설치, PC 교체, 프린터 변경 시 이 순서대로 확인하세요.",
            style="SidePanelBody.TLabel",
            wraplength=self._scaled(238),
        ).grid(row=1, column=0, sticky="w", pady=(self._scaled(5), self._scaled(11)))

        steps = (
            ("1. 프린터 설정", "장비, 라벨 크기, 인쇄후작업 저장", "열기", self.open_settings),
            ("2. DB 연결", "상품 DB 엑셀을 선택하고 인쇄 데이터 반영", "연결", self.connect_db_file),
            ("3. 실행 전 점검", "필수 파일, 설정, dry-run 결과 확인", "점검", self.run_preflight_check),
            ("4. 테스트 인쇄", "대상과 매수를 확인해 1장 출력", "인쇄", self.print_with_quantity),
        )
        responsive_rows: list[tuple[ttk.Frame, ttk.Label, ttk.Button]] = []
        for row, (title, description, button_text, command) in enumerate(steps, start=2):
            step = ttk.Frame(parent, style="Surface.TFrame")
            step.grid(row=row, column=0, sticky="ew", pady=(0, self._scaled(6)))
            step.columnconfigure(0, weight=1)
            ttk.Label(step, text=title, style="OnboardingStep.TLabel").grid(row=0, column=0, sticky="w")
            description_label = ttk.Label(
                step,
                text=description,
                style="SidePanelBody.TLabel",
                wraplength=self._scaled(180),
            )
            description_label.grid(row=1, column=0, sticky="w", pady=(self._scaled(3), 0))
            action_button = ttk.Button(step, text=button_text, command=command, style="Compact.TButton", width=6)
            action_button.grid(
                row=0,
                column=1,
                rowspan=2,
                sticky="e",
                padx=(self._scaled(8), 0),
            )
            step.bind("<Configure>", lambda event, description=description_label, button=action_button:
                      self._fit_checklist_description(event, description, button))
            if title == "4. 테스트 인쇄":
                self._print_action_widgets.append(action_button)
            responsive_rows.append((step, description_label, action_button))

        def apply_checklist_layout(compact: bool) -> None:
            for step, description_label, action_button in responsive_rows:
                step.grid_configure(pady=(0, self._scaled(3 if compact else 6)))
                if compact:
                    description_label.grid_remove()
                    action_button.grid_configure(rowspan=1)
                else:
                    description_label.grid()
                    action_button.grid_configure(rowspan=2)

        # At high Windows display scaling the checklist can request more height
        # than the work area before its first <Configure> event.  Start in the
        # compact state so the frame is mapped, then let subsequent resizes
        # choose the appropriate presentation.
        apply_checklist_layout(self._display_scale >= 1.4)

        def layout_checklist(event: tk.Event) -> None:
            compact = self._display_scale >= 1.4 or int(event.height) < self._scaled(280)
            apply_checklist_layout(compact)

        parent.bind("<Configure>", layout_checklist, add="+")

    def _fit_checklist_description(self, event: tk.Event, description: ttk.Label, button: ttk.Button) -> None:
        width = max(80, int(event.width) - button.winfo_reqwidth() - self._scaled(8) - 4)
        if int(description.cget("wraplength")) != width:
            description.configure(wraplength=width)

    def _build_job_panel(self, parent: ttk.Frame) -> None:
        ttk.Label(
            parent,
            text="출력 범위",
            style="SidePanelBody.TLabel",
        ).grid(row=1, column=0, sticky="w", pady=(self._scaled(6), self._scaled(5)))
        ttk.Label(parent, textvariable=self.print_scope_var, style="PrintScope.TLabel", anchor="w").grid(
            row=2,
            column=0,
            sticky="ew",
            pady=(0, self._scaled(9)),
        )

        selection_actions = ttk.Frame(parent, style="SidePanel.TFrame")
        selection_actions.grid(row=3, column=0, sticky="ew")
        selection_actions.columnconfigure(0, weight=1)
        selection_actions.columnconfigure(1, weight=1)
        ttk.Button(selection_actions, text="전체 선택", command=self.select_all_labels, style="Compact.TButton").grid(
            row=0,
            column=0,
            sticky="ew",
            padx=(0, self._scaled(4)),
        )
        ttk.Button(selection_actions, text="선택 해제", command=self.clear_label_selection, style="Compact.TButton").grid(
            row=0,
            column=1,
            sticky="ew",
            padx=(self._scaled(4), 0),
        )
        ttk.Separator(parent).grid(row=4, column=0, sticky="ew", pady=(self._scaled(12), self._scaled(12)))
        ttk.Button(parent, text="실행 전 점검", command=self.run_preflight_check, style="Secondary.TButton").grid(
            row=5,
            column=0,
            sticky="ew",
            pady=(0, self._scaled(8)),
        )
        ttk.Button(parent, text="프린터 설정", command=self.open_settings, style="Secondary.TButton").grid(
            row=6,
            column=0,
            sticky="ew",
        )
        ttk.Label(
            parent,
            text="지원 패키지는 도움말 메뉴에서 생성합니다.",
            style="SidePanelBody.TLabel",
            wraplength=self._scaled(238),
        ).grid(row=7, column=0, sticky="w", pady=(self._scaled(10), 0))

    def _create_table(self, tabs: ttk.Notebook, headers: tuple[str, ...], title: str) -> ttk.Treeview:
        frame = ttk.Frame(tabs, padding=SPACING.form_gap, style="Surface.TFrame")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        tree = ttk.Treeview(frame, columns=headers, show="headings", selectmode="browse")
        self._configure_table_columns(tree, headers)
        tree.tag_configure("checked", background=COLORS.accent_soft)
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        if headers and headers[0] == SELECT_HEADER:
            tree.bind("<Button-1>", self._handle_label_click)
            tree.bind("<space>", self._toggle_focused_label)
        tree.bind("<Double-1>", lambda event, active_tree=tree: self.begin_edit(active_tree, event))
        tree.bind("<Return>", lambda event, active_tree=tree: self.begin_edit(active_tree, event))
        tabs.add(frame, text=title)
        return tree

    def _configure_table_columns(self, tree: ttk.Treeview, headers: tuple[str, ...]) -> None:
        tree.configure(columns=headers)
        for header in headers:
            if header == SELECT_HEADER:
                width = 58
                minwidth = 52
                stretch = False
            elif _field_for_header(header) in {"item_name", "barcode"}:
                width = 250
                minwidth = 120
                stretch = True
            else:
                width = max(120, min(220, len(str(header)) * 16 + 70))
                minwidth = 90
                stretch = True
            tree.heading(header, text=COLUMN_LABELS.get(header, str(header)))
            anchor = "center" if header == SELECT_HEADER else "e" if str(header) in {"판매가", "가격", "price", "sale_price"} else "w"
            tree.column(header, width=width, minwidth=minwidth, stretch=stretch, anchor=anchor)

    def load_files(self) -> None:
        try:
            (
                self.db_headers,
                self.db_rows,
                self.label_headers,
                self.label_rows,
            ) = _load_active_db_state(self.db_path)
            self.label_display_headers = (SELECT_HEADER, *self.label_headers)
            self.selected_label_indexes.clear()
        except Exception as exc:
            messagebox.showerror("\ub370\uc774\ud130 \uc77d\uae30 \uc2e4\ud328", str(exc))
            return
        self.refresh_tables()
        self._set_data_dirty(False)
        self.status_var.set(f"DB {len(self.db_rows)}건을 출력 목록으로 불러왔습니다.")

    def refresh_tables(self) -> None:
        self.label_display_headers = (SELECT_HEADER, *self.label_headers)
        self._configure_table_columns(self.labels_tree, self.label_display_headers)
        self._configure_table_columns(self.db_tree, self.db_headers)
        _fill_tree(self.labels_tree, self.label_display_headers, self.label_rows, self.selected_label_indexes)
        _fill_tree(self.db_tree, self.db_headers, self.db_rows)
        self._update_print_scope()
        self._update_document_header()

    def _update_print_scope(self) -> None:
        total_count = len(self.__dict__.get("label_rows", []))
        selected_indexes = self.__dict__.get("selected_label_indexes", set())
        selected_count = sum(1 for index in selected_indexes if 0 <= index < total_count)
        if total_count == 0:
            scope_text = "인쇄할 데이터 없음"
        elif selected_count:
            scope_text = f"선택 {selected_count}건 인쇄"
        else:
            scope_text = f"전체 {total_count}건 인쇄 · 선택 없음"

        scope_var = self.__dict__.get("print_scope_var")
        if scope_var is not None:
            scope_var.set(scope_text)

        enabled = total_count > 0
        for widget in self.__dict__.get("_print_action_widgets", []):
            widget.state(["!disabled"] if enabled else ["disabled"])
        output_menu = self.__dict__.get("output_menu")
        if output_menu is not None:
            output_menu.entryconfigure(0, state="normal" if enabled else "disabled")

    def add_row(self) -> None:
        if self.tabs.index(self.tabs.select()) == 0:
            self.label_rows.append({header: "" for header in self.label_headers})
            self.refresh_tables()
            self._select_last(self.labels_tree)
        else:
            self.db_rows.append({header: "" for header in self.db_headers})
            self._sync_labels_from_db(preserve_selection=True)
            self.refresh_tables()
            self._select_last(self.db_tree)
        self._set_data_dirty(True)

    def delete_selected(self) -> None:
        tree, rows = self._active_tree_and_rows()
        selected = tree.selection()
        is_print_data = tree is self.labels_tree
        focused_index = tree.index(selected[0]) if selected else None
        if is_print_data:
            indexes = {index for index in self.selected_label_indexes if 0 <= index < len(rows)}
            if not indexes and focused_index is not None:
                indexes = {focused_index}
        else:
            indexes = {tree.index(item) for item in selected}
        if not indexes:
            self.status_var.set("삭제할 항목을 선택하세요.")
            return
        target = "인쇄 데이터" if is_print_data else "원본 DB"
        if not messagebox.askyesno(
            "선택 삭제 확인", f"{target} {len(indexes)}건을 삭제할까요?\n"
            "현재 화면에서 삭제합니다. 다른 수정이나 창 닫기 전에는 삭제 실행취소로 복구할 수 있습니다.", parent=self):
            return
        history = self.__dict__.setdefault("_delete_history", [])
        history.append({"label_rows": copy.deepcopy(self.label_rows), "db_rows": copy.deepcopy(self.db_rows),
                        "selected": set(self.selected_label_indexes), "target": target,
                        "context": self._delete_context()})
        del history[:-5]
        if tree is self.labels_tree:
            deleted_count = _delete_label_rows(self.label_rows, self.selected_label_indexes, focused_index)
            if not deleted_count:
                return
            self.refresh_tables()
        else:
            for index in sorted(indexes, reverse=True):
                del rows[index]
            self._sync_labels_from_db(preserve_selection=True)
            self.refresh_tables()
        history[-1]["expected_label_rows"] = copy.deepcopy(self.label_rows)
        history[-1]["expected_db_rows"] = copy.deepcopy(self.db_rows)
        self._set_data_dirty(True, clear_delete_history=False)
        self.status_var.set(f"{target} {len(indexes)}건 삭제 · 삭제 실행취소로 복구 가능 · 저장 필요")

    def undo_delete(self) -> None:
        history = self.__dict__.get("_delete_history", [])
        if not history:
            self.status_var.set("실행취소할 삭제가 없습니다.")
            return
        latest = history[-1]
        if (latest.get("context") != self._delete_context()
                or latest.get("expected_label_rows") != self.label_rows
                or latest.get("expected_db_rows") != self.db_rows):
            self._set_data_dirty(self.__dict__.get("_data_dirty", False))
            self.status_var.set("삭제 후 데이터가 변경되어 이전 삭제를 실행취소할 수 없습니다.")
            return
        snapshot = history.pop()
        self.label_rows = copy.deepcopy(snapshot["label_rows"])
        self.db_rows = copy.deepcopy(snapshot["db_rows"])
        self.selected_label_indexes = set(snapshot["selected"])
        self.refresh_tables()
        self._set_data_dirty(True, clear_delete_history=False)
        self.status_var.set(f"{snapshot['target']} 삭제를 실행취소했습니다. 복구 내용을 저장하세요.")

    def _delete_context(self) -> tuple[object, ...]:
        return (str(self.__dict__.get("db_path", "")),
                tuple(self.__dict__.get("db_headers", ())),
                tuple(self.__dict__.get("label_headers", ())))

    def _on_undo_delete_shortcut(self, event) -> str | None:
        if str(event.widget.winfo_class()) in {"Entry", "TEntry", "TSpinbox", "TCombobox", "Text"}:
            return None
        self.undo_delete()
        return "break"

    def _set_data_dirty(self, dirty: bool, *, clear_delete_history: bool = True) -> None:
        self._data_dirty = dirty
        if clear_delete_history:
            self._delete_history = []
        self._update_document_header()
        button = self.__dict__.get("undo_delete_button")
        if button is not None:
            button.state(["!disabled"] if self.__dict__.get("_delete_history") else ["disabled"])

    def _update_document_header(self) -> None:
        name = self.__dict__.get("document_name_var")
        state = self.__dict__.get("document_state_var")
        if name is not None:
            path = self.__dict__.get("db_path")
            name.set(str(self.__dict__.get("_db_display_name") or (path.name if path else "상품 DB")))
        if state is not None:
            if not self.__dict__.get("db_connected", True):
                state.set("상품 엑셀 미연결")
            else:
                state.set(f"{'변경사항 있음 · 저장 필요' if self.__dict__.get('_data_dirty') else '저장됨'} · 전체 {len(self.label_rows)}건")

    def _confirm_replace_data(self, action: str) -> bool:
        if not self.__dict__.get("_data_dirty", False):
            return True
        choice = messagebox.askyesnocancel("저장하지 않은 변경사항", f"{action} 전에 현재 변경사항을 저장할까요?", parent=self)
        if choice is None:
            return False
        return self.save_files() if choice else True

    def _confirm_close(self) -> None:
        if self._confirm_replace_data("닫기"):
            self.destroy()

    def save_files(self) -> bool:
        if not self.__dict__.get("db_connected", True) and not self.db_rows and not self.label_rows:
            self.status_var.set("DB 해제 상태입니다. 기존 파일 보호를 위해 저장하지 않았습니다.")
            return False
        try:
            label_headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
            db_headers = self.__dict__.get("db_headers", tuple(DB_HEADERS))
            save_label_rows(self.queue_path, _printable_rows_from_label_rows(self.label_rows, label_headers))
            _save_dynamic_db_rows(self.db_path, db_headers, self.db_rows)
        except Exception as exc:
            messagebox.showerror("\uc800\uc7a5 \uc2e4\ud328", str(exc))
            return False
        self._set_data_dirty(False, clear_delete_history=False)
        self.status_var.set("\uc800\uc7a5\ub418\uc5c8\uc2b5\ub2c8\ub2e4.")
        return True

    def connect_db_file(self) -> None:
        if not self._confirm_replace_data("상품 엑셀 연결"):
            return
        source_name = filedialog.askopenfilename(
            parent=self,
            title="DB 연결",
            filetypes=[
                ("Excel 통합 문서", "*.xlsx *.xlsm"),
                ("모든 파일", "*.*"),
            ],
        )
        if not source_name:
            return

        source = Path(source_name)
        try:
            self.db_upload_dir.mkdir(parents=True, exist_ok=True)
            destination = self.db_upload_dir / source.name
            if source.resolve() != destination.resolve():
                shutil.copy2(source, destination)
            (
                self.db_headers,
                self.db_rows,
                self.label_headers,
                self.label_rows,
            ) = _load_active_db_state(destination)
            self.label_display_headers = (SELECT_HEADER, *self.label_headers)
            self.selected_label_indexes.clear()
            _save_dynamic_db_rows(self.db_path, self.db_headers, self.db_rows)
            save_label_rows(self.queue_path, _printable_rows_from_label_rows(self.label_rows, self.label_headers))
            selected_queue = _selected_print_queue_path(self.base_dir)
            if selected_queue.exists():
                selected_queue.unlink()
        except Exception as exc:
            messagebox.showerror("DB 연결 실패", str(exc))
            return

        self.db_connected = True
        self._db_display_name = source.name
        self._set_data_dirty(False)
        self.refresh_tables()
        self.tabs.select(self.tabs.tabs()[0])
        self.status_var.set(f"DB 연결 완료: {len(self.db_rows)}건 / 저장 위치: db\\{destination.name}")

    def upload_db_excel(self) -> None:
        self.connect_db_file()

    def disconnect_db_file(self) -> None:
        if not self._confirm_replace_data("상품 DB 연결 해제"):
            return
        self.db_connected = False
        self.db_headers = tuple(DB_HEADERS)
        self.label_headers = tuple(LABEL_HEADERS)
        self.label_display_headers = (SELECT_HEADER, *self.label_headers)
        self.db_rows = []
        self.label_rows = []
        self.selected_label_indexes.clear()
        self._set_data_dirty(False)
        self.refresh_tables()
        self.status_var.set("DB 연결을 해제했습니다. 기존 파일은 삭제하지 않았습니다.")

    def scan_barcode(self, event: tk.Event | None = None) -> None:
        query = self.scan_var.get().strip()
        if not query:
            return
        db_matches = self._db_rows_by_query(query)
        if len(db_matches) > 1:
            selected_matches = self._choose_db_matches(query, db_matches)
            if selected_matches:
                db_index, db_row = selected_matches[0]
                self._select_db_query_result(query, db_index, db_row, show_tab=False)
                self._select_label_rows_for_db_matches(query, selected_matches)
            else:
                self.status_var.set(f"DB 중복 조회 취소: {query}")
        elif len(db_matches) == 1:
            db_index, db_row = db_matches[0]
            self._select_db_query_result(query, db_index, db_row, show_tab=False)
            barcode = _row_field_value(db_row, "barcode")
            if barcode:
                self.select_label_row_by_barcode(barcode)
            else:
                self.select_label_row_by_query(query)
        elif self.select_label_row_by_barcode(query):
            self.select_db_row_by_barcode(query, show_tab=False, show_warning=False)
        self.scan_var.set("")

    def select_label_row_by_barcode(self, barcode: str, show_warning: bool = True) -> bool:
        target = barcode.strip()
        for index, row in enumerate(self.label_rows):
            if _row_field_value(row, "barcode") == target:
                self.selected_label_indexes.add(index)
                self.refresh_tables()
                self.tabs.select(self.tabs.tabs()[0])
                self._select_index(self.labels_tree, index)
                label_headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
                label_value = _row_field_value(row, "item_name") or _first_non_barcode_value(row, label_headers)
                self.status_var.set(f"인쇄 데이터 선택: {target} / {label_value}")
                return True
        if show_warning:
            messagebox.showwarning("인쇄 데이터", "인쇄 데이터에서 해당 바코드를 찾을 수 없습니다.")
        return False

    def select_db_row_by_barcode(self, barcode: str, show_tab: bool = True, show_warning: bool = True) -> bool:
        target = barcode.strip()
        matches = self._db_rows_by_query(target)
        if matches:
            index, row = matches[0]
            self._select_db_query_result(target, index, row, show_tab=show_tab)
            return True
        if show_warning:
            messagebox.showwarning("\ubc14\ucf54\ub4dc DB", "DB\uc5d0 \uc5c6\ub294 \ubc14\ucf54\ub4dc \uc785\ub2c8\ub2e4.")
        return False

    def select_label_row_by_query(self, query: str, show_warning: bool = True) -> bool:
        target = query.strip()
        match = _find_row_by_query(self.label_rows, target)
        if match is None:
            if show_warning:
                messagebox.showwarning("인쇄 데이터", "인쇄 데이터에서 해당 항목을 찾을 수 없습니다.")
            return False
        index, row = match
        self.selected_label_indexes.add(index)
        self.refresh_tables()
        self.tabs.select(self.tabs.tabs()[0])
        self._select_index(self.labels_tree, index)
        barcode = _row_field_value(row, "barcode")
        label_headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
        label_value = _row_field_value(row, "item_name") or _first_non_barcode_value(row, label_headers)
        self.status_var.set(f"인쇄 데이터 선택: {barcode or target} / {label_value}")
        return True

    def _db_row_by_query(self, query: str) -> tuple[int, dict[str, str]] | None:
        matches = self._db_rows_by_query(query)
        return matches[0] if matches else None

    def _db_rows_by_query(self, query: str) -> list[tuple[int, dict[str, str]]]:
        return _matching_rows_by_query(self.db_rows, query)

    def _select_db_query_result(self, query: str, index: int, row: dict[str, str], *, show_tab: bool = True) -> None:
        if show_tab:
            self.tabs.select(self.tabs.tabs()[1])
        self._select_index(self.db_tree, index)
        barcode = _row_field_value(row, "barcode")
        item_name = _row_field_value(row, "item_name") or _first_non_barcode_value(row, self.db_headers)
        self.status_var.set(f"DB \uc870\ud68c \uc644\ub8cc: {query} -> {barcode or '-'} / {item_name}")

    def _select_label_rows_for_db_matches(self, query: str, matches: list[tuple[int, dict[str, str]]]) -> int:
        selected_indexes: set[int] = set()
        for db_index, db_row in matches:
            selected_indexes.update(self._label_indexes_for_db_match(db_index, db_row, query))
        if not selected_indexes:
            messagebox.showwarning("인쇄 데이터", "선택한 DB 항목과 일치하는 인쇄 데이터를 찾을 수 없습니다.")
            return 0
        self.selected_label_indexes = set(selected_indexes)
        self.refresh_tables()
        self.tabs.select(self.tabs.tabs()[0])
        first_index = min(selected_indexes)
        self._select_index(self.labels_tree, first_index)
        self.status_var.set(
            f"인쇄 데이터 적용 완료: {len(selected_indexes)}건 · 선택 항목 인쇄를 누르세요."
        )
        return len(selected_indexes)

    def _label_indexes_for_db_match(self, db_index: int, db_row: dict[str, str], query: str) -> list[int]:
        barcode = _row_field_value(db_row, "barcode")
        normalized_query = _normalize_search_text(query)
        if 0 <= db_index < len(self.label_rows):
            label_row = self.label_rows[db_index]
            if barcode and _row_field_value(label_row, "barcode") == barcode:
                return [db_index]
            if not barcode and _row_matches_query(label_row, normalized_query):
                return [db_index]
        if barcode:
            barcode_matches = [
                index
                for index, row in enumerate(self.label_rows)
                if _row_field_value(row, "barcode") == barcode
            ]
            return barcode_matches[:1]
        return [index for index, _row in _matching_rows_by_query(self.label_rows, query)]

    def _choose_db_matches(self, query: str, matches: list[tuple[int, dict[str, str]]]) -> list[tuple[int, dict[str, str]]]:
        dialog = tk.Toplevel(self)
        dialog.title("중복 DB 항목 선택")
        dialog.transient(self)
        dialog.grab_set()
        set_initial_window_size(
            dialog,
            preferred_width=960,
            preferred_height=620,
            minimum_width=760,
            minimum_height=520,
        )
        dialog.configure(bg=COLORS.surface)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(0, weight=1)
        selected_positions = _initial_duplicate_selection_positions(len(matches))
        result: list[tuple[int, dict[str, str]]] = []

        container = ttk.Frame(dialog, style="Surface.TFrame", padding=16)
        container.grid(row=0, column=0, sticky="nsew")
        container.columnconfigure(0, weight=1)
        container.rowconfigure(1, weight=1)
        ttk.Label(
            container,
            text=f"'{query}' 검색 결과가 {len(matches)}건 있습니다. 체크 후 확인을 누르면 인쇄 데이터에 적용됩니다.",
            style="TLabel",
            wraplength=860,
        ).grid(row=0, column=0, sticky="ew", pady=(0, 10))

        columns = (SELECT_HEADER, "__row__", *self.db_headers)
        tree_frame = ttk.Frame(container, style="Surface.TFrame")
        tree_frame.grid(row=1, column=0, sticky="nsew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        tree = ttk.Treeview(tree_frame, columns=columns, show="headings", selectmode="browse")
        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        tree.heading(SELECT_HEADER, text="선택")
        tree.column(SELECT_HEADER, width=58, minwidth=52, stretch=False, anchor="center")
        tree.heading("__row__", text="DB 행")
        tree.column("__row__", width=70, minwidth=60, stretch=False, anchor="center")
        for header in self.db_headers:
            tree.heading(header, text=COLUMN_LABELS.get(header, str(header)))
            tree.column(header, width=160, minwidth=90, stretch=True, anchor="w")

        selected_count_var = tk.StringVar(value=f"선택 {len(selected_positions)}건")

        def refresh_tree() -> None:
            tree.delete(*tree.get_children())
            for position, (db_index, row) in enumerate(matches):
                values = [
                    CHECKED if position in selected_positions else UNCHECKED,
                    str(db_index + 1),
                    *[str(row.get(header, "")).strip() for header in self.db_headers],
                ]
                tree.insert("", "end", iid=str(position), values=values, tags=("checked",) if position in selected_positions else ())
            selected_count_var.set(f"선택 {len(selected_positions)}건")

        def toggle_position(position: int) -> None:
            if position in selected_positions:
                selected_positions.remove(position)
            else:
                selected_positions.add(position)
            refresh_tree()
            tree.focus(str(position))
            tree.selection_set(str(position))

        def on_tree_click(event: tk.Event) -> None:
            row_id = tree.identify_row(event.y)
            column = tree.identify_column(event.x)
            if row_id and column == "#1":
                toggle_position(int(row_id))

        def select_all() -> None:
            selected_positions.update(range(len(matches)))
            refresh_tree()

        def clear_all() -> None:
            selected_positions.clear()
            refresh_tree()

        def accept() -> None:
            nonlocal result
            if not selected_positions:
                messagebox.showwarning("중복 조회", "확인할 항목을 먼저 체크하세요.", parent=dialog)
                return
            result = _duplicate_selection_result(matches, selected_positions)
            dialog.destroy()

        def cancel() -> None:
            dialog.destroy()

        tree.tag_configure("checked", background=COLORS.accent_soft)
        def on_tree_space(_event: tk.Event) -> str:
            focused = tree.focus()
            if focused:
                toggle_position(int(focused))
            return "break"

        tree.bind("<Button-1>", on_tree_click)
        tree.bind("<space>", on_tree_space)
        refresh_tree()

        button_row = ttk.Frame(container, style="Surface.TFrame")
        button_row.grid(row=2, column=0, sticky="ew", pady=(12, 0))
        button_row.columnconfigure(1, weight=1)
        selection_button = ttk.Menubutton(button_row, text="선택 메뉴")
        selection_menu = tk.Menu(selection_button, tearoff=0)
        selection_menu.add_command(label="전체 선택", command=select_all)
        selection_menu.add_command(label="선택 해제", command=clear_all)
        selection_button.configure(menu=selection_menu)
        selection_button.grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Label(button_row, textvariable=selected_count_var, style="ToolbarStatus.TLabel").grid(row=0, column=1, sticky="w")
        ttk.Button(button_row, text="확인", command=accept, style="Primary.TButton").grid(row=0, column=2, padx=(8, 8))
        ttk.Button(button_row, text="취소", command=cancel, style="Secondary.TButton").grid(row=0, column=3)
        dialog.protocol("WM_DELETE_WINDOW", cancel)
        dialog.bind("<Return>", lambda _event: accept())
        dialog.bind("<Escape>", lambda _event: cancel())
        self.wait_window(dialog)
        return result

    def apply_barcode_to_label_row(
        self,
        barcode: str,
        row_index: int | None = None,
        *,
        clear_selection_after: bool = False,
    ) -> int | None:
        result = _lookup_dynamic_barcode(self.db_rows, barcode)
        if result is None:
            messagebox.showwarning("\ubc14\ucf54\ub4dc DB", "DB\uc5d0 \uc5c6\ub294 \ubc14\ucf54\ub4dc \uc785\ub2c8\ub2e4.")
            return None
        if row_index is None:
            row_index = self._selected_label_index_for_scan()
        row_index, appended = _ensure_label_row_index(self.label_rows, row_index, self.label_headers)
        self.label_rows[row_index].update(_label_row_from_db_row(result, self.label_headers))
        self._set_data_dirty(True)
        self.refresh_tables()
        if clear_selection_after:
            self._see_index(self.labels_tree, row_index)
            self.labels_tree.selection_remove(self.labels_tree.selection())
        else:
            self._select_index(self.labels_tree, row_index)
        action = "\ucd9c\ub825 \ubaa9\ub85d\uc5d0 \ucd94\uac00" if appended else "\ucd9c\ub825 \ud589 \uc790\ub3d9 \ucc44\uc6c0"
        barcode_value = _row_field_value(result, "barcode")
        label_value = _row_field_value(result, "item_name") or _first_non_barcode_value(result, self.db_headers)
        self.status_var.set(f"{action}: {barcode_value} / {label_value}")
        return row_index

    def begin_edit(self, tree: ttk.Treeview, event: tk.Event) -> None:
        if self.active_editor is not None:
            self.active_editor.destroy()
            self.active_editor = None
        item = tree.identify_row(event.y) if hasattr(event, "y") else (tree.selection()[0] if tree.selection() else "")
        column_id = tree.identify_column(event.x) if hasattr(event, "x") else "#1"
        if not item or column_id == "#0":
            return
        column_index = int(column_id.replace("#", "")) - 1
        headers = self.label_display_headers if tree is self.labels_tree else self.db_headers
        if column_index < 0 or column_index >= len(headers):
            return
        header = headers[column_index]
        if header == SELECT_HEADER:
            self._toggle_label_index(tree.index(item))
            return
        bbox = tree.bbox(item, column_id)
        if not bbox:
            return
        x, y, width, height = bbox
        value = tree.set(item, header)
        editor = ttk.Entry(tree)
        editor.insert(0, value)
        editor.select_range(0, "end")
        editor.place(x=x, y=y, width=width, height=height)
        editor.focus_set()
        self.active_editor = editor
        committed = {"done": False}

        def commit(_event: tk.Event | None = None) -> None:
            if committed["done"]:
                return
            committed["done"] = True
            new_value = editor.get().strip()
            editor.destroy()
            self.active_editor = None
            self._commit_cell(tree, item, header, new_value)

        editor.bind("<Return>", commit)
        editor.bind("<FocusOut>", commit)
        editor.bind("<Escape>", lambda _event: editor.destroy())

    def _commit_cell(self, tree: ttk.Treeview, item: str, header: str, value: str) -> None:
        rows = self.label_rows if tree is self.labels_tree else self.db_rows
        index = tree.index(item)
        rows[index][header] = value
        self._set_data_dirty(True)
        if tree is self.labels_tree and _field_for_header(header) == "barcode" and value:
            self.apply_barcode_to_label_row(value, row_index=index)
            return
        if tree is self.db_tree:
            self._sync_labels_from_db(preserve_selection=True)
        self.refresh_tables()
        self._select_index(tree, index)

    def print_with_quantity(self) -> None:
        selected_only = any(0 <= index < len(self.label_rows) for index in self.selected_label_indexes)
        rows_to_print = self._print_target_rows(selected_only)
        if not rows_to_print:
            messagebox.showwarning("\ucd9c\ub825 \ubaa9\ub85d", "\ucd9c\ub825\ud560 \ud56d\ubaa9\uc774 \uc5c6\uc2b5\ub2c8\ub2e4.")
            return
        print_quantity = self.ask_print_quantity(
            _default_print_quantity(rows_to_print),
            selected_only=selected_only,
        )
        if print_quantity is None:
            self.status_var.set("인쇄 매수 선택을 취소했습니다.")
            return
        self.run_print_job(
            "--print",
            selected_only=selected_only,
            print_quantity=print_quantity,
        )

    def _print_target_rows(self, selected_only: bool) -> list[dict[str, str]]:
        rows = self._selected_print_rows() if selected_only else list(self.label_rows)
        headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
        return [row for row in rows if _printable_rows_from_label_rows([row], headers)]

    def ask_print_quantity(self, default_quantity: int = 1, *, selected_only: bool = False) -> int | None:
        dialog = tk.Toplevel(self)
        dialog.title("인쇄 매수 선택")
        set_initial_window_size(
            dialog,
            preferred_width=520,
            preferred_height=500,
            minimum_width=440,
            minimum_height=460,
        )
        dialog.configure(bg=COLORS.surface)
        dialog.transient(self)
        dialog.grab_set()
        dialog.resizable(True, True)
        dialog.columnconfigure(0, weight=1)
        dialog.rowconfigure(0, weight=1)

        frame = ttk.Frame(dialog, style="Surface.TFrame", padding=22)
        frame.grid(row=0, column=0, sticky="nsew")
        frame.columnconfigure(0, weight=1)
        scope_text = "선택 항목" if selected_only else "전체 항목 (선택 없음)"
        target_count = len(self._print_target_rows(selected_only))
        ttk.Label(frame, text="몇 장씩 인쇄할까요?", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(
            frame,
            text=f"{scope_text}의 각 항목에 같은 수량이 적용됩니다. 1장부터 100장까지 선택할 수 있습니다.",
            style="Hint.TLabel",
            wraplength=470,
            justify="left",
        ).grid(row=1, column=0, sticky="w", pady=(6, 16))

        quantity_card = tk.Frame(
            frame,
            bg=COLORS.surface_subtle,
            highlightbackground=COLORS.border_subtle,
            highlightcolor=COLORS.border,
            highlightthickness=1,
            bd=0,
        )
        quantity_card.grid(row=2, column=0, sticky="ew")
        quantity_card.columnconfigure(1, weight=1)
        tk.Label(
            quantity_card,
            text="인쇄 매수",
            bg=COLORS.surface_subtle,
            fg=COLORS.text_secondary,
            font=TYPOGRAPHY.caption,
        ).grid(row=0, column=0, columnspan=3, sticky="w", padx=16, pady=(14, 4))

        initial = max(1, min(100, default_quantity))
        qty_var = tk.StringVar(value=str(initial))
        qty_input = ttk.Spinbox(
            quantity_card,
            from_=1,
            to=100,
            textvariable=qty_var,
            width=7,
            justify="center",
            font=(TYPOGRAPHY.body[0], 24, "bold"),
        )
        qty_input.grid(row=1, column=1, sticky="ew", ipady=6, pady=(0, 14))

        def adjust(delta: int) -> None:
            try:
                value = int(qty_var.get())
            except (TypeError, ValueError):
                value = initial
            qty_var.set(str(max(1, min(100, value + delta))))

        ttk.Button(
            quantity_card,
            text="-",
            command=lambda: adjust(-1),
            style="Secondary.TButton",
            width=3,
        ).grid(row=1, column=0, sticky="e", padx=(16, 10), pady=(0, 14))
        ttk.Button(
            quantity_card,
            text="+",
            command=lambda: adjust(1),
            style="Secondary.TButton",
            width=3,
        ).grid(row=1, column=2, sticky="w", padx=(10, 16), pady=(0, 14))

        presets = ttk.Frame(frame, style="Surface.TFrame")
        presets.grid(row=3, column=0, sticky="ew", pady=(12, 0))
        ttk.Label(presets, text="빠른 선택", style="Hint.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 8))
        for column, value in enumerate((1, 3, 5, 10), start=1):
            ttk.Button(
                presets,
                text=f"{value}장",
                command=lambda selected=value: qty_var.set(str(selected)),
                style="Secondary.TButton",
                width=5,
            ).grid(row=0, column=column, padx=(0, 6), sticky="ew")

        summary_var = tk.StringVar()

        def refresh_summary(*_args: object) -> None:
            try:
                value = int(qty_var.get())
            except (TypeError, ValueError):
                value = initial
            value = max(1, min(100, value))
            summary_var.set(f"{scope_text} {target_count}건 × 각 {value}장 = 총 {target_count * value}장\n"
                            "인쇄 시작을 누르면 프린터로 전송합니다.")

        qty_var.trace_add("write", refresh_summary)
        refresh_summary()
        ttk.Label(frame, textvariable=summary_var, style="Hint.TLabel", wraplength=440).grid(row=4, column=0, sticky="w", pady=(12, 0))

        result: dict[str, int | None] = {"quantity": None}

        def cancel() -> None:
            dialog.destroy()

        def confirm() -> None:
            try:
                quantity = int(qty_var.get())
            except (TypeError, ValueError):
                messagebox.showwarning("인쇄 매수", "인쇄 매수는 1부터 100 사이 숫자로 입력해 주세요.", parent=dialog)
                return
            if quantity < 1 or quantity > 100:
                messagebox.showwarning("인쇄 매수", "인쇄 매수는 1부터 100 사이로 선택해 주세요.", parent=dialog)
                return
            result["quantity"] = quantity
            dialog.destroy()

        buttons = ttk.Frame(frame, style="Surface.TFrame")
        buttons.grid(row=5, column=0, sticky="ew", pady=(18, 0))
        buttons.columnconfigure(0, weight=1)
        ttk.Button(buttons, text="취소", command=cancel, style="Secondary.TButton").grid(row=0, column=1, padx=(0, 8))
        ttk.Button(buttons, text="인쇄 시작", command=confirm, style="Primary.TButton").grid(row=0, column=2)
        dialog.bind("<Return>", lambda _event: confirm())
        dialog.bind("<Escape>", lambda _event: cancel())
        qty_input.focus_set()
        qty_input.select_range(0, "end")
        self.wait_window(dialog)
        return result["quantity"]

    def run_print_job(self, action: str, selected_only: bool = False, print_quantity: int | None = None) -> None:
        label_headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
        rows_to_print = self._print_target_rows(selected_only)
        if not rows_to_print:
            messagebox.showwarning("\ucd9c\ub825 \ubaa9\ub85d", "\ucd9c\ub825\ud560 \ud56d\ubaa9\uc774 \uc5c6\uc2b5\ub2c8\ub2e4.")
            return
        readiness_errors = _print_readiness_errors(
            rows_to_print,
            label_headers,
            self.config_path,
            print_quantity=print_quantity,
        )
        if readiness_errors:
            self.status_var.set("인쇄 전 점검에서 수정할 항목이 발견되었습니다.")
            messagebox.showwarning("인쇄 전 점검", "\n".join(f"- {error}" for error in readiness_errors))
            return
        try:
            # Printing snapshots the current screen without saving the customer's DB or default queue.
            excel_path = _selected_print_queue_path(self.base_dir)
            excel_path.parent.mkdir(parents=True, exist_ok=True)
            save_label_rows(excel_path, _rows_with_print_quantity(
                _printable_rows_from_label_rows(rows_to_print, label_headers), print_quantity))
        except Exception as exc:
            messagebox.showerror("\uc800\uc7a5 \uc2e4\ud328", str(exc))
            return
        if self.print_exe is None or not self.print_exe.exists():
            messagebox.showerror("\ucd9c\ub825 \uc2e4\ud328", f"라벨출력엔진.exe\ub97c \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.\n{self.print_exe}")
            return
        args = _build_print_args(self.print_exe, self.config_path, action, excel_path)
        try:
            result = self._invoke_print_engine(args)
        except OSError as exc:
            messagebox.showerror("\ucd9c\ub825 \uc2e4\ud328", str(exc))
            return
        if action == "--print" and result.returncode != 0:
            try:
                recovered = self._recover_print_progress(args, result)
            except (OSError, PrintProgressError) as exc:
                self.status_var.set("인쇄 오류")
                messagebox.showerror("인쇄 오류", str(exc))
                return
            if recovered is None:
                return
            result = recovered
        if result.returncode != 0:
            self.status_var.set("인쇄 오류")
            messagebox.showerror("인쇄 오류", (result.stderr or result.stdout or "").strip())
            return
        title = "출력 파일 생성 완료" if action == "--dry-run" else "프린터 전송 완료 · 실제 출력 확인 필요"
        if selected_only:
            title = f"\uc120\ud0dd {len(rows_to_print)}건 {title}"
        detail = (result.stdout or "").strip()
        self.status_var.set(f"{title}: {detail}" if detail else title)
        if action == "--print":
            messagebox.showinfo(
                "프린터 전송 완료 · 출력 확인 필요",
                f"{len(rows_to_print)}건의 인쇄 명령을 프린터로 전송했습니다.\n실제 라벨 출력 여부는 프린터에서 확인하세요.",
            )

    def _invoke_print_engine(self, args: list[str]) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            args,
            cwd=self.base_dir,
            text=True,
            capture_output=True,
            check=False,
            creationflags=_creationflags(),
        )

    def _recover_print_progress(
        self,
        args: list[str],
        result: subprocess.CompletedProcess[str],
    ) -> subprocess.CompletedProcess[str] | None:
        output = (result.stderr or result.stdout or "").strip()
        if NEW_JOB_REQUIRED_CODE in output:
            result = self._invoke_print_engine([*args, "--new-job"])
            output = (result.stderr or result.stdout or "").strip()

        if result.returncode == 0:
            return result
        if UNKNOWN_ITEMS_CODE not in output and TRANSPORT_UNCERTAIN_CODE not in output:
            return result

        progress_path = load_config(self.config_path).data.output_dir / PROGRESS_FILE_NAME
        PrintProgress.load(progress_path)
        self.status_var.set("인쇄 오류 · 이전 인쇄 작업의 전송 상태를 확인해야 합니다.")
        messagebox.showerror(
            "인쇄 오류",
            (
                "이전 인쇄 작업 중 일부의 전송 상태를 확인할 수 없습니다.\n"
                "중복 출력을 막기 위해 자동 재전송하지 않았습니다.\n\n"
                f"작업 상태 파일: {progress_path}\n"
                "프린터에서 실제 출력 여부를 확인한 뒤 새 인쇄 작업을 시작하세요."
            ),
        )
        return None

    def open_settings(self) -> None:
        if self.settings_exe is None:
            messagebox.showwarning("\ud504\ub9b0\ud130 \uc124\uc815", "\ud504\ub9b0\ud130\uc124\uc815.exe\ub97c \ucc3e\uc744 \uc218 \uc5c6\uc2b5\ub2c8\ub2e4.")
            return
        subprocess.Popen([str(self.settings_exe), "--config", str(self.config_path)], cwd=self.base_dir, creationflags=_creationflags())

    def open_designer(self) -> None:
        if self.designer_exe is None:
            messagebox.showwarning("라벨 디자인", "라벨디자이너.exe를 찾을 수 없습니다. 배포 폴더 전체를 복사했는지 확인하세요.")
            return
        subprocess.Popen([str(self.designer_exe)], cwd=self.base_dir, creationflags=_creationflags())

    def _open_customer_data_script(self, name: str) -> None:
        script = self.base_dir / name
        if not script.is_file():
            messagebox.showwarning("고객 데이터", f"{name} 파일을 찾을 수 없습니다.")
            return
        subprocess.Popen(["cmd.exe", "/c", str(script)], cwd=self.base_dir)

    def open_backup(self) -> None:
        self._open_customer_data_script("고객데이터_백업.cmd")

    def open_restore(self) -> None:
        self._open_customer_data_script("고객데이터_복원.cmd")

    def open_quick_guide(self) -> None:
        self._open_document(self.quick_guide_path, "빠른 사용안내")

    def open_manual(self) -> None:
        self._open_document(self.manual_path, "고객용 매뉴얼")

    def _open_document(self, path: Path | None, title: str) -> None:
        if path is None or not path.exists():
            messagebox.showwarning(title, f"{title} 파일을 찾을 수 없습니다.")
            return
        try:
            _open_file(path)
        except OSError as exc:
            messagebox.showerror(title, f"파일을 열 수 없습니다.\n{path}\n\n{exc}")

    def run_preflight_check(self) -> None:
        if self.preflight_exe is None or not self.preflight_exe.exists():
            messagebox.showwarning("실행 전 점검", "고객환경점검.exe를 찾을 수 없습니다.")
            return
        args = _build_preflight_check_args(self.preflight_exe, self.base_dir)
        self.status_var.set("실행 전 점검을 진행하는 중입니다.")
        self.update_idletasks()
        try:
            result = subprocess.run(
                args,
                cwd=self.base_dir,
                text=True,
                capture_output=True,
                check=False,
                timeout=180,
                creationflags=_creationflags(),
                encoding=locale.getpreferredencoding(False),
                errors="replace",
            )
        except subprocess.TimeoutExpired:
            self.status_var.set("실행 전 점검 시간이 초과되었습니다.")
            messagebox.showerror("실행 전 점검 실패", "점검 시간이 초과되었습니다. 프린터 연결과 파일 잠금 상태를 확인해 주세요.")
            return
        except OSError as exc:
            self.status_var.set("실행 전 점검에 실패했습니다.")
            messagebox.showerror("실행 전 점검 실패", str(exc))
            return

        report_path, _package_path = _support_package_paths(self.base_dir)
        summary = _read_preflight_summary(report_path)
        if result.returncode != 0:
            detail = summary or _process_output_text(result) or "고객환경점검.exe가 오류를 반환했습니다."
            self.status_var.set("실행 전 점검에서 오류가 발견되었습니다.")
            messagebox.showerror(
                "실행 전 점검 실패",
                f"{detail}\n\n진단 보고서: {report_path}\n도움말 > 지원 패키지 생성으로 지원 ZIP을 만든 뒤 문의하세요.",
            )
            return

        if not report_path.exists():
            self.status_var.set("실행 전 점검 보고서를 찾지 못했습니다.")
            messagebox.showerror("실행 전 점검 실패", "점검은 끝났지만 진단 보고서 파일이 없습니다.")
            return

        self.status_var.set(f"실행 전 점검 완료: {summary or report_path}")
        messagebox.showinfo(
            "실행 전 점검 완료",
            f"{summary or '점검이 완료되었습니다.'}\n\n"
            f"진단 보고서: {report_path}\n\n"
            "다음 단계: 프린터 설정을 확인한 뒤 1장 테스트 출력하세요.",
        )

    def create_support_package(self) -> None:
        if self.preflight_exe is None or not self.preflight_exe.exists():
            messagebox.showwarning("지원 패키지", "고객환경점검.exe를 찾을 수 없습니다.")
            return
        args = _build_support_package_args(self.preflight_exe, self.base_dir)
        self.status_var.set("지원 패키지를 생성하는 중입니다.")
        self.update_idletasks()
        try:
            result = subprocess.run(
                args,
                cwd=self.base_dir,
                text=True,
                capture_output=True,
                check=False,
                timeout=180,
                creationflags=_creationflags(),
                encoding=locale.getpreferredencoding(False),
                errors="replace",
            )
        except subprocess.TimeoutExpired:
            self.status_var.set("지원 패키지 생성 시간이 초과되었습니다.")
            messagebox.showerror("지원 패키지 생성 실패", "점검 시간이 초과되었습니다. 프린터 연결과 파일 잠금 상태를 확인해 주세요.")
            return
        except OSError as exc:
            self.status_var.set("지원 패키지 생성에 실패했습니다.")
            messagebox.showerror("지원 패키지 생성 실패", str(exc))
            return
        if result.returncode != 0:
            detail = _process_output_text(result)
            self.status_var.set("지원 패키지 생성에 실패했습니다.")
            messagebox.showerror("지원 패키지 생성 실패", detail or "고객환경점검.exe가 오류를 반환했습니다.")
            return

        report_path, package_path = _support_package_paths(self.base_dir)
        missing = [path.name for path in (report_path, package_path) if not path.exists()]
        if missing:
            self.status_var.set("지원 패키지 결과 파일을 찾지 못했습니다.")
            messagebox.showerror("지원 패키지 생성 실패", "점검은 끝났지만 결과 파일이 없습니다.\n" + "\n".join(missing))
            return

        self.status_var.set(f"지원 패키지 생성 완료: {package_path}")
        messagebox.showinfo(
            "지원 패키지 생성",
            "지원 패키지를 생성했습니다.\n\n"
            f"진단 보고서: {report_path}\n"
            f"지원 ZIP: {package_path}\n\n"
            "문의 시 지원 ZIP 파일을 전달하세요.",
        )

    def open_path(self, path: Path) -> None:
        path.mkdir(parents=True, exist_ok=True)
        subprocess.Popen(["explorer.exe", str(path)])

    def select_all_labels(self) -> None:
        self.selected_label_indexes = set(range(len(self.label_rows)))
        self.refresh_tables()
        self.status_var.set(f"전체 {len(self.selected_label_indexes)}건을 선택했습니다.")

    def clear_label_selection(self) -> None:
        self.selected_label_indexes.clear()
        self.refresh_tables()
        self.status_var.set("선택을 해제했습니다.")

    def _handle_label_click(self, event: tk.Event) -> str | None:
        if self.labels_tree.identify_column(event.x) != "#1":
            return None
        item = self.labels_tree.identify_row(event.y)
        if not item:
            return None
        self._toggle_label_index(self.labels_tree.index(item))
        return "break"

    def _toggle_focused_label(self, _event: tk.Event) -> str:
        item = self.labels_tree.focus()
        if item:
            self._toggle_label_index(self.labels_tree.index(item))
        return "break"

    def _toggle_label_index(self, index: int) -> None:
        if index in self.selected_label_indexes:
            self.selected_label_indexes.remove(index)
        else:
            self.selected_label_indexes.add(index)
        self.refresh_tables()
        self._select_index(self.labels_tree, index)
        self.status_var.set(f"선택 {len(self.selected_label_indexes)}건")

    def _sync_labels_from_db(self, preserve_selection: bool = True) -> None:
        selected_barcodes = self._selected_label_barcodes() if preserve_selection else set()
        self.label_rows = _dynamic_label_rows_from_db_rows(self.db_rows, self.label_headers)
        if preserve_selection:
            self.selected_label_indexes = {
                index
                for index, row in enumerate(self.label_rows)
                if _row_field_value(row, "barcode") in selected_barcodes
            }
        else:
            self.selected_label_indexes.clear()

    def _selected_label_barcodes(self) -> set[str]:
        return {
            _row_field_value(self.label_rows[index], "barcode")
            for index in self.selected_label_indexes
            if 0 <= index < len(self.label_rows) and _row_field_value(self.label_rows[index], "barcode")
        }

    def _selected_print_rows(self) -> list[dict[str, str]]:
        label_headers = self.__dict__.get("label_headers", tuple(LABEL_HEADERS))
        return [
            row
            for index, row in enumerate(self.label_rows)
            if index in self.selected_label_indexes and any(str(row.get(header, "")).strip() for header in label_headers)
        ]

    def _active_tree_and_rows(self) -> tuple[ttk.Treeview, list[dict[str, str]]]:
        if self.tabs.index(self.tabs.select()) == 0:
            return self.labels_tree, self.label_rows
        return self.db_tree, self.db_rows

    def _selected_label_index(self) -> int | None:
        selected = self.labels_tree.selection()
        return self.labels_tree.index(selected[0]) if selected else None

    def _selected_label_index_for_scan(self) -> int | None:
        if self.tabs.index(self.tabs.select()) != 0:
            return None
        return self._selected_label_index()

    def _first_empty_label_index(self) -> int:
        for index, row in enumerate(self.label_rows):
            if not any(str(row.get(header, "")).strip() for header in LABEL_HEADERS):
                return index
        self.label_rows.append({header: "" for header in LABEL_HEADERS})
        return len(self.label_rows) - 1

    def _select_last(self, tree: ttk.Treeview) -> None:
        children = tree.get_children()
        if children:
            tree.selection_set(children[-1])
            tree.see(children[-1])

    def _select_index(self, tree: ttk.Treeview, index: int) -> None:
        children = tree.get_children()
        if 0 <= index < len(children):
            tree.selection_set(children[index])
            if callable(getattr(tree, "focus", None)):
                tree.focus(children[index])
            tree.see(children[index])

    def _see_index(self, tree: ttk.Treeview, index: int) -> None:
        children = tree.get_children()
        if 0 <= index < len(children):
            tree.see(children[index])


def _label_rows_from_db_rows(db_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    return _dynamic_label_rows_from_db_rows(db_rows, LABEL_HEADERS)


def _dynamic_label_rows_from_db_rows(db_rows: list[dict[str, str]], label_headers: tuple[str, ...]) -> list[dict[str, str]]:
    return [
        _label_row_from_db_row(row, label_headers)
        for row in db_rows
        if any(str(row.get(header, "")).strip() for header in label_headers)
    ]


def _label_row_from_db_row(row: dict[str, str], label_headers: tuple[str, ...]) -> dict[str, str]:
    label_row = {header: str(row.get(header, "")).strip() for header in label_headers}
    for header in label_headers:
        if _field_for_header(header) == "print_qty" and not label_row[header]:
            label_row[header] = DEFAULT_PRINT_QTY
    return label_row


def _ensure_label_row_index(label_rows: list[dict[str, str]], row_index: int | None, label_headers: tuple[str, ...] = LABEL_HEADERS) -> tuple[int, bool]:
    if row_index is None or row_index < 0:
        row_index = len(label_rows)
    appended = False
    while row_index >= len(label_rows):
        label_rows.append({header: "" for header in label_headers})
        appended = True
    return row_index, appended


def _delete_label_rows(label_rows: list[dict[str, str]], selected_indexes: set[int], focused_index: int | None) -> int:
    if selected_indexes:
        delete_indexes = {index for index in selected_indexes if 0 <= index < len(label_rows)}
    elif focused_index is not None and 0 <= focused_index < len(label_rows):
        delete_indexes = {focused_index}
    else:
        delete_indexes = set()
    if not delete_indexes:
        return 0
    label_rows[:] = [row for index, row in enumerate(label_rows) if index not in delete_indexes]
    selected_indexes.clear()
    return len(delete_indexes)


def _fill_tree(
    tree: ttk.Treeview,
    headers: tuple[str, ...],
    rows: list[dict[str, str]],
    selected_indexes: set[int] | None = None,
) -> None:
    selected_indexes = selected_indexes or set()
    for item in tree.get_children():
        tree.delete(item)
    for index, row in enumerate(rows):
        values = []
        for header in headers:
            if header == SELECT_HEADER:
                values.append(CHECKED if index in selected_indexes else UNCHECKED)
            else:
                values.append(row.get(header, ""))
        tags = ("checked",) if index in selected_indexes else ()
        tree.insert("", "end", values=values, tags=tags)


def _first_existing(*paths: Path) -> Path | None:
    for path in paths:
        if path.exists():
            return path
    return None


def _creationflags() -> int:
    return subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0


def _open_file(path: Path) -> None:
    if sys.platform == "win32":
        os.startfile(str(path))  # type: ignore[attr-defined]
        return
    subprocess.Popen(["xdg-open", str(path)])


def _build_preflight_check_args(preflight_exe: Path, base_dir: Path) -> list[str]:
    return [str(preflight_exe), "--base-dir", str(base_dir), "--no-support-package"]


def _build_support_package_args(preflight_exe: Path, base_dir: Path) -> list[str]:
    return [str(preflight_exe), "--base-dir", str(base_dir)]


def _support_package_paths(base_dir: Path) -> tuple[Path, Path]:
    out_dir = base_dir / "out"
    return out_dir / SUPPORT_REPORT_FILE, out_dir / SUPPORT_PACKAGE_FILE


def _read_preflight_summary(report_path: Path) -> str:
    if not report_path.exists():
        return ""
    try:
        lines = report_path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    except OSError:
        return ""
    for line in lines:
        if line.startswith("점검 결과:"):
            return line
    return ""


def _process_output_text(result: subprocess.CompletedProcess[str]) -> str:
    text = (result.stderr or result.stdout or "").strip()
    if len(text) > 1200:
        return text[:1200].rstrip() + "\n..."
    return text


def _ask_unknown_resolution(parent: tk.Misc, item_index: int) -> str | None:
    dialog = tk.Toplevel(parent)
    dialog.title("출력 상태 확인")
    dialog.transient(parent)
    dialog.grab_set()
    dialog.resizable(False, False)
    frame = ttk.Frame(dialog, padding=(20, 18))
    frame.grid(row=0, column=0, sticky="nsew")
    ttk.Label(frame, text=f"{item_index}번 항목의 실제 출력 여부를 확인하세요.", style="Title.TLabel").grid(
        row=0, column=0, columnspan=3, sticky="w"
    )
    ttk.Label(
        frame,
        text="RAW 전송 오류만으로 물리 라벨 출력 여부를 알 수 없습니다. 프린터와 라벨을 직접 확인하세요.",
        style="Hint.TLabel",
        wraplength=480,
    ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(8, 18))
    result: dict[str, str | None] = {"value": None}

    def finish(value: str | None) -> None:
        result["value"] = value
        dialog.destroy()

    ttk.Button(frame, text="출력됨", command=lambda: finish("sent"), style="Primary.TButton").grid(
        row=2, column=0, padx=(0, 8)
    )
    ttk.Button(frame, text="출력 안 됨", command=lambda: finish("pending"), style="Secondary.TButton").grid(
        row=2, column=1, padx=(0, 8)
    )
    ttk.Button(frame, text="취소", command=lambda: finish(None), style="Secondary.TButton").grid(row=2, column=2)
    dialog.protocol("WM_DELETE_WINDOW", lambda: finish(None))
    dialog.bind("<Escape>", lambda _event: finish(None))
    parent.wait_window(dialog)
    return result["value"]


def _selected_print_queue_path(base_dir: Path) -> Path:
    return base_dir / "out" / "selected_print_queue.xlsx"


def _load_dynamic_db_rows(path: Path) -> tuple[tuple[str, ...], list[dict[str, str]]]:
    if not path.exists():
        # load_db_rows creates the built-in customer DB. Keep its visible
        # sales columns instead of exposing queue-only lot/quantity columns
        # during the very first launch.
        rows, headers = load_db_source(path)
        return headers, rows
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook["BarcodeDB"] if "BarcodeDB" in workbook.sheetnames else workbook.active
        values = iter(sheet.iter_rows())
        header_cells = next(values, None)
        if header_cells is None:
            return tuple(DB_HEADERS), []
        headers = _normalize_dynamic_headers(tuple(cell.value for cell in header_cells))
        rows: list[dict[str, str]] = []
        for row_number, source_row in enumerate(values, start=2):
            row = {
                header: _db_cell_text(source_row[index], row_number, header)
                if index < len(source_row) else ""
                for index, header in enumerate(headers)
            }
            if any(row.values()):
                rows.append(row)
        return headers, rows
    finally:
        workbook.close()


def _load_active_db_state(
    path: Path,
) -> tuple[tuple[str, ...], list[dict[str, str]], tuple[str, ...], list[dict[str, str]]]:
    """Build one complete print state from one DB so old DB data cannot leak."""
    db_headers, db_rows = _load_dynamic_db_rows(path)
    label_headers = _label_headers_from_db_headers(db_headers)
    label_rows = _dynamic_label_rows_from_db_rows(db_rows, label_headers)
    return db_headers, db_rows, label_headers, label_rows


def _save_dynamic_db_rows(path: Path, headers: tuple[str, ...], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BarcodeDB"
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)


def _normalize_dynamic_headers(raw_headers: tuple[object, ...]) -> tuple[str, ...]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, raw_header in enumerate(raw_headers, start=1):
        header = _cell_to_text(raw_header) or f"컬럼{index}"
        base_header = header
        occurrence = seen.get(base_header, 0) + 1
        seen[base_header] = occurrence
        if occurrence > 1:
            header = f"{base_header}_{occurrence}"
        headers.append(header)
    return tuple(headers) or tuple(DB_HEADERS)


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _label_headers_from_db_headers(db_headers: tuple[str, ...]) -> tuple[str, ...]:
    headers = tuple(header for header in db_headers if header != SELECT_HEADER)
    return headers or tuple(LABEL_HEADERS)


def _field_for_header(header: str) -> str | None:
    normalized = _normalize_header_key(header)
    for field, aliases in FIELD_ALIASES.items():
        if normalized in {_normalize_header_key(alias) for alias in aliases}:
            return field
    return None


def _normalize_header_key(value: str) -> str:
    return "".join(ch for ch in str(value).strip().lower() if ch.isalnum())


def _row_field_value(row: dict[str, str], field: str) -> str:
    for header, value in row.items():
        if _field_for_header(header) == field:
            return str(value).strip()
    return str(row.get(field, "")).strip()


def _first_non_barcode_value(row: dict[str, str], headers: tuple[str, ...]) -> str:
    for header in headers:
        if _field_for_header(header) == "barcode":
            continue
        value = str(row.get(header, "")).strip()
        if value:
            return value
    return ""


def _lookup_dynamic_barcode(rows: list[dict[str, str]], barcode: str) -> dict[str, str] | None:
    target = barcode.strip()
    for row in rows:
        if _row_field_value(row, "barcode") == target:
            return row
    return None


def _find_row_by_query(rows: list[dict[str, str]], query: str) -> tuple[int, dict[str, str]] | None:
    matches = _matching_rows_by_query(rows, query)
    return matches[0] if matches else None


def _matching_rows_by_query(rows: list[dict[str, str]], query: str) -> list[tuple[int, dict[str, str]]]:
    target = _normalize_search_text(query)
    if not target:
        return []
    exact_barcode_matches = [
        (index, row)
        for index, row in enumerate(rows)
        if _normalize_search_text(_row_field_value(row, "barcode")) == target
    ]
    if exact_barcode_matches:
        return exact_barcode_matches
    return [
        (index, row)
        for index, row in enumerate(rows)
        if _row_matches_query(row, target)
    ]


def _row_matches_query(row: dict[str, str], normalized_query: str) -> bool:
    return any(normalized_query in _normalize_search_text(value) for value in row.values())


def _normalize_search_text(value: object) -> str:
    return str(value).strip().casefold()


def _printable_rows_from_label_rows(rows: list[dict[str, str]], label_headers: tuple[str, ...] = LABEL_HEADERS) -> list[dict[str, str]]:
    printable_rows: list[dict[str, str]] = []
    for row in rows:
        if not any(
            str(row.get(header, "")).strip()
            for header in label_headers
            if _field_for_header(header) != "print_qty"
        ):
            continue
        barcode = _row_field_value(row, "barcode")
        printable_row = {
            "item_code": _row_field_value(row, "item_code"),
            "item_name": _row_field_value(row, "item_name"),
            "barcode": barcode,
            "lot_no": _row_field_value(row, "lot_no"),
            "qty": _row_field_value(row, "qty") or "1",
            "print_qty": _row_field_value(row, "print_qty") or DEFAULT_PRINT_QTY,
        }
        for header in label_headers:
            value = str(row.get(header, "")).strip()
            if not value:
                continue
            field = _field_for_header(header)
            if header in printable_row and field in {"item_code", "item_name", "lot_no", "qty"}:
                printable_row[COLUMN_LABELS[field]] = value
            elif header not in printable_row:
                printable_row[header] = value
        if any(printable_row.values()):
            printable_rows.append(printable_row)
    return printable_rows


def _rows_with_print_quantity(rows: list[dict[str, str]], print_quantity: int | None) -> list[dict[str, str]]:
    job_rows = [dict(row) for row in rows]
    if print_quantity is None:
        return job_rows
    safe_quantity = max(1, min(100, int(print_quantity)))
    for row in job_rows:
        row["print_qty"] = str(safe_quantity)
    return job_rows


def _print_readiness_errors(
    rows: list[dict[str, str]],
    label_headers: tuple[str, ...] = LABEL_HEADERS,
    config_path: Path | None = None,
    *,
    print_quantity: int | None = None,
) -> list[str]:
    printable_rows = _rows_with_print_quantity(_printable_rows_from_label_rows(rows, label_headers), print_quantity)
    errors: list[str] = []

    if not printable_rows:
        errors.append("인쇄할 유효한 행이 없습니다. 바코드 또는 품목 정보를 입력하세요.")
        return errors

    for index, row in enumerate(printable_rows, start=1):
        barcode = str(row.get("barcode", "")).strip()
        if not barcode:
            errors.append(f"{index}행 바코드가 비어 있습니다.")
        raw_quantity = str(row.get("print_qty", "") or DEFAULT_PRINT_QTY).strip()
        try:
            quantity = int(float(raw_quantity))
        except ValueError:
            errors.append(f"{index}행 출력 매수는 숫자로 입력하세요.")
            continue
        if quantity < 1 or quantity > 1000:
            errors.append(f"{index}행 출력 매수는 1부터 1000 사이로 입력하세요.")

    if config_path is None:
        return errors
    if not config_path.exists():
        errors.append("프린터 설정 파일(config.ini)을 찾을 수 없습니다. 프린터설정.exe에서 먼저 저장하세요.")
        return errors

    try:
        settings = load_printer_settings(config_path)
    except Exception as exc:
        errors.append(f"프린터 설정을 읽을 수 없습니다: {exc}")
        return errors

    report = validate_printer_settings(settings)
    errors.extend(f"프린터 설정: {error}" for error in report.errors)
    return errors


def _default_print_quantity(rows: list[dict[str, str]]) -> int:
    for row in rows:
        raw_value = _row_field_value(row, "print_qty")
        if not raw_value:
            continue
        try:
            quantity = int(float(raw_value))
        except ValueError:
            continue
        if quantity > 0:
            return max(1, min(100, quantity))
    return 1


def _build_print_args(print_exe: Path, config_path: Path, action: str, excel_path: Path) -> list[str]:
    args = [str(print_exe), "--config", str(config_path), "--excel", str(excel_path), action]
    if action == "--print":
        args.append("--yes")
    return args


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Label print management app.")
    parser.add_argument("--base-dir", default=None, help="Folder containing config.ini and 라벨출력엔진.exe")
    parser.add_argument("--smoke-test", action="store_true", help="Load data files and exit")
    parser.add_argument("--ui-smoke-test", action="store_true", help="Create the Tk UI once and exit")
    args = parser.parse_args(argv)
    base_dir = Path(args.base_dir).resolve() if args.base_dir else app_base_dir()
    if args.ui_smoke_test:
        app = LabelManagerApp(base_dir)
        app.update_idletasks()
        app.destroy()
        return 0
    if args.smoke_test:
        _load_dynamic_db_rows(base_dir / "barcode_db.xlsx")
        load_label_rows(base_dir / "print_queue.xlsx")
        return 0
    app = LabelManagerApp(base_dir)
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

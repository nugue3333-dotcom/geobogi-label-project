from __future__ import annotations

import argparse
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .customer_backup import (
    BACKUP_DIRS,
    BACKUP_FILES,
    MANIFEST_FILE,
    CustomerBackupError,
    RestoreResult,
    create_customer_backup as _create_customer_backup,
    restore_customer_backup as _restore_customer_backup,
)
from .data_store import LABEL_HEADERS
from .runtime_paths import executable_dir, runtime_base_dir


MODE_CHOICES = ("DryRun", "Print", "OpenOutputFolder", "OpenLastRunLog", "Backup", "Restore")
BRAND_TITLE = "채움랩 라벨 출력"
RESTORE_ALLOWED_FILES = frozenset(BACKUP_FILES)
RESTORE_ALLOWED_DIR_PREFIXES = ("templates/", "db/", "assets/images/")
RESTORE_MANIFEST_FILE = MANIFEST_FILE


class JobRunnerError(RuntimeError):
    """Customer-facing execution failure with a message safe to show."""


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run customer label output jobs without PowerShell.")
    parser.add_argument("mode", choices=MODE_CHOICES)
    parser.add_argument("-Quiet", "--quiet", action="store_true", help="Print messages to stdout instead of popups")
    parser.add_argument("--install-dir", default=None, help=argparse.SUPPRESS)
    parser.add_argument("--backup-zip", default=None, help="Customer backup ZIP to restore")
    args = parser.parse_args(argv)

    try:
        install_dir = Path(args.install_dir).resolve() if args.install_dir else executable_dir()
        run_job(args.mode, install_dir=install_dir, quiet=args.quiet, backup_zip=args.backup_zip)
        return 0
    except Exception as exc:
        message = str(exc) or exc.__class__.__name__
        _show_message(message, title=f"{BRAND_TITLE} 오류", quiet=args.quiet, error=True)
        return 1


def run_job(mode: str, *, install_dir: Path, quiet: bool = False, backup_zip: str | Path | None = None) -> None:
    install_dir = install_dir.resolve()
    base_dir = runtime_base_dir(install_dir)
    if mode == "DryRun":
        _invoke_engine(install_dir, base_dir, "--dry-run")
        _show_message(
            f"출력 전 검증 완료\n\n{base_dir / 'print_queue.xlsx'}\n{base_dir / 'out'}\n{base_dir / 'last_run.log'}",
            title=BRAND_TITLE,
            quiet=quiet,
        )
        return
    if mode == "Print":
        if not _confirm_print_send(quiet=quiet):
            _show_message("실제 출력 전송을 취소했습니다.", title=BRAND_TITLE, quiet=quiet)
            return
        _invoke_engine(install_dir, base_dir, "--print --yes")
        _show_message(
            "프린터 전송 완료\n\n프린터 실제 출력 여부는 장비 상태와 라벨 배출을 확인하세요.",
            title="프린터 전송 완료",
            quiet=quiet,
        )
        return
    if mode == "OpenOutputFolder":
        output_dir = base_dir / "out"
        output_dir.mkdir(parents=True, exist_ok=True)
        _open_path(output_dir)
        return
    if mode == "OpenLastRunLog":
        log_path = base_dir / "last_run.log"
        if not log_path.exists():
            _show_message("last_run.log 파일이 아직 없습니다.", title=BRAND_TITLE, quiet=quiet)
            return
        _open_path(log_path)
        return
    if mode == "Backup":
        backup_path = create_customer_backup(base_dir)
        _show_message(
            f"고객 데이터 백업 완료\n\n{backup_path}\n\nPC 교체나 재설치 전에 이 ZIP 파일을 보관하세요.",
            title=BRAND_TITLE,
            quiet=quiet,
        )
        return
    if mode == "Restore":
        if backup_zip is None:
            backup_zip = _ask_backup_zip_path()
        result = restore_customer_backup(base_dir, Path(backup_zip))
        pre_restore = f"\n복원 전 기존 데이터 백업: {result.pre_restore_backup}" if result.pre_restore_backup else ""
        _show_message(
            f"고객 데이터 복원 완료\n\n복원 파일 수: {result.restored_files}{pre_restore}",
            title=BRAND_TITLE,
            quiet=quiet,
        )
        return
    raise JobRunnerError(f"지원하지 않는 실행 모드입니다: {mode}")


def export_print_queue(base_dir: Path) -> Path:
    source_path = _labels_workbook_path(base_dir)
    rows = _read_label_source_rows(source_path)
    if not rows:
        raise JobRunnerError("출력할 데이터 행을 찾지 못했습니다.")

    queue_path = base_dir / "print_queue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Labels"
    sheet.append(list(LABEL_HEADERS))
    for row in rows:
        sheet.append([row[header] for header in LABEL_HEADERS])
    for column_index, header in enumerate(LABEL_HEADERS, start=1):
        sheet.column_dimensions[sheet.cell(1, column_index).column_letter].width = 28 if header in {"item_name", "barcode"} else 16
    queue_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(queue_path)
    return queue_path


def create_customer_backup(base_dir: Path) -> Path:
    try:
        return _create_customer_backup(base_dir)
    except CustomerBackupError as exc:
        raise JobRunnerError(str(exc)) from exc


def restore_customer_backup(base_dir: Path, backup_path: Path) -> RestoreResult:
    try:
        return _restore_customer_backup(base_dir, backup_path)
    except CustomerBackupError as exc:
        raise JobRunnerError(str(exc)) from exc


def _invoke_engine(install_dir: Path, base_dir: Path, action: str) -> None:
    print_exe = install_dir / "라벨출력엔진.exe"
    config_path = base_dir / "config.ini"
    log_path = base_dir / "last_run.log"
    if not print_exe.is_file():
        raise JobRunnerError(f"라벨출력엔진.exe 파일이 없습니다.\n{print_exe}")
    if not config_path.is_file():
        raise JobRunnerError(f"config.ini 파일이 없습니다.\n{config_path}")

    export_print_queue(base_dir)

    command = [str(print_exe), "--config", str(config_path), *action.split()]
    completed = subprocess.run(
        command,
        cwd=str(base_dir),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    log_path.write_text((completed.stdout or "") + (completed.stderr or ""), encoding="utf-8")
    if completed.returncode != 0:
        raise JobRunnerError(f"실행 실패입니다. last_run.log 파일을 확인하세요.\n{log_path}")


def _ask_backup_zip_path() -> str:
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        selected = filedialog.askopenfilename(
            title="복원할 고객 데이터 백업 ZIP 선택",
            filetypes=(("채움랩 백업 ZIP", "chaeumlab_customer_backup_*.zip"), ("ZIP 파일", "*.zip"), ("모든 파일", "*.*")),
        )
        root.destroy()
        if selected:
            return selected
    except Exception:
        pass
    raise JobRunnerError("복원할 백업 ZIP 파일을 선택하지 않았습니다.")


def _confirm_print_send(*, quiet: bool) -> bool:
    message = (
        "현재 인쇄 데이터가 실제 프린터로 전송됩니다.\n\n"
        "라벨 용지, 프린터 전원, 연결 상태, 라벨 크기를 확인한 뒤 진행하세요."
    )
    if quiet:
        raise JobRunnerError("Quiet 모드에서는 실제 출력 전송을 실행할 수 없습니다. 화면 확인 후 다시 실행하세요.")
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        confirmed = messagebox.askyesno("실제 출력 전송 확인", message, parent=root)
        root.destroy()
    except Exception as exc:
        raise JobRunnerError(f"실제 출력 전송 확인 창을 열 수 없습니다: {exc}") from exc
    return bool(confirmed)


def _labels_workbook_path(base_dir: Path) -> Path:
    for name in ("labels.xlsm", "labels.xlsx"):
        path = base_dir / name
        if path.is_file():
            return path
    raise JobRunnerError(f"labels.xlsm 또는 labels.xlsx 파일을 찾을 수 없습니다.\n{base_dir}")


def _read_label_source_rows(source_path: Path) -> list[dict[str, str]]:
    try:
        workbook = load_workbook(source_path, data_only=True, keep_vba=False)
    except Exception as exc:
        raise JobRunnerError(f"엑셀 입력 파일을 열 수 없습니다: {exc}") from exc
    try:
        sheet = workbook.active
        header_map = _header_map(sheet)
        rows: list[dict[str, str]] = []
        for row_number in range(2, sheet.max_row + 1):
            values = {header: _cell_text(sheet.cell(row_number, header_map[header]).value) for header in LABEL_HEADERS}
            if not any(values.values()):
                continue
            if not values["barcode"]:
                continue
            if not values["print_qty"]:
                values["print_qty"] = "1"
            rows.append(values)
        return rows
    finally:
        workbook.close()


def _header_map(sheet: object) -> dict[str, int]:
    found: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):  # type: ignore[index]
        value = _cell_text(cell.value)
        if value:
            found[value] = column_index
    missing = [header for header in LABEL_HEADERS if header not in found]
    if missing:
        raise JobRunnerError(f"필수 컬럼이 없습니다.\n\n- " + "\n- ".join(missing))
    return {header: found[header] for header in LABEL_HEADERS}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _open_path(path: Path) -> None:
    try:
        os.startfile(str(path))  # type: ignore[attr-defined]
    except AttributeError:
        subprocess.Popen(["xdg-open", str(path)])
    except OSError as exc:
        raise JobRunnerError(f"파일을 열 수 없습니다.\n{path}\n\n{exc}") from exc


def _show_message(text: str, *, title: str, quiet: bool, error: bool = False) -> None:
    if quiet:
        stream = sys.stderr if error else sys.stdout
        print(text, file=stream)
        return
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        if error:
            messagebox.showerror(title, text, parent=root)
        else:
            messagebox.showinfo(title, text, parent=root)
        root.destroy()
    except Exception:
        stream = sys.stderr if error else sys.stdout
        print(f"{title}: {text}", file=stream)


if __name__ == "__main__":
    raise SystemExit(main())

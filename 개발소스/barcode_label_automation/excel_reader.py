from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import re

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("item_code", "item_name", "barcode", "lot_no", "qty", "print_qty")
PRICE_HEADERS = ("price", "판매가", "가격", "판매 가격", "금액", "단가")
DISPLAY_FIELD_LABELS = {
    "item_code": "상품코드",
    "item_name": "상품명",
    "lot_no": "LOT",
    "qty": "수량",
}


@dataclass(frozen=True)
class LabelRow:
    item_code: str
    item_name: str
    barcode: str
    lot_no: str
    qty: int
    print_qty: int
    source_fields: tuple[tuple[str, str], ...] = ()


def read_labels(path: str | Path) -> list[LabelRow]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Excel file not found: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = _read_headers(sheet)
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    labels: list[LabelRow] = []
    for row_number in range(2, sheet.max_row + 1):
        values = {column: sheet.cell(row=row_number, column=index).value for column, index in headers.items()}
        if all(value is None or str(value).strip() == "" for value in values.values()):
            continue
        labels.append(
            LabelRow(
                item_code=_as_text(values["item_code"]),
                item_name=_as_text(values["item_name"]),
                barcode=_as_text(values["barcode"]),
                lot_no=_as_text(values["lot_no"]),
                qty=_as_positive_int(values["qty"], "qty", row_number),
                print_qty=_as_required_int(values["print_qty"], "print_qty", row_number),
                source_fields=_source_fields(headers, values),
            )
        )
    return labels


def _read_headers(sheet: Any) -> dict[str, int]:
    headers: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):
        if cell.value is None:
            continue
        header = str(cell.value).strip()
        if header:
            headers[header] = column_index
    return headers


def _source_fields(headers: dict[str, int], values: dict[str, object]) -> tuple[tuple[str, str], ...]:
    extra_headers = [header for header in headers if header not in REQUIRED_COLUMNS]
    # The manager adds the connected DB's visible columns as queue extras.
    # Preserve that exact order instead of collapsing every DB to a fixed
    # product/price/code layout.
    represented_fields = {_field_for_header(header) for header in extra_headers}
    if extra_headers:
        source_headers: list[str] = []
        if "item_name" not in represented_fields:
            source_headers.append("item_name")
        source_headers.extend(extra_headers)
        if "item_code" not in represented_fields:
            source_headers.append("item_code")
    else:
        source_headers = ["item_name", "item_code", "lot_no", "qty"]
    fields: list[tuple[str, str]] = []
    for header in source_headers:
        value = _as_text(values.get(header))
        if not value:
            continue
        if _field_for_header(header) in {"barcode", "print_qty"}:
            continue
        if _is_price_header(header):
            value = _format_korean_won(value)
        label = DISPLAY_FIELD_LABELS.get(header, header)
        fields.append((label, value))
    return tuple(fields)


def _field_for_header(header: str) -> str | None:
    normalized = _normalize_header_key(header)
    aliases = {
        "barcode": ("barcode", "bar_code", "바코드", "상품바코드", "제품바코드"),
        "item_code": ("item_code", "itemcode", "품목코드", "상품코드", "제품코드", "코드", "품번"),
        "item_name": ("item_name", "itemname", "품목명", "품명", "상품명", "제품명", "이름", "명칭"),
        "lot_no": ("lot_no", "lot", "lot번호", "lot 번호", "로트", "로트번호", "로트 번호", "lotno"),
        "qty": ("qty", "quantity", "수량", "입수", "개수"),
        "print_qty": ("print_qty", "print quantity", "출력매수", "인쇄매수", "발행매수"),
    }
    for field, names in aliases.items():
        if normalized in {_normalize_header_key(name) for name in names}:
            return field
    return None


def _is_price_header(header: str) -> bool:
    normalized = _normalize_header_key(header)
    return normalized in {_normalize_header_key(name) for name in PRICE_HEADERS}


def _normalize_header_key(value: str) -> str:
    return "".join(ch for ch in str(value).strip().lower() if ch.isalnum())


def _format_korean_won(value: str) -> str:
    """Use a Korean currency suffix instead of the fragile won glyph."""
    cleaned = value.strip().replace("₩", "").strip()
    if cleaned.endswith("원"):
        return cleaned
    if re.fullmatch(r"[+-]?\d[\d,]*", cleaned):
        return f"{int(cleaned.replace(',', '')):,}원"
    return value


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_positive_int(value: object, column: str, row_number: int) -> int:
    number = _as_required_int(value, column, row_number)
    if number < 1:
        raise ValueError(f"Row {row_number}: {column} must be greater than 0")
    return number


def _as_required_int(value: object, column: str, row_number: int) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Row {row_number}: {column} is required")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Row {row_number}: {column} must be an integer") from exc
    return number

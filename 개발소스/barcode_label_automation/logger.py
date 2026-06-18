from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .excel_reader import LabelRow


LOG_HEADERS = (
    "item_code",
    "barcode",
    "lot_no",
    "qty",
    "print_qty",
    "language",
    "mode",
    "status",
    "output_file",
    "error_message",
)


def append_print_log(output_dir: str | Path, entries: list[dict[str, object]]) -> Path:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    log_path = output_path / "print_log.xlsx"

    if log_path.exists():
        workbook = load_workbook(log_path)
        sheet = workbook.active
        _ensure_error_message_header(sheet)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "print_log"
        sheet.append(LOG_HEADERS)

    for entry in entries:
        row: LabelRow = entry["label"]  # type: ignore[assignment]
        sheet.append(
            [
                row.item_code,
                row.barcode,
                row.lot_no,
                row.qty,
                row.print_qty,
                entry.get("language"),
                entry.get("mode"),
                entry.get("status"),
                str(entry.get("output_file", "")),
                str(entry.get("error_message", "")),
            ]
        )

    workbook.save(log_path)
    return log_path


def _ensure_error_message_header(sheet: object) -> None:
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=False), None)  # type: ignore[attr-defined]
    if first_row is None:
        sheet.append(LOG_HEADERS)  # type: ignore[attr-defined]
        return
    headers = [cell.value for cell in first_row]
    if "error_message" not in headers:
        sheet.cell(row=1, column=len(headers) + 1, value="error_message")  # type: ignore[attr-defined]

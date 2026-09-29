from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .brand_assets import HEADER_LOGO_FILE
from .config import AppConfig, load_config
from .data_store import load_label_rows
from .errors import BarcodeLabelAutomationError
from .font_assets import APP_FONT_RELATIVE_PATH, APP_FONT_SHA256, register_bundled_font
from .release_manifest import validate_release_manifest
from .runtime_paths import executable_dir, runtime_base_dir


BRAND_NAME = "채움랩"
DEFAULT_PRINT_QTY = "1"
PYINSTALLER_RUNTIME_DIR = Path(r"C:\Users\Public\ChaeumLABRuntime")
PYINSTALLER_STALE_RUNTIME_SECONDS = 6 * 60 * 60
REQUIRED_EXECUTABLE_FILES = (
    "라벨디자이너.exe",
    "라벨출력관리.exe",
    "프린터설정.exe",
    "라벨작업실행기.exe",
    "라벨출력엔진.exe",
)
REQUIRED_DISTRIBUTION_FILES = (
    "release_manifest.json",
    "배포_파일목록.txt",
    "버전정보.txt",
    "시작하기.cmd",
    "처음실행_점검.cmd",
    "고객데이터_백업.cmd",
    "고객데이터_복원.cmd",
    "00_install_trusted_location.cmd",
    "register_label_filetype.cmd",
    "scripts/install_chaeumlab_font.ps1",
    APP_FONT_RELATIVE_PATH.as_posix(),
)
REQUIRED_DATA_FILES = (
    "config.ini",
    "barcode_db.xlsx",
    "print_queue.xlsx",
    "labels.xlsm",
    "templates/default_label.json",
)
REQUIRED_BRAND_ASSET_FILES = (
    HEADER_LOGO_FILE.as_posix(),
    "assets/brand/chaeumlab_logo_compact.png",
    "assets/brand/chaeumlab_app_icon.ico",
    "assets/brand/chaeumlab_app_icon_white.ico",
    "assets/brand/chaeumlab_label_file_icon_white.ico",
)
FORBIDDEN_DISTRIBUTION_FILES = (
    "label_designer.exe",
    "label_job_runner.exe",
    "label_manager.exe",
    "print_labels.exe",
    "printer_settings.exe",
    "customer_preflight.exe",
    "db/123.xlsm",
    "db/test.xlsx",
)
FORBIDDEN_DISTRIBUTION_FILE_PATTERNS = (
    "labels.before-*.xlsm",
)
FORBIDDEN_DISTRIBUTION_DIRS = (
    "tools/tools",
    "ChaeumLAB",
)
FORBIDDEN_OCR_RUNTIME_FILES = (
    "ambiguous_words.exe",
    "classifier_tester.exe",
    "cntraining.exe",
    "combine_lang_model.exe",
    "combine_tessdata.exe",
    "dawg2wordlist.exe",
    "lstmeval.exe",
    "lstmtraining.exe",
    "merge_unicharsets.exe",
    "mftraining.exe",
    "set_unicharset_properties.exe",
    "shapeclustering.exe",
    "tesseract-uninstall.exe",
    "text2image.exe",
    "unicharset_extractor.exe",
    "wordlist2dawg.exe",
    "tessdata/osd.traineddata",
)
FORBIDDEN_OCR_RUNTIME_SUFFIXES = {".html", ".jar"}


@dataclass(frozen=True)
class CheckResult:
    name: str
    ok: bool
    message: str


@dataclass(frozen=True)
class PreflightReport:
    install_dir: Path
    data_dir: Path
    results: list[CheckResult]

    @property
    def ok(self) -> bool:
        return all(result.ok for result in self.results)


def run_preflight(
    install_dir: str | Path | None = None,
    *,
    run_dry_run: bool = True,
    dry_run_timeout_sec: int = 45,
) -> PreflightReport:
    base_dir = Path(install_dir).resolve() if install_dir else executable_dir()
    data_dir = runtime_base_dir(base_dir)
    results: list[CheckResult] = []

    results.extend(_check_required_files(base_dir, data_dir))
    results.append(_check_distribution_hygiene(base_dir))
    results.append(_check_runtime_temp_hygiene())
    results.append(_check_ocr_runtime_hygiene(base_dir))
    results.append(_check_brand_assets(base_dir))
    results.append(_check_app_font(base_dir))
    config = _check_config(data_dir / "config.ini", results)
    results.extend(_check_workbooks(data_dir))
    results.append(_check_print_data_readiness(data_dir))
    results.append(_check_default_template(data_dir / "templates" / "default_label.json"))
    results.append(_check_writable_dir(data_dir))
    results.append(_check_release_manifest(base_dir))
    if run_dry_run:
        results.append(_check_dry_run(base_dir, data_dir, data_dir / "config.ini", dry_run_timeout_sec))

    if config is not None:
        results.append(_check_label_settings(config))

    return PreflightReport(install_dir=base_dir, data_dir=data_dir, results=results)


def print_report(report: PreflightReport) -> None:
    print(format_report(report))


def format_report(report: PreflightReport, *, redact_paths: bool = True) -> str:
    lines = [
        f"{BRAND_NAME} 라벨 출력 패키지 실행 전 점검",
        f"점검 시간: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"설치 폴더: {report.install_dir}",
        f"데이터 폴더: {report.data_dir}",
        "",
    ]
    for result in report.results:
        status = "OK" if result.ok else "오류"
        lines.append(f"[{status}] {result.name}: {result.message}")
    ok_count = sum(1 for result in report.results if result.ok)
    fail_count = len(report.results) - ok_count
    lines.append("")
    lines.append(f"점검 결과: 정상 {ok_count}개 / 오류 {fail_count}개")
    if report.ok:
        lines.append("바로 사용 가능합니다. 실제 출력 전에는 프린터 설정과 라벨 1장 테스트를 확인하세요.")
    else:
        lines.append("오류 항목을 수정한 뒤 다시 실행하세요.")
    text = "\n".join(lines)
    return _redact_text(text, report) if redact_paths else text


def default_report_path(report: PreflightReport) -> Path:
    return report.data_dir / "out" / "customer_preflight_report.txt"


def write_report(report: PreflightReport, path: str | Path | None = None) -> Path:
    report_path = Path(path) if path else default_report_path(report)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(format_report(report) + "\n", encoding="utf-8-sig")
    return report_path


def default_support_package_path(report: PreflightReport) -> Path:
    return report.data_dir / "out" / "customer_support_package.zip"


def write_support_package(report: PreflightReport, path: str | Path | None = None) -> Path:
    package_path = Path(path) if path else default_support_package_path(report)
    package_path.parent.mkdir(parents=True, exist_ok=True)
    contents = {
        "customer_preflight_report.txt": format_report(report) + "\n",
        "config_summary.txt": _format_config_summary(report),
        "environment_summary.txt": _format_environment_summary(report),
        "file_inventory.txt": _format_file_inventory(report),
        "release_manifest_summary.txt": _format_release_manifest_summary(report),
        "print_log_summary.txt": _format_print_log_summary(report),
        "privacy_note.txt": (
            "This support package does not include barcode_db.xlsx, print_queue.xlsx, labels.xlsm, "
            "or raw print_log.xlsx rows.\n"
            "Local Windows paths are redacted before writing package text files.\n"
        ),
    }
    with zipfile.ZipFile(package_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, text in contents.items():
            archive.writestr(name, text)
    return package_path


def _redact_text(text: str, report: PreflightReport) -> str:
    replacements: list[tuple[str, str]] = []
    install_dir = str(report.install_dir)
    data_dir = str(report.data_dir)
    if data_dir and data_dir != install_dir:
        replacements.append((data_dir, "<DATA_DIR>"))
    if install_dir:
        replacements.append((install_dir, "<INSTALL_DIR>"))
    for env_name in ("USERPROFILE", "LOCALAPPDATA", "APPDATA", "TEMP", "TMP"):
        raw = os.environ.get(env_name)
        if raw:
            replacements.append((str(Path(raw)), f"%{env_name}%"))
            replacements.append((raw, f"%{env_name}%"))
    sanitized = text
    for raw, replacement in sorted(set(replacements), key=lambda item: len(item[0]), reverse=True):
        if raw:
            sanitized = sanitized.replace(raw, replacement)
    return sanitized


def _format_config_summary(report: PreflightReport) -> str:
    try:
        config = load_config(report.data_dir / "config.ini")
    except Exception as exc:
        return f"config.ini summary unavailable: {exc}\n"
    lines = [
        "config.ini safe summary",
        f"brand={config.printer.brand}",
        f"language={config.printer.language}",
        f"mode={config.printer.mode}",
        f"print_method={config.printer.print_method}",
        f"media_handling={config.printer.media_handling}",
        f"media_type={config.label.media_type}",
        f"label={config.label.width_mm}x{config.label.height_mm}mm",
        f"dpi={config.label.dpi}",
        f"gap_mm={config.label.gap_mm}",
        f"barcode_type={config.barcode.barcode_type}",
        f"printer_ip={_mask_ip(config.printer.ip)}",
        f"printer_port={config.printer.port}",
        f"windows_printer_name={'configured' if config.printer.windows_printer_name else 'empty'}",
        "",
    ]
    return "\n".join(lines)


def _format_file_inventory(report: PreflightReport) -> str:
    rows = ["file inventory", "scope\tname\texists\tsize_bytes"]
    for file_name in REQUIRED_EXECUTABLE_FILES:
        path = report.install_dir / file_name
        rows.append(_inventory_row("executable", file_name, path))
    for file_name in REQUIRED_DISTRIBUTION_FILES:
        path = report.install_dir / file_name
        rows.append(_inventory_row("distribution", file_name, path))
    for file_name in REQUIRED_DATA_FILES:
        path = report.data_dir / file_name
        rows.append(_inventory_row("data", file_name, path))
    rows.append("")
    return "\n".join(rows)


def _format_release_manifest_summary(report: PreflightReport) -> str:
    ok, message = validate_release_manifest(report.install_dir)
    return "\n".join(
        [
            "release manifest summary",
            f"ok={str(ok).lower()}",
            f"message={message}",
            "",
        ]
    )


def _format_environment_summary(report: PreflightReport) -> str:
    same_runtime_dir = report.install_dir.resolve() == report.data_dir.resolve()
    lines = [
        "environment summary",
        f"generated_at={datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"os={platform.platform()}",
        f"machine={platform.machine()}",
        f"python_frozen={str(getattr(sys, 'frozen', False)).lower()}",
        f"install_dir={report.install_dir}",
        f"data_dir={report.data_dir}",
        f"runtime_storage={'install_folder' if same_runtime_dir else 'user_local_appdata'}",
        f"cwd={Path.cwd()}",
        "",
    ]
    return _redact_text("\n".join(lines), report)


def _inventory_row(scope: str, file_name: str, path: Path) -> str:
    exists = path.is_file()
    size = path.stat().st_size if exists else 0
    return f"{scope}\t{file_name}\t{str(exists).lower()}\t{size}"


def _format_print_log_summary(report: PreflightReport) -> str:
    path = report.data_dir / "print_log.xlsx"
    if not path.exists():
        return "print_log.xlsx not found.\n"
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except Exception as exc:
        return f"print_log.xlsx summary unavailable: {exc}\n"
    if not rows:
        return "print_log.xlsx is empty.\n"
    headers = [str(value or "").strip() for value in rows[0]]
    status_index = headers.index("status") if "status" in headers else -1
    counter: Counter[str] = Counter()
    for row in rows[1:]:
        status = str(row[status_index] or "").strip() if status_index >= 0 and status_index < len(row) else "unknown"
        counter[status or "blank"] += 1
    lines = [
        "print_log.xlsx safe summary",
        f"total_rows={max(len(rows) - 1, 0)}",
    ]
    for status, count in sorted(counter.items()):
        lines.append(f"status.{status}={count}")
    lines.append("")
    return "\n".join(lines)


def _mask_ip(value: str) -> str:
    parts = value.strip().split(".")
    if len(parts) == 4 and all(part.isdigit() and 0 <= int(part) <= 255 for part in parts):
        return ".".join(parts[:3] + ["xxx"])
    return "configured" if value.strip() else "empty"


def _check_required_files(install_dir: Path, data_dir: Path) -> list[CheckResult]:
    results: list[CheckResult] = []
    for file_name in REQUIRED_EXECUTABLE_FILES:
        path = install_dir / file_name
        results.append(_file_result(f"실행 파일 {file_name}", path))
    for file_name in REQUIRED_DISTRIBUTION_FILES:
        path = install_dir / file_name
        results.append(_file_result(f"배포 파일 {file_name}", path))
    for file_name in REQUIRED_DATA_FILES:
        path = data_dir / file_name
        results.append(_file_result(f"데이터 파일 {file_name}", path))
    return results


def _check_distribution_hygiene(install_dir: Path) -> CheckResult:
    problems: list[str] = []
    for relative in FORBIDDEN_DISTRIBUTION_FILES:
        if (install_dir / relative).exists():
            problems.append(relative)
    for pattern in FORBIDDEN_DISTRIBUTION_FILE_PATTERNS:
        problems.extend(path.name for path in sorted(install_dir.glob(pattern))[:3] if path.is_file())
    for relative in FORBIDDEN_DISTRIBUTION_DIRS:
        if (install_dir / relative).exists():
            problems.append(relative)
    active_meipass = _active_pyinstaller_temp_dir(install_dir)
    mei_dirs = sorted(
        path.name
        for path in install_dir.glob("_MEI*")
        if path.is_dir() and path.resolve() != active_meipass
    )
    problems.extend(mei_dirs[:3])
    if problems:
        extra = "" if len(problems) <= 5 else f" 외 {len(problems) - 5}개"
        return CheckResult("배포 폴더 정리", False, "불필요한 배포/임시 파일이 남아 있습니다: " + ", ".join(problems[:5]) + extra)
    return CheckResult("배포 폴더 정리", True, "중복 실행파일과 임시 폴더 없음")


def _check_runtime_temp_hygiene(
    runtime_dir: Path = PYINSTALLER_RUNTIME_DIR,
    *,
    max_age_seconds: int = PYINSTALLER_STALE_RUNTIME_SECONDS,
    now: float | None = None,
    allow_cleanup: bool | None = None,
) -> CheckResult:
    if allow_cleanup is None:
        allow_cleanup = bool(getattr(sys, "frozen", False))
        if not allow_cleanup and runtime_dir == PYINSTALLER_RUNTIME_DIR:
            return CheckResult("런타임 임시 폴더", True, "소스 실행 중 점검 생략")
    if not runtime_dir.exists():
        return CheckResult("런타임 임시 폴더", True, "공용 런타임 임시 폴더 없음")

    active_dir = _active_pyinstaller_temp_dir(runtime_dir)
    current_time = time.time() if now is None else now
    removed: list[str] = []
    failed: list[str] = []
    recent_count = 0
    active_count = 0
    skipped_stale_count = 0

    for path in sorted(runtime_dir.glob("_MEI*")):
        if not path.is_dir():
            continue
        try:
            resolved = path.resolve()
        except OSError:
            failed.append(path.name)
            continue
        if active_dir is not None and resolved == active_dir:
            active_count += 1
            continue
        try:
            age_seconds = current_time - path.stat().st_mtime
        except OSError:
            failed.append(path.name)
            continue
        if age_seconds < max_age_seconds:
            recent_count += 1
            continue
        if not allow_cleanup:
            skipped_stale_count += 1
            continue
        try:
            shutil.rmtree(path)
            removed.append(path.name)
        except OSError as exc:
            failed.append(f"{path.name}: {exc}")

    if failed:
        extra = "" if len(failed) <= 3 else f" 외 {len(failed) - 3}개"
        return CheckResult(
            "런타임 임시 폴더",
            True,
            "Windows 보호 권한으로 정리할 수 없는 이전 실행 폴더를 보류했습니다: "
            + ", ".join(failed[:3])
            + extra
            + ". 프로그램 실행과 고객 데이터에는 영향이 없습니다.",
        )

    details: list[str] = []
    if removed:
        details.append(f"오래된 폴더 {len(removed)}개 정리")
    if active_count:
        details.append(f"실행 중 {active_count}개 보류")
    if recent_count:
        details.append(f"최근 폴더 {recent_count}개 보류")
    if skipped_stale_count:
        details.append(f"소스 실행 중 오래된 폴더 {skipped_stale_count}개 확인")
    return CheckResult("런타임 임시 폴더", True, ", ".join(details) if details else "정리할 런타임 임시 폴더 없음")


def _check_ocr_runtime_hygiene(install_dir: Path) -> CheckResult:
    ocr_dir = install_dir / "tools" / "ocr"
    if not ocr_dir.exists():
        return CheckResult("OCR 런타임 정리", True, "OCR 런타임 폴더 없음")
    problems: list[str] = []
    for relative in FORBIDDEN_OCR_RUNTIME_FILES:
        if (ocr_dir / relative).exists():
            problems.append(relative)
    for path in ocr_dir.rglob("*"):
        if path.is_file() and path.suffix.lower() in FORBIDDEN_OCR_RUNTIME_SUFFIXES:
            problems.append(path.relative_to(ocr_dir).as_posix())
    if problems:
        extra = "" if len(problems) <= 5 else f" 외 {len(problems) - 5}개"
        return CheckResult(
            "OCR 런타임 정리",
            False,
            "고객 런타임에 불필요한 OCR 훈련/문서 파일이 포함되어 있습니다: " + ", ".join(problems[:5]) + extra,
        )
    return CheckResult("OCR 런타임 정리", True, "고객 런타임 파일만 포함")


def _check_brand_assets(install_dir: Path) -> CheckResult:
    missing_or_empty = [
        relative
        for relative in REQUIRED_BRAND_ASSET_FILES
        if not (install_dir / relative).is_file() or (install_dir / relative).stat().st_size <= 0
    ]
    if missing_or_empty:
        return CheckResult(
            "브랜드 로고",
            False,
            "프로그램 상단 로고 파일이 없습니다: " + ", ".join(missing_or_empty),
        )
    return CheckResult("브랜드 로고", True, "채움랩 로고 파일 확인")


def _check_app_font(install_dir: Path) -> CheckResult:
    font_path = install_dir / APP_FONT_RELATIVE_PATH
    if not font_path.is_file() or font_path.stat().st_size <= 0:
        return CheckResult("기본 글꼴", False, f"채움랩 기본 글꼴 파일이 없습니다: {font_path}")
    digest = hashlib.sha256(font_path.read_bytes()).hexdigest()
    if digest != APP_FONT_SHA256:
        return CheckResult("기본 글꼴", False, "채움랩 기본 글꼴 파일 해시가 다릅니다.")
    if not register_bundled_font(base_dir=install_dir, install_dir=install_dir):
        return CheckResult("기본 글꼴", False, "채움랩 기본 글꼴을 Windows GDI에 등록하지 못했습니다.")
    return CheckResult("기본 글꼴", True, "머니그라피 Rounded 파일 및 앱 전용 등록 확인")


def _active_pyinstaller_temp_dir(install_dir: Path) -> Path | None:
    bundle_dir = getattr(sys, "_MEIPASS", None)
    if not bundle_dir:
        return None
    try:
        path = Path(str(bundle_dir)).resolve()
        if path.name.startswith("_MEI") and path.parent == install_dir.resolve():
            return path
    except OSError:
        return None
    return None


def _file_result(name: str, path: Path) -> CheckResult:
    if path.is_file() and path.stat().st_size > 0:
        return CheckResult(name, True, "확인됨")
    return CheckResult(name, False, f"파일이 없거나 비어 있습니다: {path}")


def _check_config(config_path: Path, results: list[CheckResult]) -> AppConfig | None:
    try:
        config = load_config(config_path)
    except BarcodeLabelAutomationError as exc:
        results.append(CheckResult("config.ini 설정", False, str(exc)))
        return None
    except Exception as exc:  # defensive boundary for customer-facing diagnostics
        results.append(CheckResult("config.ini 설정", False, f"읽기 실패: {exc}"))
        return None
    results.append(
        CheckResult(
            "config.ini 설정",
            True,
            (
                f"{config.printer.brand}/{config.printer.language}, "
                f"{config.label.width_mm}x{config.label.height_mm}mm, {config.label.dpi}dpi"
            ),
        )
    )
    return config


def _check_workbooks(data_dir: Path) -> list[CheckResult]:
    return [
        _check_workbook(data_dir / "barcode_db.xlsx", "바코드 DB"),
        _check_workbook(data_dir / "print_queue.xlsx", "인쇄 데이터"),
        _check_workbook(data_dir / "labels.xlsm", "엑셀 입력 파일"),
    ]


def _check_workbook(path: Path, name: str) -> CheckResult:
    if not path.exists():
        return CheckResult(name, False, f"파일이 없습니다: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=False)
        sheet_count = len(workbook.sheetnames)
        workbook.close()
    except Exception as exc:
        return CheckResult(name, False, f"엑셀 파일을 열 수 없습니다: {exc}")
    return CheckResult(name, True, f"시트 {sheet_count}개 확인")


def _check_print_data_readiness(data_dir: Path) -> CheckResult:
    queue_path = data_dir / "print_queue.xlsx"
    if not queue_path.exists():
        return CheckResult("인쇄 데이터 준비", False, f"print_queue.xlsx 파일이 없습니다: {queue_path}")
    try:
        rows = load_label_rows(queue_path)
    except Exception as exc:
        return CheckResult("인쇄 데이터 준비", False, f"print_queue.xlsx 필수 열을 읽을 수 없습니다: {exc}")

    if not rows:
        return CheckResult(
            "인쇄 데이터 준비",
            False,
            "인쇄할 유효한 행이 없습니다. 라벨출력관리.exe에서 DB를 연결하거나 인쇄 데이터를 저장하세요.",
        )

    errors: list[str] = []
    for index, row in enumerate(rows, start=1):
        barcode = str(row.get("barcode", "")).strip()
        if not barcode:
            errors.append(f"{index}행 바코드가 비어 있습니다.")

        raw_quantity = str(row.get("print_qty", "") or DEFAULT_PRINT_QTY).strip()
        try:
            quantity = int(float(raw_quantity))
        except ValueError:
            errors.append(f"{index}행 출력 매수는 숫자로 입력하세요.")
            continue
        if quantity < 1 or quantity > 1000:
            errors.append(f"{index}행 출력 매수는 1부터 1000 사이로 입력하세요.")

    if errors:
        extra = "" if len(errors) <= 5 else f" 외 {len(errors) - 5}개"
        return CheckResult("인쇄 데이터 준비", False, " ".join(errors[:5]) + extra)
    return CheckResult("인쇄 데이터 준비", True, f"출력 가능 행 {len(rows)}건")


def _check_default_template(path: Path) -> CheckResult:
    if not path.exists():
        return CheckResult("기본 템플릿", False, f"파일이 없습니다: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        return CheckResult("기본 템플릿", False, f"JSON 읽기 실패: {exc}")
    elements = payload.get("elements")
    if not isinstance(elements, list):
        return CheckResult("기본 템플릿", False, "elements 배열이 없습니다.")
    if elements:
        return CheckResult("기본 템플릿", False, "기본 템플릿 elements는 빈 배열이어야 합니다. 저장한 라벨은 .gblabel 파일로 따로 보관하세요.")
    return CheckResult("기본 템플릿", True, f"개체 {len(elements)}개")


def _check_writable_dir(data_dir: Path) -> CheckResult:
    try:
        data_dir.mkdir(parents=True, exist_ok=True)
        probe = data_dir / f".preflight_{uuid.uuid4().hex}.tmp"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
    except OSError as exc:
        return CheckResult("출력 폴더 쓰기 권한", False, str(exc))
    return CheckResult("출력 폴더 쓰기 권한", True, "고객 데이터 폴더 쓰기 가능")


def _check_release_manifest(install_dir: Path) -> CheckResult:
    ok, message = validate_release_manifest(install_dir)
    return CheckResult("배포 파일 목록", ok, message)


def _check_dry_run(install_dir: Path, data_dir: Path, config_path: Path, timeout_sec: int) -> CheckResult:
    runner_source = install_dir / "라벨작업실행기.exe"
    print_source = install_dir / "라벨출력엔진.exe"
    if not runner_source.exists() and not print_source.exists():
        return CheckResult("dry-run 출력", False, f"출력 실행 파일이 없습니다: {runner_source}")

    # A preflight check must not rewrite the release folder. The job runner
    # creates print_queue.xlsx from labels.xlsm and the engine creates out/*;
    # both are valid runtime work, but neither belongs in the package being
    # checked. Run the real executable against a small disposable copy instead.
    with tempfile.TemporaryDirectory(prefix="chaeumlab_preflight_") as temp_dir:
        isolated_dir = Path(temp_dir)
        try:
            if runner_source.exists():
                shutil.copy2(runner_source, isolated_dir / runner_source.name)
            if print_source.exists():
                shutil.copy2(print_source, isolated_dir / print_source.name)
            shutil.copy2(config_path, isolated_dir / config_path.name)
            label_source = next(
                (data_dir / name for name in ("labels.xlsm", "labels.xlsx") if (data_dir / name).is_file()),
                None,
            )
            if label_source is None:
                return CheckResult("dry-run 출력", False, f"labels.xlsm 또는 labels.xlsx 파일이 없습니다: {data_dir}")
            shutil.copy2(label_source, isolated_dir / label_source.name)
        except OSError as exc:
            return CheckResult("dry-run 출력", False, f"임시 점검 폴더 준비 실패: {exc}")

        runner_exe = isolated_dir / runner_source.name
        if runner_source.exists():
            command = [str(runner_exe), "DryRun", "-Quiet"]
        else:
            command = [str(isolated_dir / print_source.name), "--config", str(isolated_dir / config_path.name), "--dry-run", "--limit", "1"]
        try:
            completed = subprocess.run(
                command,
                cwd=str(isolated_dir),
                text=True,
                capture_output=True,
                timeout=timeout_sec,
                check=False,
            )
        except subprocess.TimeoutExpired:
            return CheckResult("dry-run 출력", False, f"{timeout_sec}초 안에 끝나지 않았습니다.")
        except OSError as exc:
            return CheckResult("dry-run 출력", False, f"실행 실패: {exc}")
        if completed.returncode != 0:
            message = (completed.stderr or completed.stdout).strip()
            return CheckResult("dry-run 출력", False, message or f"exit code {completed.returncode}")
        summary = (completed.stdout or "").strip().splitlines()
        return CheckResult("dry-run 출력", True, summary[0] if summary else "작업 실행기 dry-run 확인")


def _check_label_settings(config: AppConfig) -> CheckResult:
    if config.label.width_mm <= 0 or config.label.height_mm <= 0:
        return CheckResult("라벨 크기", False, "가로/세로는 0보다 커야 합니다.")
    if config.label.dpi not in {203, 300, 600}:
        return CheckResult("라벨 DPI", False, "지원 DPI는 203, 300, 600입니다.")
    return CheckResult("라벨 크기/DPI", True, "정상 범위")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Check customer deployment readiness.")
    parser.add_argument("--base-dir", default=None, help="Customer deployment folder to check")
    parser.add_argument("--skip-dry-run", action="store_true", help="Skip print_labels dry-run check")
    parser.add_argument("--no-report", action="store_true", help="Do not write customer_preflight_report.txt")
    parser.add_argument("--report-path", default=None, help="Override diagnostic report output path")
    parser.add_argument("--no-support-package", action="store_true", help="Do not write customer_support_package.zip")
    parser.add_argument("--support-package-path", default=None, help="Override support package output path")
    args = parser.parse_args(argv)

    report = run_preflight(args.base_dir, run_dry_run=not args.skip_dry_run)
    print_report(report)
    if not args.no_report:
        try:
            report_path = write_report(report, args.report_path)
            print()
            print(f"진단 보고서 저장: {report_path}")
        except OSError as exc:
            print()
            print(f"진단 보고서 저장 실패: {exc}", file=sys.stderr)
    if not args.no_support_package:
        try:
            package_path = write_support_package(report, args.support_package_path)
            print(f"지원 패키지 저장: {package_path}")
        except OSError as exc:
            print(f"지원 패키지 저장 실패: {exc}", file=sys.stderr)
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
